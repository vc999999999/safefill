"""No-roster, no-terminal SafeFill workflow regression."""
import base64
import json
import os
import re
import stat
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "safefill-fill" / "scripts"))

import collection
import fill


PNG_1X1 = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")


def submit(form, answers, out_dir, *, previous=None):
    path = out_dir.parent / (out_dir.name + "-answers.json")
    path.write_text(json.dumps(answers, ensure_ascii=False), encoding="utf-8")
    if os.name != "nt":
        path.chmod(0o600)
    result = fill.cmd_submit(SimpleNamespace(
        form=form, answers=str(path), out=None, out_dir=str(out_dir), credential=None,
        previous=str(previous) if previous else None, delete_answers=True,
    ))
    assert not path.exists()
    return result


def workbook_rows(path):
    workbook = load_workbook(path, read_only=True)
    rows = list(workbook.active.values)
    workbook.close()
    return rows


def test_reply_filename_is_readable_and_safe():
    assert re.fullmatch(r"张_三-[A-Z0-9]{6}\.yintian", fill.reply_filename(' 张/三:*? '))


def test_open_template_to_excel(tmp_path):
    initialized = tmp_path / "initialized.json"
    collection.cmd_init_config(SimpleNamespace(out=str(initialized), force=False, mode="open"))
    assert [field["id"] for field in collection.load_json(initialized)["fields"]] == ["name"]
    config = {
        "title": "员工联系方式",
        "purpose": "用于紧急联络",
        "deadline": "2098-12-31T23:59:59Z",
        "retention_until": "2099-01-31T23:59:59Z",
        "contact": "人事部 hr@example.com",
        "correction": "联系人事部后使用本人旧回执重新提交",
        "template_version": "1.0",
        "fields": [
            {"id": "name", "label": "姓名", "type": "text", "required": True, "sensitive": True},
            {"id": "phone", "label": "手机号", "type": "phone_cn", "required": True, "sensitive": True},
            {"id": "photo", "label": "照片", "type": "image_attachment", "required": False, "sensitive": True, "ocr_fields": []},
        ],
    }
    config_path = tmp_path / "collection.json"
    collection.dump_json(config_path, config)
    created = collection.create_task(None, config_path, tmp_path / "tasks", None, require_terminal=False, mode="open")
    task_dir = Path(created["task_dir"])
    private_blob = (task_dir / "private.pem.enc").read_bytes()
    assert private_blob.lstrip().startswith(b"{") and b"PRIVATE KEY" not in private_blob
    local_key = collection.local_task_secret_path(task_dir, created["task_id"])
    assert local_key.is_file() and task_dir not in local_key.parents
    if os.name != "nt":
        assert stat.S_IMODE(local_key.stat().st_mode) == 0o600

    photo = tmp_path / "photo.png"
    photo.write_bytes(PNG_1X1)
    incoming = tmp_path / "incoming"
    incoming.mkdir()
    replies = {}
    for name, phone, attachments in (
        ("张三", "13800138000", {"photo": str(photo)}),
        ("李四", "13900139000", {}),
    ):
        result = submit(created["form"], {"consent_confirmed": True, "values": {"name": name, "phone": phone}, "attachments": attachments}, incoming)
        reply = Path(result["out"])
        replies[name] = reply
        assert name in reply.name
        assert name not in reply.read_text(encoding="utf-8")

    result = collection.collect_open(task_dir, incoming, tmp_path / "result.xlsx")
    assert result["rows"] == 2 and result["excluded"] == 0 and result["attachments"] == 1
    rows = workbook_rows(result["xlsx"])
    assert rows[0] == ("姓名", "手机号", "照片")
    workbook = load_workbook(result["xlsx"])
    assert all(cell.data_type == "s" for cell in workbook.active[1])
    workbook.close()
    if os.name != "nt":
        assert stat.S_IMODE(Path(result["xlsx"]).stat().st_mode) == 0o600
        assert stat.S_IMODE(Path(result["attachments_dir"]).stat().st_mode) == 0o700
    records = {row[0]: row[1:] for row in rows[1:]}
    assert records["李四"] == ("13900139000", None)
    photo_ref = records["张三"][1]
    assert records["张三"][0] == "13800138000"
    assert (Path(result["xlsx"]).parent / photo_ref).read_bytes() == PNG_1X1

    correction_incoming = tmp_path / "correction-incoming"
    correction_incoming.mkdir()
    corrected = submit(
        created["form"],
        {"consent_confirmed": True, "values": {"name": "张三", "phone": "13700137000"}, "attachments": {"photo": str(photo)}},
        correction_incoming,
        previous=replies["张三"],
    )
    assert corrected["invite_id"] == json.loads(replies["张三"].read_bytes())["invite_id"]
    corrected_export = collection.collect_open(task_dir, correction_incoming, tmp_path / "corrected.xlsx")
    corrected_rows = workbook_rows(corrected_export["xlsx"])
    corrected_records = {row[0]: row[1:] for row in corrected_rows[1:]}
    assert corrected_records["张三"][0] == "13700137000"
    assert collection.status_task(task_dir)["total"] == 2

    handoff = tmp_path / "task.yintian-task"
    secret = collection.load_local_task_secret(task_dir, created["task_id"])
    collection.export_task(task_dir, handoff, handoff_password="handoff-password-2026")
    assert secret.encode() not in handoff.read_bytes()
    imported = collection.import_task(handoff, tmp_path / "imported", handoff_password="handoff-password-2026")
    imported_dir = Path(imported["task_dir"])
    assert collection.load_local_task_secret(imported_dir, created["task_id"]) == secret
    empty = tmp_path / "empty"
    empty.mkdir()
    imported_export = collection.collect_open(imported_dir, empty, tmp_path / "imported.xlsx")
    assert workbook_rows(imported_export["xlsx"])[0] == ("姓名", "手机号", "照片")

    if os.name != "nt":
        unsafe = tmp_path / "unsafe-answers.json"
        unsafe.write_text(json.dumps({"consent_confirmed": True, "values": {"name": "王五", "phone": "13600136000"}, "attachments": {}}), encoding="utf-8")
        unsafe.chmod(0o644)
        try:
            fill.cmd_submit(SimpleNamespace(form=created["form"], answers=str(unsafe), out=None, out_dir=str(tmp_path / "unsafe-out"), credential=None, previous=None))
        except fill.FillError as exc:
            assert "ANSWERS_PERMISSIONS" in str(exc)
        else:
            raise AssertionError("submit 必须拒绝权限过宽的临时明文")
        assert not unsafe.exists()

    invalid = {**config, "fields": [{**config["fields"][0], "required": False}, *config["fields"][1:]]}
    try:
        collection.validate_config(invalid, mode="open")
    except ValueError:
        pass
    else:
        raise AssertionError("开放模板必须拒绝非必填姓名")
    form = fill.load_form(created["form"])
    try:
        fill._check_values(form, {"name": "名" * (collection.MAX_NAME_CHARS + 1), "phone": "13800138000"})
    except fill.FillError:
        pass
    else:
        raise AssertionError("填写端必须拒绝超长姓名")


if __name__ == "__main__":
    test_reply_filename_is_readable_and_safe()
    with tempfile.TemporaryDirectory(prefix="safefill-open-") as directory:
        test_open_template_to_excel(Path(directory))
    print("open workflow ok")
