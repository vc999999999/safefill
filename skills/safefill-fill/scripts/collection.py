"""SafeFill · 端到端加密的私密信息收集管理 CLI。"""
from __future__ import annotations

import argparse
import base64
import csv
import getpass
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import shutil
import sqlite3
import string
import sys
import tempfile
import unicodedata
import zipfile
from contextlib import closing, contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path, PurePosixPath
from typing import Any

import secure_io
import evidence_routing

FORMAT_VERSION = "yintian-submission/2"
GROUP_FORMAT_VERSION = "yintian-submission/3"
OPEN_FORMAT_VERSION = "yintian-submission/4"
AUTH_VERSION = "hmac-sha256-token/1"
CREDENTIAL_FORMAT = "yintian-credential/1"
DB_VERSION = 1
LEGACY_FORMAT_VERSION = "yintian-submission/1"
LEGACY_TASK_PACKAGE_VERSION = "yintian-task/1"
TASK_PACKAGE_VERSION = "yintian-task/2"
ENCRYPTED_TASK_PACKAGE_VERSION = "yintian-task/3"
KEY_ENVELOPE_VERSION = "yintian-key/1"
FORM_FORMAT_VERSION = "yintian-form/1"
LEGACY_OPEN_FORM_FORMAT_VERSION = "yintian-form/3"
OPEN_REQUEST_FORMAT_VERSION = "yintian-request/1"
GROUP_INVITE_PREFIX = "GRP-"
OPEN_INVITE_PREFIX = "OPEN-"
TASK_MODES = {"directed", "group", "open"}
MASK_RULES = {"last4", "mid4"}
LEGACY_KEY_PEM_PREFIX = b"-----BEGIN ENCRYPTED PRIVATE KEY-----"
SCRYPT_N = 2**15
SCRYPT_R = 8
SCRYPT_P = 1
ALLOWED_TYPES = {"text", "phone_cn", "cn_id", "date", "address", "single_choice", "image_attachment", "pdf_attachment"}
ATTACHMENT_TYPES = {"image_attachment", "pdf_attachment"}
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TOTAL_BYTES = 15 * 1024 * 1024
MAX_PDF_PAGES = 20
PDF_RENDER_SCALE = 2
MAX_ENVELOPE_BYTES = 32 * 1024 * 1024
MAX_V3_ENVELOPE_BYTES = 200 * 1024 * 1024
V3_SNIFF_BYTES = 1024 * 1024
MAX_PACKAGE_FILES = 20_000
MAX_PACKAGE_BYTES = 2 * 1024 * 1024 * 1024
MAX_PACKAGE_MEMBER_BYTES = 128 * 1024 * 1024
MAX_IMAGE_PIXELS = int(os.getenv("YINTIAN_MAX_IMAGE_PIXELS", "40000000"))
MAX_FIELDS = 100
MAX_LABEL_CHARS = 200
MAX_VALUE_CHARS = 10_000
MAX_NAME_CHARS = 200
MAX_OPTIONS = 100
MAX_ATTACHMENTS = 60


def positive_int_env(name: str, default: int) -> int:
    try:
        value = int(os.getenv(name, ""))
    except ValueError:
        return default
    return value if value > 0 else default


MAX_VERSIONS_PER_INVITE = positive_int_env("YINTIAN_MAX_VERSIONS_PER_INVITE", 10)
MAX_OPEN_INVITES = positive_int_env("YINTIAN_MAX_OPEN_INVITES", 10_000)
MAX_TASK_SUBMISSION_BYTES = positive_int_env("YINTIAN_MAX_TASK_SUBMISSION_BYTES", 1024 * 1024 * 1024)
TASK_ID_RE = re.compile(r"^YT-[0-9]{8}-[A-Z0-9]{6}$")
INVITE_ID_RE = re.compile(r"^INV-[A-Z0-9]{10}$")
OPEN_INVITE_ID_RE = re.compile(r"^OPEN-[A-Z0-9]{16}$")
FIELD_ID_RE = re.compile(r"^[a-z][a-z0-9_]{1,63}$")
ILLEGAL_XML_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def terminal_text(value):
    return ''.join(char if unicodedata.category(char) not in {'Cc', 'Cf'} else ' ' for char in str(value))


def safe_filename_component(value: Any, fallback: str = "item", max_chars: int = 60, max_bytes: int = 120) -> str:
    text = unicodedata.normalize("NFKC", terminal_text(value))
    text = re.sub(r'[<>:"/\\|?*]+', "_", text).strip(" ._")[:max_chars].strip(" ._")
    while len(text.encode("utf-8")) > max_bytes:
        text = text[:-1]
    return text or fallback


def parse_time(value: str) -> datetime:
    value = value.strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        value += "T23:59:59+00:00"
    elif value.endswith("Z"):
        value = value[:-1] + "+00:00"
    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def normalize_submitted_at(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    try:
        parsed = parse_time(value)
    except (TypeError, ValueError):
        return None
    return parsed.replace(microsecond=0).isoformat().replace("+00:00", "Z")


def valid_date(value: str) -> bool:
    try:
        return datetime.strptime(value, "%Y-%m-%d").strftime("%Y-%m-%d") == value
    except (TypeError, ValueError):
        return False


def canonical(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(secure_io.read_bytes(path, MAX_PACKAGE_MEMBER_BYTES))


def atomic_write(path: Path, data: bytes) -> None:
    secure_io.atomic_write(path, data)


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(secure_io.read_bytes(path, MAX_V3_ENVELOPE_BYTES))


def dump_json(path: Path, value: Any) -> None:
    atomic_write(path, json.dumps(value, ensure_ascii=False, indent=2).encode("utf-8") + b"\n")


def random_id(prefix: str, length: int) -> str:
    alphabet = string.ascii_uppercase + string.digits
    return prefix + "".join(secrets.choice(alphabet) for _ in range(length))


def generate_password(length: int = 28) -> str:
    alphabet = string.ascii_letters + string.digits + "-_.~"
    while True:
        value = "".join(secrets.choice(alphabet) for _ in range(length))
        if any(c.islower() for c in value) and any(c.isupper() for c in value) and any(c.isdigit() for c in value):
            return value


def open_control_terminal():
    """打开控制终端用于显示一次性密码（POSIX 用 /dev/tty，Windows 用 CON）；打不开时返回 None。"""
    try:
        return open("CON" if os.name == "nt" else "/dev/tty", "w", encoding="utf-8", errors="replace")
    except OSError:
        return None


def print_once_secret(heading: str, secret: str, footnote: str) -> None:
    """一次性密码只写控制终端；控制终端不可用时回退 stderr，绝不写可能被管道采集的 stdout。"""
    stream = open_control_terminal()
    try:
        target = stream if stream is not None else sys.stderr
        print(f"\n{heading}", file=target)
        print(secret, file=target)
        print(footnote + "\n", file=target)
    finally:
        if stream is not None:
            stream.close()


def require_tty(action: str) -> None:
    if not (sys.stdin.isatty() and sys.stdout.isatty()):
        raise RuntimeError(f"{action} 必须在交互终端中运行；TTY 检查不能证明该终端未被 Agent 控制。")


def task_expired(task: dict[str, Any]) -> bool:
    return datetime.now(timezone.utc) > parse_time(task["retention_until"])


def task_late(task: dict[str, Any], received_at: str) -> bool:
    return parse_time(received_at) > parse_time(task["deadline"])


def load_task(task_dir: str | Path) -> tuple[Path, dict[str, Any]]:
    root = secure_io.checked_path(task_dir)
    task_path = root / "task.json"
    if not task_path.is_file():
        raise FileNotFoundError(f"不是有效任务目录: {root}")
    task = load_json(task_path)
    if not isinstance(task, dict) or not TASK_ID_RE.fullmatch(task.get("task_id", "")):
        raise ValueError("task.json 中的 task_id 无效")
    if task.get("format_version") not in {LEGACY_FORMAT_VERSION, FORMAT_VERSION, GROUP_FORMAT_VERSION, OPEN_FORMAT_VERSION}:
        raise ValueError("任务格式版本不受支持")
    if task.get("mode", "directed") not in TASK_MODES:
        raise ValueError("task.json 中的 mode 无效")
    if task.get("format_version") == GROUP_FORMAT_VERSION and (task_mode(task) != "group" or task.get("submission_auth") != AUTH_VERSION):
        raise ValueError("GROUP_AUTH_REQUIRED: 群发任务缺少认证配置")
    if task.get("format_version") == OPEN_FORMAT_VERSION and task_mode(task) != "open":
        raise ValueError("开放请求任务格式与 mode 不匹配")
    if task_mode(task) == "open" and task.get("format_version") != OPEN_FORMAT_VERSION:
        raise ValueError("开放请求任务必须使用当前提交协议")
    required = ("key_id", "title", "purpose", "deadline", "retention_until", "contact", "correction", "template_version", "schema_hash", "notice_hash", "fields")
    if any(key not in task for key in required) or not isinstance(task["fields"], list):
        raise ValueError("task.json 缺少必要配置")
    for key in ("title", "purpose", "deadline", "retention_until", "contact", "correction", "template_version"):
        if not isinstance(task[key], str) or not task[key] or len(task[key]) > MAX_VALUE_CHARS:
            raise ValueError(f"task.json 文本字段无效: {key}")
    if not isinstance(task["key_id"], str) or not re.fullmatch(r"[0-9a-f]{24}", task["key_id"]) or any(not isinstance(task[key], str) or not re.fullmatch(r"[0-9a-f]{64}", task[key]) for key in ("schema_hash", "notice_hash")):
        raise ValueError("task.json 公钥或摘要标识无效")
    if parse_time(task["retention_until"]) <= parse_time(task["deadline"]):
        raise ValueError("task.json 保存期限无效")
    if validate_field_definitions(task["fields"], task_mode(task)) != task["fields"]:
        raise ValueError("task.json 字段模板未规范化")
    if sha256_bytes(canonical(task["fields"])) != task["schema_hash"]:
        raise ValueError("task.json 字段模板哈希不匹配")
    notice = {key: task[key] for key in ("title", "purpose", "deadline", "retention_until", "contact", "correction")}
    if sha256_bytes(canonical(notice)) != task["notice_hash"]:
        raise ValueError("task.json 告知内容哈希不匹配")
    return root, task


def require_current_format(task: dict[str, Any], action: str) -> None:
    if task["format_version"] not in {FORMAT_VERSION, GROUP_FORMAT_VERSION, OPEN_FORMAT_VERSION}:
        raise RuntimeError(f"旧版任务不支持{action}；请新建 v2 任务继续收集")
    if task_mode(task) == "group" and task.get("format_version") != GROUP_FORMAT_VERSION:
        raise RuntimeError("GROUP_AUTH_REQUIRED: 旧群发任务没有个人认证，保留只读；请新建群发任务")


def task_mode(task: dict[str, Any]) -> str:
    """任务模式只从本地 task.json 读取（load_task 已校验取值），提交信封与载荷无法伪造。"""
    return task.get("mode", "directed")


def valid_invite_identifier(value):
    if not isinstance(value, str):
        return False
    if INVITE_ID_RE.fullmatch(value):
        return True
    if OPEN_INVITE_ID_RE.fullmatch(value):
        return True
    try:
        parse_group_employee_id(value)
        return True
    except ValueError:
        return False


def group_invite_id(employee_id: str) -> str:
    return GROUP_INVITE_PREFIX + employee_id


def parse_group_employee_id(invite_id: str) -> str:
    if not invite_id.startswith(GROUP_INVITE_PREFIX):
        raise ValueError("group 提交的 invite_id 必须以 GRP- 开头")
    employee_id = invite_id[len(GROUP_INVITE_PREFIX):]
    if not employee_id or employee_id != employee_id.strip() or len(employee_id) > 128 or any(ord(char) < 0x20 for char in employee_id):
        raise ValueError("group 提交的 invite_id 格式无效")
    return employee_id


@contextmanager
def task_lock(root: Path, *, migration: bool = False):
    with secure_io.file_lock(root / ".write.lock"):
        if not migration:
            with closing(connect_db(root)) as db:
                if db.execute("PRAGMA user_version").fetchone()[0] != DB_VERSION:
                    raise RuntimeError("MIGRATION_REQUIRED: 请先运行 collection.py migrate TASK_DIR")
        yield


def connect_db(root: Path) -> sqlite3.Connection:
    path = secure_io.checked_path(root / "state.sqlite3", root)
    for suffix in ("-wal", "-shm", "-journal"):
        secure_io.checked_path(str(path) + suffix, root)
    if not path.is_file():
        raise FileNotFoundError("DATABASE_MISSING: 任务数据库不存在")
    db = sqlite3.connect(path, timeout=5.0)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA busy_timeout=5000")
    db.execute("PRAGMA foreign_keys=ON")
    return db


def init_db(root: Path, rows: list[dict[str, str]]) -> None:
    db_path = root / "state.sqlite3"
    secure_io.atomic_write(db_path, b"", overwrite=False)
    with closing(connect_db(root)) as db, db:
        db.executescript(
            """
            CREATE TABLE invites (
                invite_id TEXT PRIMARY KEY,
                employee_id TEXT NOT NULL UNIQUE,
                name TEXT NOT NULL,
                token_hash TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'invited',
                current_submission_id INTEGER,
                created_at TEXT NOT NULL
            );
            CREATE TABLE submissions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                invite_id TEXT NOT NULL REFERENCES invites(invite_id),
                version INTEGER NOT NULL,
                sha256 TEXT NOT NULL UNIQUE,
                path TEXT NOT NULL,
                received_at TEXT NOT NULL,
                submitted_at TEXT,
                reviewed_at TEXT,
                late INTEGER NOT NULL DEFAULT 0,
                status TEXT NOT NULL,
                missing_fields TEXT NOT NULL DEFAULT '[]',
                conflict_fields TEXT NOT NULL DEFAULT '[]',
                attachment_count INTEGER NOT NULL DEFAULT 0,
                consent_confirmed INTEGER NOT NULL DEFAULT 0,
                UNIQUE(invite_id, version)
            );
            CREATE TABLE audit (
                id INTEGER PRIMARY KEY,
                action TEXT NOT NULL,
                submission_id INTEGER,
                at TEXT NOT NULL,
                result TEXT NOT NULL,
                reason TEXT NOT NULL,
                operator TEXT NOT NULL DEFAULT ''
            );
            PRAGMA user_version=1;
            """
        )
        db.executemany(
            "INSERT INTO invites(invite_id,employee_id,name,token_hash,status,created_at) VALUES(:invite_id,:employee_id,:name,:token_hash,'invited',:created_at)",
            rows,
        )
    os.chmod(root / "state.sqlite3", 0o600)


def audit(db, action, submission_id=None, result="ok", reason="", operator=""):
    db.execute("INSERT INTO audit(action,submission_id,at,result,reason,operator) VALUES(?,?,?,?,?,?)",
               (action, submission_id, now_iso(), result, reason, operator))


def cmd_migrate(args):
    root, task = load_task(args.task_dir)
    require_current_format(task, "迁移")
    with task_lock(root, migration=True), closing(connect_db(root)) as db:
        version = db.execute("PRAGMA user_version").fetchone()[0]
        if version == DB_VERSION:
            return {"schema_version": version, "changed": False}
        if version != 0:
            raise RuntimeError("SCHEMA_UNSUPPORTED: 不支持此数据库版本")
        backup_path = secure_io.checked_path(root / "migration-backup.sqlite3", root)
        if backup_path.exists():
            backup_path = root / ('migration-backup-' + secrets.token_hex(6) + '.sqlite3')
        secure_io.atomic_write(backup_path, b"", overwrite=False)
        with sqlite3.connect(backup_path) as backup:
            db.backup(backup)
        try:
            db.execute("BEGIN IMMEDIATE")
            db.execute("CREATE TABLE audit(id INTEGER PRIMARY KEY, action TEXT NOT NULL, submission_id INTEGER, at TEXT NOT NULL, result TEXT NOT NULL, reason TEXT NOT NULL, operator TEXT NOT NULL DEFAULT '')")
            db.execute("UPDATE submissions SET status='needs_review',conflict_fields=? WHERE status='verified'", (json.dumps(['runtime:migration_recheck']),))
            db.execute("UPDATE invites SET status='needs_review' WHERE current_submission_id IN (SELECT id FROM submissions WHERE status='needs_review')")
            db.execute(f"PRAGMA user_version={DB_VERSION}")
            audit(db, "migrate", reason="schema_1")
            db.commit()
        except BaseException:
            db.rollback()
            raise
    return {"schema_version": DB_VERSION, "changed": True}


def cmd_doctor(args):
    import importlib.util
    import importlib.metadata
    modules = {'cryptography': 'cryptography', 'PIL': 'pillow', 'openpyxl': 'openpyxl', 'mcp': 'mcp',
               'pypdfium2': 'pypdfium2', 'rapidocr_openvino': 'rapidocr-openvino', 'tkinter': None}
    checks = {}
    for module, package in modules.items():
        available = importlib.util.find_spec(module) is not None
        if module == 'tkinter' and available:
            try:
                __import__('tkinter')
            except ImportError:
                available = False
        try:
            version = importlib.metadata.version(package) if available and package else None
        except importlib.metadata.PackageNotFoundError:
            version = None
        checks[module] = {'available': available, 'version': version}
    vault_path = getattr(args, 'vault', None)
    permissions_ok = None
    if vault_path:
        path = secure_io.checked_path(vault_path)
        permissions_ok = path.is_dir() and (os.name == 'nt' or path.stat().st_mode & 0o077 == 0)
    return {'python': sys.version.split()[0], 'python_supported': sys.version_info >= (3, 11),
            'dependencies': checks, 'vault_permissions_ok': permissions_ok,
            'recognition': {'local_optional': True, 'agent': 'host_provided_candidates', 'manual': checks['tkinter']['available'], 'requires_api_key': False},
            'core_ready': sys.version_info >= (3, 11) and all(checks[k]['available'] for k in ('cryptography', 'PIL', 'openpyxl', 'mcp'))}


def default_config() -> dict[str, Any]:
    deadline = datetime.now(timezone.utc).replace(microsecond=0) + timedelta(days=30)
    return {
        "title": "基础身份信息收集",
        "purpose": "请填写本次工作所需的基础身份信息。请将用途修改为具体、必要的业务目的。",
        "deadline": deadline.isoformat().replace("+00:00", "Z"),
        "retention_until": (deadline + timedelta(days=30)).isoformat().replace("+00:00", "Z"),
        "contact": "请填写联系人和联系方式",
        "correction": "如需更正，请联系任务发起人并使用原邀请重新提交。",
        "template_version": "1.0",
        "fields": [
            {"id": "name", "label": "姓名", "type": "text", "required": True, "sensitive": False},
            {"id": "phone", "label": "手机号", "type": "phone_cn", "required": True, "sensitive": True},
            {"id": "id_number", "label": "身份证号", "type": "cn_id", "required": True, "sensitive": True},
            {"id": "address", "label": "住址", "type": "address", "required": True, "sensitive": True},
            {"id": "id_front", "label": "身份证正面", "type": "image_attachment", "required": True, "sensitive": True, "ocr_fields": ["name", "id_number", "address"], "ocr_backend": "auto"},
            {"id": "id_back", "label": "身份证反面", "type": "image_attachment", "required": True, "sensitive": True, "ocr_fields": [], "ocr_backend": "auto"},
        ],
    }


def validate_field_definitions(raw_fields: Any, mode: str) -> list[dict[str, Any]]:
    if mode not in TASK_MODES:
        raise ValueError("字段定义的任务模式无效")
    if not isinstance(raw_fields, list) or not (1 <= len(raw_fields) <= MAX_FIELDS):
        raise ValueError(f"fields 必须包含 1 到 {MAX_FIELDS} 个字段")
    seen, labels, fields = set(), set(), []
    for raw in raw_fields:
        if not isinstance(raw, dict):
            raise ValueError("每个字段定义必须是对象")
        field = dict(raw)
        field_id, field_type = field.get("id", ""), field.get("type", "")
        if not isinstance(field_id, str) or not FIELD_ID_RE.fullmatch(field_id) or field_id in seen:
            raise ValueError(f"字段 id 无效或重复: {field_id}")
        if field_type not in ALLOWED_TYPES:
            raise ValueError(f"不支持的字段类型: {field_type}")
        if not isinstance(field.get("label"), str) or not field["label"] or len(field["label"]) > MAX_LABEL_CHARS or ILLEGAL_XML_RE.search(field["label"]):
            raise ValueError(f"字段缺少 label: {field_id}")
        if field["label"] in labels:
            raise ValueError(f"字段 label 重复: {field['label']}")
        if any(key in field and not isinstance(field[key], bool) for key in ("required", "sensitive")):
            raise ValueError(f"字段 required/sensitive 必须是布尔值: {field_id}")
        if field_type == "single_choice":
            options = field.get("options")
            if not isinstance(options, list) or not (1 <= len(options) <= MAX_OPTIONS) or any(not isinstance(value, str) or not value or len(value) > MAX_LABEL_CHARS or ILLEGAL_XML_RE.search(value) for value in options) or len(set(options)) != len(options):
                raise ValueError(f"单选字段 options 无效: {field_id}")
        elif "options" in field:
            raise ValueError(f"非单选字段不能包含 options: {field_id}")
        if "multiple" in field and (field_type not in ATTACHMENT_TYPES or not isinstance(field["multiple"], bool)):
            raise ValueError(f"multiple 只允许附件字段使用布尔值: {field_id}")
        field["required"] = bool(field.get("required"))
        field["sensitive"] = bool(field.get("sensitive", field_type in {"phone_cn", "cn_id", "address"} or field_type in ATTACHMENT_TYPES))
        if field_type in ATTACHMENT_TYPES:
            field["multiple"] = bool(field.get("multiple", False))
        fields.append(field)
        seen.add(field_id)
        labels.add(field["label"])
    definitions = {field["id"]: field for field in fields}
    if definitions.get("name", {}).get("type") != "text" or definitions["name"].get("required") is not True:
        raise ValueError("任务字段必须包含必填 text 字段 id=name")
    if mode == "group" and (definitions.get("employee_id", {}).get("type") != "text" or definitions["employee_id"].get("required") is not True):
        raise ValueError("group 模式必须包含必填 text 字段 id=employee_id")
    scalar_fields = {field["id"] for field in fields if field["type"] not in ATTACHMENT_TYPES}
    for field in fields:
        if "ocr_fields" in field:
            bindings = field["ocr_fields"]
            if field["type"] not in ATTACHMENT_TYPES or not isinstance(bindings, list) or any(not isinstance(key, str) or key not in scalar_fields for key in bindings) or len(set(bindings)) != len(bindings):
                raise ValueError("OCR_BINDING_INVALID: ocr_fields 必须引用不重复的表单值字段")
    evidence_routing.validate_fields(fields)
    return fields


def validate_config(config: dict[str, Any], mode: str = "directed") -> dict[str, Any]:
    if mode not in TASK_MODES:
        raise ValueError("mode 必须是 directed、group 或 open")
    if not isinstance(config, dict):
        raise ValueError("配置必须是 JSON 对象")
    for key in ("title", "purpose", "deadline", "retention_until", "contact", "correction", "template_version", "fields"):
        if not config.get(key):
            raise ValueError(f"配置缺少必填项: {key}")
    for key in ("title", "purpose", "deadline", "retention_until", "contact", "correction", "template_version"):
        if not isinstance(config[key], str) or len(config[key]) > MAX_VALUE_CHARS:
            raise ValueError(f"配置字段必须是长度不超过 {MAX_VALUE_CHARS} 的文本: {key}")
    deadline, retention = parse_time(config["deadline"]), parse_time(config["retention_until"])
    if deadline <= datetime.now(timezone.utc):
        raise ValueError("deadline 必须晚于当前时间")
    if retention <= deadline:
        raise ValueError("retention_until 必须晚于 deadline")
    defaults = default_config()
    for key in ("purpose", "contact", "correction"):
        if config[key] == defaults[key]:
            raise ValueError(f"请先把 {key} 的占位内容改为真实、具体的信息")
    fields = validate_field_definitions(config["fields"], mode)
    result = dict(config)
    result["fields"] = fields
    return result


def read_roster(path: Path, *, allow_empty: bool = False) -> list[dict[str, str]]:
    with io.StringIO(secure_io.read_bytes(path).decode("utf-8-sig"), newline="") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames or not {"employee_id", "name"}.issubset(reader.fieldnames):
            raise ValueError("名单 CSV 必须包含 employee_id,name 表头")
        rows = []
        ids = set()
        for line, raw in enumerate(reader, 2):
            employee_id, name = (raw.get("employee_id") or "").strip(), (raw.get("name") or "").strip()
            if not employee_id or not name:
                raise ValueError(f"名单第 {line} 行缺少 employee_id 或 name")
            if employee_id in ids:
                raise ValueError(f"名单 employee_id 重复: {employee_id}")
            ids.add(employee_id)
            rows.append({"employee_id": employee_id, "name": name})
    if not rows and not allow_empty:
        raise ValueError("名单为空")
    return rows


def derive_scrypt_key(password: str, salt: bytes, n: int = SCRYPT_N, r: int = SCRYPT_R, p: int = SCRYPT_P) -> bytes:
    from cryptography.hazmat.primitives.kdf.scrypt import Scrypt

    return Scrypt(salt=salt, length=32, n=n, r=r, p=p).derive(password.encode("utf-8"))


def validate_kdf_params(kdf: Any) -> tuple[bytes, int, int, int]:
    if not isinstance(kdf, dict) or kdf.get("name") != "scrypt":
        raise ValueError("KDF 参数不受支持")
    try:
        salt = base64.b64decode(kdf.get("salt", ""), validate=True)
        n, r, p = int(kdf["n"]), int(kdf["r"]), int(kdf["p"])
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError("KDF 参数无效") from exc
    if len(salt) != 16 or not (2**10 <= n <= 2**16 and n & (n - 1) == 0) or not (1 <= r <= 8 and 1 <= p <= 2):
        raise ValueError("KDF 参数超出允许范围")
    return salt, n, r, p


def aes_gcm_seal(format_name: str, plaintext: bytes, password: str, aad: bytes | None = None) -> dict[str, Any]:
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    salt, nonce = secrets.token_bytes(16), secrets.token_bytes(12)
    key = derive_scrypt_key(password, salt)
    ciphertext = AESGCM(key).encrypt(nonce, plaintext, aad if aad is not None else format_name.encode("utf-8"))
    return {
        "format": format_name,
        "kdf": {"name": "scrypt", "salt": base64.b64encode(salt).decode("ascii"), "n": SCRYPT_N, "r": SCRYPT_R, "p": SCRYPT_P},
        "cipher": "AES-256-GCM",
        "nonce": base64.b64encode(nonce).decode("ascii"),
        "ciphertext": base64.b64encode(ciphertext).decode("ascii"),
    }


def aes_gcm_open(envelope: dict[str, Any], password: str, aad: bytes) -> bytes:
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    if not isinstance(envelope, dict) or envelope.get("cipher") != "AES-256-GCM":
        raise ValueError("加密信封格式不受支持")
    salt, n, r, p = validate_kdf_params(envelope.get("kdf"))
    try:
        nonce = base64.b64decode(envelope["nonce"], validate=True)
        ciphertext = base64.b64decode(envelope["ciphertext"], validate=True)
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError("加密信封缺少 nonce 或密文") from exc
    if len(nonce) != 12:
        raise ValueError("加密信封 nonce 无效")
    key = derive_scrypt_key(password, salt, n, r, p)
    return AESGCM(key).decrypt(nonce, ciphertext, aad)


def encrypt_private_key(private_der: bytes, password: str) -> bytes:
    envelope = aes_gcm_seal(KEY_ENVELOPE_VERSION, private_der, password)
    return json.dumps(envelope, ensure_ascii=False, indent=2).encode("utf-8") + b"\n"


def decrypt_private_key(envelope: dict[str, Any], password: str) -> bytes:
    if envelope.get("format") != KEY_ENVELOPE_VERSION:
        raise ValueError("私钥信封格式不受支持")
    return aes_gcm_open(envelope, password, KEY_ENVELOPE_VERSION.encode("utf-8"))


def valid_private_key_blob(data: bytes) -> bool:
    blob = data.lstrip()
    if blob.startswith(LEGACY_KEY_PEM_PREFIX):
        return True
    if blob.startswith(b"{"):
        try:
            envelope = json.loads(blob)
        except (json.JSONDecodeError, UnicodeDecodeError):
            return False
        return isinstance(envelope, dict) and envelope.get("format") == KEY_ENVELOPE_VERSION
    return False


def generate_keys(password: str) -> tuple[bytes, bytes, str]:
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa

    private = rsa.generate_private_key(public_exponent=65537, key_size=3072)
    public = private.public_key()
    public_pem = public.public_bytes(serialization.Encoding.PEM, serialization.PublicFormat.SubjectPublicKeyInfo)
    public_der = public.public_bytes(serialization.Encoding.DER, serialization.PublicFormat.SubjectPublicKeyInfo)
    private_der = private.private_bytes(serialization.Encoding.DER, serialization.PrivateFormat.PKCS8, serialization.NoEncryption())
    return public_pem, encrypt_private_key(private_der, password), hashlib.sha256(public_der).hexdigest()[:24]


def local_task_secret_path(root: Path, task_id: str) -> Path:
    return secure_io.checked_path(root.parent / ".safefill-keys" / f"{task_id}.key", root.parent)


def save_local_task_secret(root: Path, task_id: str, secret: str) -> Path:
    path = local_task_secret_path(root, task_id)
    secure_io.atomic_write(path, secret.encode("ascii"), overwrite=False)
    return path


def load_local_task_secret(root: Path, task_id: str) -> str:
    path = local_task_secret_path(root, task_id)
    if not path.is_file() or (os.name != "nt" and path.stat().st_mode & 0o077):
        raise RuntimeError("LOCAL_KEY_UNAVAILABLE: 开放任务的本地密钥缺失或权限不安全")
    try:
        secret = secure_io.read_bytes(path, 256).decode("ascii")
    except (OSError, UnicodeDecodeError) as exc:
        raise RuntimeError("LOCAL_KEY_UNAVAILABLE: 开放任务的本地密钥无法读取") from exc
    if not re.fullmatch(r"[A-Za-z0-9_.~-]{28,128}", secret):
        raise RuntimeError("LOCAL_KEY_INVALID: 开放任务的本地密钥格式无效")
    return secret


def csv_text(value: str) -> str:
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def cmd_init_config(args) -> dict[str, Any]:
    out = Path(args.out).expanduser()
    if out.exists() and not args.force:
        raise FileExistsError(f"文件已存在: {out}")
    config = default_config()
    mode = getattr(args, 'mode', 'open')
    if mode == 'open':
        config['correction'] = '如需更正，请联系任务联系人，并使用本人上一次 .yintian 重新提交。'
        config['fields'] = [config['fields'][0]]
    elif mode == 'group':
        config['fields'].insert(0, {'id': 'employee_id', 'label': '工号', 'type': 'text', 'required': True, 'sensitive': False})
    dump_json(out, config)
    return {"config": str(out.resolve())}


def create_task(roster_path: Path | None, config_path: Path, out_parent: Path, password: str | None, require_terminal: bool = True, mode: str = "directed") -> dict[str, Any]:
    if require_terminal:
        require_tty("create")
    config = validate_config(load_json(config_path), mode=mode)
    roster = [] if mode == "open" else read_roster(roster_path)
    if mode == 'group':
        for person in roster:
            parse_group_employee_id(group_invite_id(person['employee_id']))
    task_id = "YT-" + datetime.now().strftime("%Y%m%d") + "-" + random_id("", 6)
    root = secure_io.checked_path(out_parent) / task_id
    if root.exists():
        raise FileExistsError(f"任务目录已存在: {root}")
    root.mkdir(parents=True, mode=0o700)
    saved_secret_path = None
    try:
        (root / "invites").mkdir(mode=0o700)
        (root / "submissions").mkdir(mode=0o700)
        (root / "reports").mkdir(mode=0o700)

        local_secret = secrets.token_urlsafe(32) if mode == "open" else None
        if mode != "open" and not isinstance(password, str):
            raise ValueError("非开放任务必须提供任务密码")
        public_pem, private_pem, key_id = generate_keys(local_secret if local_secret is not None else password)
        schema_hash = sha256_bytes(canonical(config["fields"]))
        notice = {key: config[key] for key in ("title", "purpose", "deadline", "retention_until", "contact", "correction")}
        notice_hash = sha256_bytes(canonical(notice))
        task = {
            "format_version": GROUP_FORMAT_VERSION if mode == "group" else OPEN_FORMAT_VERSION if mode == "open" else FORMAT_VERSION,
            "mode": mode,
            "task_id": task_id,
            "key_id": key_id,
            "title": config["title"],
            "purpose": config["purpose"],
            "deadline": config["deadline"],
            "retention_until": config["retention_until"],
            "contact": config["contact"],
            "correction": config["correction"],
            "template_version": config["template_version"],
            "schema_hash": schema_hash,
            "notice_hash": notice_hash,
            "fields": config["fields"],
            "created_at": now_iso(),
        }
        if mode == "group":
            task["submission_auth"] = AUTH_VERSION
        dump_json(root / "task.json", task)
        atomic_write(root / "public.pem", public_pem)
        atomic_write(root / "private.pem.enc", private_pem)
        roster_bytes = secure_io.read_bytes(roster_path) if roster_path else b"employee_id,name\n"
        atomic_write(root / "roster.csv", roster_bytes)

        db_rows, index_rows = [], []
        if mode == "open":
            request = {"format": OPEN_REQUEST_FORMAT_VERSION, "kind": "agent_request", "target_skill": "safefill-fill", **task, "public_key_pem": public_pem.decode("ascii")}
            dump_json(root / "REQUEST.yintian-request", request)
        elif mode == "group":
            form = {"format": "yintian-form/2", **task, "public_key_pem": public_pem.decode("ascii")}
            dump_json(root / "FORM.yintian-form", form)
            (root / "credentials").mkdir(mode=0o700)
            for item in roster:
                invite_id = group_invite_id(item["employee_id"])
                parse_group_employee_id(invite_id)  # 创建期闭环校验：名单工号必须能构成合法的 GRP- 标识
                token = secrets.token_urlsafe(32)
                dump_json(root / "credentials" / (invite_id + ".yintian-credential"), {
                    "format": CREDENTIAL_FORMAT, "task_id": task_id, "key_id": key_id,
                    "invite_id": invite_id, "employee_id": item["employee_id"], "name": item["name"],
                    "invite_token": token, "schema_hash": schema_hash,
                })
                db_rows.append({**item, "invite_id": invite_id, "token_hash": sha256_bytes(token.encode()), "created_at": now_iso()})
                index_rows.append({**item, "invite_id": invite_id, "invite_file": "FORM.yintian-form"})
        else:
            for item in roster:
                invite_id = random_id("INV-", 10)
                invite_token = secrets.token_urlsafe(32)
                invite_name = invite_id + ".yintian-form"
                invite_config = dict(task)
                invite_config.update({"invite_id": invite_id, "invite_token": invite_token, "name": item["name"], "public_key_pem": public_pem.decode("ascii")})
                dump_json(root / "invites" / invite_name, {"format": FORM_FORMAT_VERSION, **invite_config})
                created_at = now_iso()
                db_rows.append({**item, "invite_id": invite_id, "token_hash": sha256_bytes(invite_token.encode()), "created_at": created_at})
                index_rows.append({**item, "invite_id": invite_id, "invite_file": f"invites/{invite_name}"})
        index_path = root / "invite-index.csv"
        index_buffer = io.StringIO()
        writer = csv.DictWriter(index_buffer, fieldnames=["employee_id", "name", "invite_id", "invite_file"])
        writer.writeheader()
        writer.writerows([{key: csv_text(value) for key, value in row.items()} for row in index_rows])
        atomic_write(index_path, index_buffer.getvalue().encode("utf-8-sig"))
        init_db(root, db_rows)
        if local_secret is not None:
            saved_secret_path = save_local_task_secret(root, task_id, local_secret)
        artifact = root / ("REQUEST.yintian-request" if mode == "open" else "FORM.yintian-form")
        return {"task_id": task_id, "task_dir": str(root), "request": str(artifact) if mode == "open" else None, "form": str(artifact) if mode == "group" else None, "invite_count": len(index_rows)}
    except BaseException:
        shutil.rmtree(root, ignore_errors=True)
        if saved_secret_path is not None:
            saved_secret_path.unlink(missing_ok=True)
        raise


def cmd_create(args) -> dict[str, Any]:
    require_tty("create")
    password = generate_password()
    result = create_task(Path(args.roster), Path(args.config), Path(args.out), password, require_terminal=False, mode=getattr(args, "mode", "directed"))
    print_once_secret("任务密码（仅显示一次，丢失不可恢复）：", password, "请通过独立安全渠道交给授权处理人员，不要写入任务目录或聊天提示词。")
    return result


def cmd_create_open(args) -> dict[str, Any]:
    return create_task(None, Path(args.config), Path(args.out), None, require_terminal=False, mode="open")


def validate_envelope_header(envelope: dict[str, Any], task: dict[str, Any]) -> str:
    required = {"format_version", "task_id", "invite_id", "schema_hash", "key_id", "algorithms", "encrypted_key_b64", "iv_b64", "ciphertext_b64"}
    allowed = required | ({"auth_tag"} if task_mode(task) == "group" else set())
    if not isinstance(envelope, dict) or set(envelope) != allowed:
        raise ValueError("提交包字段集合无效")
    for key in required - {"algorithms"}:
        if not isinstance(envelope.get(key), str) or not envelope[key]:
            raise ValueError(f"提交包缺少字段: {key}")
    if envelope["format_version"] != task["format_version"] or envelope["task_id"] != task["task_id"]:
        raise ValueError("提交包属于其他任务或格式版本")
    if envelope.get("algorithms") != {"content": "AES-256-GCM", "key_wrap": "RSA-OAEP-3072-SHA256"}:
        raise ValueError("提交包算法套件不受支持")
    if envelope["schema_hash"] != task["schema_hash"] or envelope["key_id"] != task["key_id"]:
        raise ValueError("提交包字段模板或公钥不匹配")
    invite_id = envelope["invite_id"]
    if task_mode(task) == "group":
        parse_group_employee_id(invite_id)
    elif task_mode(task) == "open":
        if not OPEN_INVITE_ID_RE.fullmatch(invite_id):
            raise ValueError("open 提交的 invite_id 格式无效")
    elif not INVITE_ID_RE.fullmatch(invite_id):
        raise ValueError("invite_id 格式无效")
    for key in ("encrypted_key_b64", "iv_b64", "ciphertext_b64"):
        base64.b64decode(envelope[key], validate=True)
    return invite_id


def submission_auth_tag(envelope, token_hash):
    content = {key: value for key, value in envelope.items() if key != "auth_tag"}
    return hmac.new(bytes.fromhex(token_hash), canonical(content), hashlib.sha256).hexdigest()


def verify_submission_auth(task, envelope, invite):
    if task.get("submission_auth") == AUTH_VERSION:
        tag = envelope.get("auth_tag")
        if not isinstance(tag, str) or not secrets.compare_digest(tag, submission_auth_tag(envelope, invite["token_hash"])):
            raise ValueError("SUBMISSION_AUTH_FAILED: 提交者认证失败")


def ingest_task(task_dir: str | Path, submissions_dir: str | Path) -> dict[str, Any]:
    root, task = load_task(task_dir)
    require_current_format(task, "接收")
    if task_expired(task):
        raise RuntimeError("任务已超过保存期限，停止接收新提交")
    source = secure_io.checked_path(submissions_dir)
    if not source.is_dir():
        raise NotADirectoryError(source)
    paths = []
    for path in source.iterdir():
        if path.suffix == ".yintian":
            paths.append(path)
            if len(paths) > MAX_PACKAGE_FILES:
                raise ValueError(f"INBOX_LIMIT: 单次收件最多处理 {MAX_PACKAGE_FILES} 个回执")
    summary = {"accepted": 0, "duplicates": 0, "rejected": 0, "errors": []}
    with task_lock(root), closing(connect_db(root)) as db, db:
        stored_bytes = 0
        for stored in db.execute("SELECT path FROM submissions"):
            stored_path = secure_io.checked_path(root / stored["path"], root)
            stored_bytes += stored_path.stat().st_size
        for index, path in enumerate(sorted(paths), start=1):
            digest = None
            created_path = None
            db.execute("SAVEPOINT receive_one")
            try:
                if not path.is_file() or path.is_symlink():
                    raise ValueError("提交项不是普通文件")
                if path.stat().st_size > MAX_ENVELOPE_BYTES:
                    raise ValueError("提交包超过 32MB 上限")
                raw = secure_io.read_bytes(path, MAX_ENVELOPE_BYTES)
                digest = sha256_bytes(raw)
                if db.execute("SELECT 1 FROM submissions WHERE sha256=?", (digest,)).fetchone():
                    summary["duplicates"] += 1
                    db.execute("RELEASE receive_one")
                    continue
                envelope = json.loads(raw)
                invite_id = validate_envelope_header(envelope, task)
                invite = db.execute("SELECT * FROM invites WHERE invite_id=?", (invite_id,)).fetchone()
                if not invite and task_mode(task) == "open":
                    if db.execute("SELECT COUNT(*) FROM invites").fetchone()[0] >= MAX_OPEN_INVITES:
                        raise ValueError(f"OPEN_INVITE_LIMIT: 开放任务最多接收 {MAX_OPEN_INVITES} 个不同回执编号")
                    db.execute(
                        "INSERT INTO invites(invite_id,employee_id,name,token_hash,status,created_at) VALUES(?,?,?,'','invited',?)",
                        (invite_id, invite_id, "", now_iso()),
                    )
                    invite = db.execute("SELECT * FROM invites WHERE invite_id=?", (invite_id,)).fetchone()
                if not invite:
                    raise ValueError("invite_id 不属于当前任务")
                verify_submission_auth(task, envelope, invite)
                existing_versions = db.execute("SELECT COUNT(*) FROM submissions WHERE invite_id=?", (invite_id,)).fetchone()[0]
                if existing_versions >= MAX_VERSIONS_PER_INVITE:
                    raise ValueError(f"该邀请的提交版本数已达上限 {MAX_VERSIONS_PER_INVITE}，拒绝继续存储")
                if stored_bytes + len(raw) > MAX_TASK_SUBMISSION_BYTES:
                    raise ValueError("TASK_STORAGE_LIMIT: 任务密文总量达到安全上限")
                version = db.execute("SELECT COALESCE(MAX(version),0)+1 FROM submissions WHERE invite_id=?", (invite_id,)).fetchone()[0]
                target_dir = secure_io.checked_path(root / "submissions" / invite_id, root)
                target_dir.mkdir(parents=True, exist_ok=True, mode=0o700)
                target = target_dir / f"v{version:04d}_{digest[:12]}.yintian"
                if target.exists() and sha256_file(target) != digest:
                    raise ValueError("RECOVERY_CONFLICT: 收件路径已存在不同内容")
                if not target.exists():
                    created_path = target
                atomic_write(target, raw)
                received_at = now_iso()
                db.execute(
                    "INSERT INTO submissions(invite_id,version,sha256,path,received_at,late,status) VALUES(?,?,?,?,?,?,?)",
                    (invite_id, version, digest, target.relative_to(root).as_posix(), received_at, int(task_late(task, received_at)), "submitted"),
                )
                db.execute("UPDATE invites SET status='submitted' WHERE invite_id=?", (invite_id,))
                audit(db, "ingest", db.execute("SELECT last_insert_rowid()").fetchone()[0])
                db.execute("RELEASE receive_one")
                summary["accepted"] += 1
                stored_bytes += len(raw)
            except Exception as exc:
                db.execute("ROLLBACK TO receive_one")
                db.execute("RELEASE receive_one")
                if created_path is not None:
                    created_path.unlink(missing_ok=True)
                summary["rejected"] += 1
                summary["errors"].append({"file_ref": digest[:12] if digest else f"entry-{index}", "error": type(exc).__name__,
                                          "code": "INGEST_IO" if isinstance(exc, OSError) else "SUBMISSION_REJECTED",
                                          "retryable": isinstance(exc, OSError), "next_action": "retry" if isinstance(exc, OSError) else "check_invitation"})
    return summary


def cmd_ingest(args) -> dict[str, Any]:
    return ingest_task(args.task_dir, args.submissions_dir)


def unlock_private_key(root: Path, password: str | None):
    from cryptography.hazmat.primitives import serialization

    data = secure_io.read_bytes(root / "private.pem.enc", 128 * 1024)
    blob = data.lstrip()
    if blob.startswith(b"-----BEGIN PRIVATE KEY-----"):
        raise ValueError("未加密私钥不受支持")
    if blob.startswith(b"{"):
        if password is None:
            task = load_task(root)[1]
            if task_mode(task) != "open":
                raise ValueError("任务私钥需要密码")
            password = load_local_task_secret(root, task["task_id"])
        private_der = decrypt_private_key(json.loads(data), password)
        return serialization.load_der_private_key(private_der, password=None)
    if not blob.startswith(LEGACY_KEY_PEM_PREFIX):
        raise ValueError("私钥文件既不是 yintian-key/1 信封，也不是加密的 PKCS#8 PEM，拒绝解锁")
    return serialization.load_pem_private_key(data, password=password.encode())


def aad_for(envelope: dict[str, Any]) -> bytes:
    return canonical([envelope["format_version"], envelope["task_id"], envelope["invite_id"], envelope["schema_hash"], envelope["key_id"]])


def decrypt_envelope(root: Path, task: dict[str, Any], path: Path, private_key, expected_invite_id: str) -> dict[str, Any]:
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.asymmetric import padding
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    path = secure_io.checked_path(path, root)
    submissions_root = secure_io.checked_path(root / "submissions", root)
    if submissions_root not in path.parents or not path.is_file():
        raise ValueError("提交路径不在任务 submissions 目录内")
    envelope = load_json(path)
    invite_id = validate_envelope_header(envelope, task)
    if invite_id != expected_invite_id:
        raise ValueError("提交包与数据库邀请不匹配")
    encrypted_key = base64.b64decode(envelope["encrypted_key_b64"], validate=True)
    aes_key = private_key.decrypt(encrypted_key, padding.OAEP(mgf=padding.MGF1(algorithm=hashes.SHA256()), algorithm=hashes.SHA256(), label=None))
    plaintext = AESGCM(aes_key).decrypt(base64.b64decode(envelope["iv_b64"], validate=True), base64.b64decode(envelope["ciphertext_b64"], validate=True), aad_for(envelope))
    payload = json.loads(plaintext)
    if not isinstance(payload, dict):
        raise ValueError("解密载荷必须是对象")
    payload_fields = {"format_version", "task_id", "invite_id", "schema_hash", "notice_hash", "template_version", "submitted_at", "consent_confirmed", "values", "attachments"}
    if task_mode(task) != "open":
        payload_fields.add("invite_token")
    if set(payload) - payload_fields - {"agent_ocr", "agent_ocr_confirmed"}:
        raise ValueError("解密载荷包含协议外字段")
    for key in ("format_version", "task_id", "invite_id", "schema_hash"):
        if payload.get(key) != envelope[key]:
            raise ValueError(f"解密载荷 {key} 与外层不一致")
    return payload


def normalize_value(field_type: str, value: str) -> str:
    value = str(value or "").strip()
    if field_type == "phone_cn":
        value = re.sub(r"[\s-]", "", value)
        value = re.sub(r"^\+?86", "", value)
    if field_type == "cn_id":
        value = value.upper()
    return value


def validate_scalar_values(values: Any, allowed_ids: set[str]) -> None:
    if not isinstance(values, dict) or len(values) > MAX_FIELDS:
        raise ValueError("values 必须是字段数量受限的对象")
    if set(values) - allowed_ids:
        raise ValueError("PAYLOAD_FIELDS_INVALID: 提交含模板外字段")
    for field_id, value in values.items():
        limit = MAX_NAME_CHARS if field_id == "name" else MAX_VALUE_CHARS
        if not isinstance(value, str) or len(value) > limit or ILLEGAL_XML_RE.search(value):
            raise ValueError(f"字段必须是长度不超过 {limit} 的文本: {field_id}")


def validate_payload(task: dict[str, Any], invite: sqlite3.Row, payload: dict[str, Any]) -> tuple[list[str], list[str], list[dict[str, Any]]]:
    from ocr_matcher import validate_chinese_id

    required = {"notice_hash", "template_version", "submitted_at", "consent_confirmed", "values", "attachments"}
    if task_mode(task) != "open":
        required.add("invite_token")
    if not isinstance(payload, dict) or not required.issubset(payload):
        raise ValueError("解密载荷字段集合无效")
    missing, conflicts, attachment_items = [], [], []
    if payload.get("notice_hash") != task["notice_hash"]:
        conflicts.append("notice_hash")
    if payload.get("template_version") != task["template_version"]:
        conflicts.append("template_version")
    if payload.get("consent_confirmed") is not True:
        conflicts.append("consent")
    token = payload.get("invite_token", "")
    if task_mode(task) != "open":
        if not isinstance(token, str) or not secrets.compare_digest(sha256_bytes(token.encode()), invite["token_hash"]):
            raise ValueError("邀请认证令牌无效")
    if normalize_submitted_at(payload.get("submitted_at")) is None:
        conflicts.append("submitted_at")
    values, attachments = payload.get("values", {}), payload.get("attachments", {})
    if not isinstance(attachments, dict) or len(attachments) > MAX_FIELDS:
        raise ValueError("values/attachments 格式无效")
    value_ids = {f['id'] for f in task['fields'] if f['type'] not in ATTACHMENT_TYPES}
    attachment_ids = {f['id'] for f in task['fields'] if f['type'] in ATTACHMENT_TYPES}
    validate_scalar_values(values, value_ids)
    if set(attachments) - attachment_ids:
        raise ValueError('PAYLOAD_FIELDS_INVALID: 提交含模板外字段')
    total_size = 0
    total_attachments = 0
    for field in task["fields"]:
        field_id, field_type = field["id"], field["type"]
        if field_type in ATTACHMENT_TYPES:
            items = attachments.get(field_id, [])
            if not isinstance(items, list) or len(items) > 20:
                raise ValueError(f"附件字段格式无效: {field_id}")
            total_attachments += len(items)
            if total_attachments > MAX_ATTACHMENTS:
                raise ValueError(f"附件总数超过 {MAX_ATTACHMENTS} 个")
            if field.get("required") and not items:
                missing.append(field_id)
            if not field.get("multiple") and len(items) > 1:
                conflicts.append(field_id)
            for item in items:
                if not isinstance(item, dict):
                    raise ValueError(f"附件项目格式无效: {field_id}")
                name = item.get("name", "")
                if not isinstance(name, str) or len(name) > 255:
                    raise ValueError(f"附件文件名无效: {field_id}")
                raw = base64.b64decode(item.get("data_b64", ""), validate=True)
                mime = item.get("type", "")
                if len(raw) != int(item.get("size", -1)) or len(raw) > MAX_FILE_BYTES or sha256_bytes(raw) != item.get("sha256"):
                    raise ValueError(f"附件大小或哈希无效: {field_id}")
                if field_type == "image_attachment" and mime not in {"image/jpeg", "image/png", "image/webp"}:
                    raise ValueError(f"图片类型不支持: {mime}")
                if field_type == "pdf_attachment" and mime != "application/pdf":
                    raise ValueError("PDF 附件类型无效")
                if field_type == "image_attachment":
                    from PIL import Image

                    try:
                        with Image.open(io.BytesIO(raw)) as image:
                            expected = {"image/jpeg": "JPEG", "image/png": "PNG", "image/webp": "WEBP"}[mime]
                            if image.format != expected or image.width * image.height > MAX_IMAGE_PIXELS:
                                raise ValueError(f"图片内容或像素尺寸无效: {field_id}")
                            image.verify()
                    except (OSError, Image.DecompressionBombError):
                        raise ValueError('ATTACHMENT_INVALID: 图片数据无效') from None
                elif not raw.startswith(b"%PDF-"):
                    raise ValueError("PDF 文件头无效")
                total_size += len(raw)
                attachment_items.append({"field_id": field_id, "name": name, "type": mime, "data": raw})
        else:
            value = normalize_value(field_type, values.get(field_id, ""))
            if field.get("required") and not value:
                missing.append(field_id)
            if value and field_type == "phone_cn" and not re.fullmatch(r"1[3-9]\d{9}", value):
                conflicts.append(field_id)
            if value and field_type == "cn_id" and not validate_chinese_id(value):
                conflicts.append(field_id)
            if value and field_type == "date" and not valid_date(value):
                conflicts.append(field_id)
            if value and field_type == "single_choice" and value not in field.get("options", []):
                conflicts.append(field_id)
    if total_size > MAX_TOTAL_BYTES:
        raise ValueError("附件总大小超过 15MB")
    if task_mode(task) != "open" and normalize_value("text", values.get("name", "")) != invite["name"]:
        conflicts.append("name_roster_mismatch")
    if task_mode(task) == "group" and normalize_value("text", values.get("employee_id", "")) != invite["employee_id"]:
        conflicts.append("employee_id_roster_mismatch")
    if "agent_ocr" in payload:
        evidence_routing.validate_result(task, attachment_items, payload["agent_ocr"])
        if payload.get("agent_ocr_confirmed") is not True:
            conflicts.append("agent_ocr_consent")
    elif "agent_ocr_confirmed" in payload:
        conflicts.append("agent_ocr_consent")
    return sorted(set(missing)), sorted(set(conflicts)), attachment_items


def ocr_attachment(item: dict[str, Any]) -> tuple[str, list[str]]:
    try:
        import ocr_engine
    except Exception as exc:
        return "", [f"ocr_unavailable:{type(exc).__name__}"]
    try:
        if item["type"].startswith("image/"):
            result = ocr_engine.scan_bytes(item["data"])
            flags = ["ocr_low_confidence"] if any(block.get("low_confidence") for block in result["blocks"]) else []
            text = result["full_text"]
            if not text.strip():
                flags.append("ocr_no_text")
            return text, flags
        if item["type"] == "application/pdf":
            try:
                import pypdfium2 as pdfium
            except ImportError:
                return "", ["pdf_ocr_unavailable"]
            document = pdfium.PdfDocument(item["data"])
            flags = []
            if len(document) > MAX_PDF_PAGES:
                flags.append("pdf_page_limit")
            texts = []
            for page_index in range(min(len(document), MAX_PDF_PAGES)):
                page = document[page_index]
                width, height = page.get_size()
                rendered_w, rendered_h = width * PDF_RENDER_SCALE, height * PDF_RENDER_SCALE
                if rendered_w * rendered_h > MAX_IMAGE_PIXELS:
                    flags.append("pdf_pixel_limit")
                    continue
                bitmap = page.render(scale=PDF_RENDER_SCALE)
                image = bitmap.to_pil()
                buffer = io.BytesIO()
                image.save(buffer, format="PNG")
                result = ocr_engine.scan_bytes(buffer.getvalue())
                texts.append(result["full_text"])
                if any(block.get("low_confidence") for block in result["blocks"]):
                    flags.append("ocr_low_confidence")
            text = "\n".join(texts)
            if not text.strip():
                flags.append("ocr_no_text")
            return text, sorted(set(flags))
    except Exception as exc:
        return "", [f"ocr_error:{type(exc).__name__}"]
    return "", []


def is_identity_attachment_field(field_id: str) -> bool:
    """附件字段按下划线分词后含整词 front/back 才视为证件正反面（避免误伤 backdrop、feedback_scan 等）。"""
    return any(part in {"front", "back"} for part in field_id.split("_"))


def compare_ocr(payload, attachments, field_defs, *, provenance=None):
    """Check each explicitly bound piece of evidence; no empty-set/any-match approval."""
    from ocr_matcher import extract_from_text

    definitions = {f["id"]: f for f in field_defs}
    flags = []
    for item in attachments:
        definition = definitions.get(item["field_id"])
        if definition is None:
            continue
        binding = evidence_routing.bindings(definition, field_defs)
        if not binding:
            continue
        text, warnings, agent_candidates, source, reason = evidence_routing.select(payload, item, definition, ocr_attachment)
        if provenance is not None:
            provenance.append({"field_id": item['field_id'], "source": source, "reason": reason})
        if source == 'manual':
            flags.append(f"ocr:manual_required:{item['field_id']}")
            continue
        if source == 'agent':
            flags.append(f"ocr:agent_confirmation:{item['field_id']}")
        for warning in warnings:
            kind = "runtime:ocr" if any(k in warning for k in ("unavailable", "error")) else "ocr:quality"
            flags.append(f"{kind}:{item['field_id']}")
        candidates = {}
        extract_from_text(text, "attachment", candidates)
        for field_id in binding:
            field = definitions[field_id]
            value = normalize_value(field["type"], payload.get("values", {}).get(field_id, ""))
            if not value:
                continue
            if agent_candidates is not None:
                found_values = agent_candidates.get(field_id, [])
                valid = {normalize_value(field['type'], v) for v in found_values if evidence_routing.candidate_valid(field, v)}
                if any(not evidence_routing.candidate_valid(field, v) for v in found_values):
                    flags.append(f"ocr:invalid:{field_id}")
            else:
                extracted = "name" if field_id == "name" else {"cn_id": "id_number", "address": "address", "phone_cn": "phone"}.get(field["type"])
                if extracted is None:
                    flags.append(f"ocr:unsupported:{field_id}")
                    continue
                found = candidates.get(extracted, [])
                valid = {normalize_value(field["type"], candidate.value) for candidate in found if candidate.valid is not False}
                if any(candidate.valid is False for candidate in found):
                    flags.append(f"ocr:invalid:{field_id}")
            if not valid:
                flags.append(f"ocr:missing:{field_id}")
            elif len(valid) > 1:
                flags.append(f"ocr:ambiguous:{field_id}")
            elif value not in valid:
                flags.append(f"ocr:conflict:{field_id}")
    return sorted(set(flags))


def stored_payload(root, task, row, private_key):
    path = secure_io.checked_path(root / row["path"], root)
    if sha256_file(path) != row["sha256"]:
        raise ValueError("CIPHERTEXT_CHANGED: 已接收的文件发生变化")
    payload = decrypt_envelope(root, task, path, private_key, row["invite_id"])
    return payload


def review_task(task_dir, password, retry_needs_review=False, invite_id=None):
    from cryptography.exceptions import InvalidTag
    root, task = load_task(task_dir)
    require_current_format(task, "复核")
    if task_expired(task):
        raise RuntimeError("TASK_EXPIRED: 任务已超过保存期限")
    try:
        private_key = unlock_private_key(root, password)
    except Exception as exc:
        raise RuntimeError("KEY_UNLOCK_FAILED: 任务密码错误或私钥损坏") from exc
    summary = {"verified": 0, "needs_review": 0, "invalid": 0}
    with task_lock(root), closing(connect_db(root)) as db, db:
        states = "('submitted','needs_review')" if retry_needs_review else "('submitted')"
        query = f"SELECT s.*,i.employee_id,i.name,i.token_hash FROM submissions s JOIN invites i ON i.invite_id=s.invite_id WHERE s.status IN {states}"
        params = []
        if invite_id:
            query += " AND s.invite_id=?"
            params.append(invite_id)
        # Only the newest version remains actionable; resolved history is immutable.
        query += " AND (s.status='submitted' OR s.id=(SELECT MAX(n.id) FROM submissions n WHERE n.invite_id=s.invite_id)) ORDER BY s.id"
        for row in db.execute(query, params).fetchall():
            if task_expired(load_task(root)[1]):
                raise RuntimeError("TASK_EXPIRED: 任务已超过保存期限")
            missing, conflicts, attachments, payload, provenance = [], [], [], {}, []
            try:
                payload = stored_payload(root, task, row, private_key)
                missing, conflicts, attachments = validate_payload(task, row, payload)
                if not missing and not conflicts:
                    conflicts = compare_ocr(payload, attachments, task["fields"], provenance=provenance)
                state = "needs_review" if missing or conflicts else "verified"
            except (OSError, ImportError):
                state, conflicts = "needs_review", ["runtime:dependency_or_io"]
            except (ValueError, InvalidTag, TypeError, KeyError):
                state, conflicts = "invalid", ["ciphertext_or_payload_invalid"]
            except Exception:
                state, conflicts = "needs_review", ["runtime:review"]
            db.execute(
                "UPDATE submissions SET submitted_at=?,reviewed_at=?,status=?,missing_fields=?,conflict_fields=?,attachment_count=?,consent_confirmed=? WHERE id=?",
                (normalize_submitted_at(payload.get("submitted_at")), now_iso(), state, json.dumps(missing), json.dumps(conflicts), len(attachments), int(payload.get("consent_confirmed") is True), row["id"]))
            if state != "invalid":
                db.execute("UPDATE invites SET current_submission_id=? WHERE invite_id=? AND (current_submission_id IS NULL OR current_submission_id<=?)", (row["id"], row["invite_id"], row["id"]))
                if task_mode(task) == "open":
                    db.execute("UPDATE invites SET name=? WHERE invite_id=?", (normalize_value("text", payload.get("values", {}).get("name", "")), row["invite_id"]))
            db.execute("UPDATE invites SET status=? WHERE invite_id=? AND ?=(SELECT MAX(id) FROM submissions WHERE invite_id=?)", (state, row["invite_id"], row["id"], row["invite_id"]))
            audit(db, "ocr_route", row["id"], json.dumps(provenance, separators=(",", ":")), "recognition_route")
            audit(db, "review", row["id"], state, ",".join(missing + conflicts))
            summary[state] += 1
    return summary


def cmd_review(args):
    require_tty("review")
    return review_task(args.task_dir, getpass.getpass("任务密码: "),
                       getattr(args, "retry_needs_review", False), getattr(args, "invite", None))


def cmd_decide(args):
    require_tty("decide")
    root, task = load_task(args.task_dir)
    require_current_format(task, "人工复核")
    if task_expired(task):
        raise RuntimeError("TASK_EXPIRED: 任务已到期")
    if not re.fullmatch(r"[A-Za-z0-9_.@-]{1,64}", args.operator):
        raise ValueError("OPERATOR_INVALID: 操作者标识限字母、数字、_.@-")
    private_key = unlock_private_key(root, getpass.getpass("任务密码: "))
    with task_lock(root), closing(connect_db(root)) as db:
        row = db.execute("SELECT s.*,i.name,i.employee_id,i.token_hash FROM submissions s JOIN invites i ON i.invite_id=s.invite_id WHERE s.invite_id=? ORDER BY s.version DESC LIMIT 1", (args.invite_id,)).fetchone()
        if row is None or row['version'] != args.version or row['status'] in {'verified_manual', 'returned'}:
            raise RuntimeError("STATE_CHANGED: 版本不符或已人工结案")
        snapshot = export_snapshot(root)
        issues = json.loads(row['conflict_fields'])
        if args.action == 'confirm':
            payload = stored_payload(root, task, row, private_key)
            missing, hard_conflicts, attachments = validate_payload(task, row, payload)
            if row['status'] != 'needs_review' or missing or hard_conflicts or not issues or any(not item.startswith('ocr:') for item in issues):
                raise RuntimeError("MANUAL_NOT_ALLOWED: 人工只能裁定 OCR 问题；输入、认证和运行错误不能放行")
    if args.action == 'confirm':
        from review_evidence import confirm_evidence
        candidate_args = {'agent_candidates': payload['agent_ocr']['items']} if payload.get('agent_ocr') else {}
        if not confirm_evidence(payload['values'], attachments, issues, **candidate_args):
            raise RuntimeError("CANCELLED: 未完成逐项确认")
        state, reason = 'verified_manual', 'evidence_checked'
    else:
        if input(f"退回版本 {args.version}，输入任务 ID {task['task_id']}: ").strip() != task['task_id']:
            raise RuntimeError('CANCELLED: 已取消退回')
        state, reason = 'returned', 'correction_requested'
    with task_lock(root), closing(connect_db(root)) as db, db:
        if task_expired(load_task(root)[1]) or snapshot != export_snapshot(root):
            raise RuntimeError('STATE_CHANGED: 任务已变化或到期，请刷新后重试')
        if args.action == 'confirm':
            stored_payload(root, task, row, private_key)
        db.execute('UPDATE submissions SET status=?,reviewed_at=? WHERE id=?', (state, now_iso(), row['id']))
        db.execute('UPDATE invites SET status=?,current_submission_id=? WHERE invite_id=?', (state, row['id'], args.invite_id))
        for code in issues if args.action == 'confirm' else [reason]:
            audit(db, 'manual_' + args.action, row['id'], state, code, args.operator)
    return {'version': args.version, 'status': state}


def report_rows(root):
    with closing(connect_db(root)) as db:
        rows = db.execute("""
            SELECT i.employee_id,i.name,i.invite_id,s.id AS submission_id,
                   COALESCE(s.status,'invited') AS status,
                   s.version AS submission_version,s.submitted_at,s.reviewed_at,s.received_at,s.late,
                   s.missing_fields,s.conflict_fields,s.attachment_count,s.consent_confirmed,
                   a.version AS approved_version,a.status AS approved_status,
                   (SELECT COUNT(*) FROM submissions n WHERE n.invite_id=i.invite_id
                    AND n.id=s.id AND n.status IN ('submitted','needs_review')) AS pending_count
            FROM invites i
            LEFT JOIN submissions s ON s.id=(SELECT MAX(n.id) FROM submissions n WHERE n.invite_id=i.invite_id)
            LEFT JOIN submissions a ON a.id=(SELECT MAX(n.id) FROM submissions n WHERE n.invite_id=i.invite_id AND n.status IN ('verified','verified_manual'))
            ORDER BY i.employee_id
        """).fetchall()
        routes = {}
        if db.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='audit'").fetchone():
            routes = {r['submission_id']: json.loads(r['result']) for r in db.execute("SELECT submission_id,result FROM audit WHERE id IN (SELECT MAX(id) FROM audit WHERE action='ocr_route' GROUP BY submission_id)")}
    output = []
    for row in rows:
        conflicts = json.loads(row["conflict_fields"] or "[]")
        output.append({
            "employee_id": row["employee_id"], "name": row["name"], "status": row["status"],
            "submission_version": row["submission_version"], "latest_version": row["submission_version"],
            "approved_version": row["approved_version"], "received_at": row["received_at"],
            "pending_count": row["pending_count"],
            "approval_method": "manual" if row["approved_status"] == "verified_manual" else "automatic" if row["approved_status"] else None,
            "retryable": row["status"] == "needs_review" and any(c.startswith("runtime:") for c in conflicts),
            "submitted_at": row["submitted_at"], "reviewed_at": row["reviewed_at"],
            "late": bool(row["late"]), "missing_fields": json.loads(row["missing_fields"] or "[]"),
            "conflict_fields": conflicts, "attachment_count": row["attachment_count"] or 0,
            "consent_confirmed": bool(row["consent_confirmed"]),
            "recognition_sources": routes.get(row["submission_id"], []),
        })
    return output


def status_task(task_dir):
    root, task = load_task(task_dir)
    rows = report_rows(root)
    counts = {}
    for row in rows:
        counts[row["status"]] = counts.get(row["status"], 0) + 1
    with closing(connect_db(root)) as db:
        schema_version = db.execute("PRAGMA user_version").fetchone()[0]
    return {"task_id": task["task_id"], "task_status": "expired" if task_expired(task) else "active",
            "deadline": task["deadline"], "retention_until": task["retention_until"], "counts": counts,
            "total": len(rows), "pending_count": sum(row["pending_count"] for row in rows),
            "manual_review_count": counts.get("needs_review", 0),
            "latest_received_at": max((row["received_at"] for row in rows if row["received_at"]), default=None),
            "schema_version": schema_version, "migration_required": schema_version != DB_VERSION}


def cmd_status(args):
    return status_task(args.task_dir)


def _write_reports(task_dir: str | Path, formats: list[str]) -> dict[str, str]:
    root, task = load_task(task_dir)
    if task_expired(task):
        raise RuntimeError("任务已超过保存期限，停止生成报告；请执行清理")
    unknown = set(formats) - {"xlsx", "json"}
    if unknown:
        raise ValueError(f"不支持的报告格式: {sorted(unknown)}")
    rows = report_rows(root)
    reports = root / "reports"
    reports.mkdir(exist_ok=True)
    written = {}
    if "json" in formats:
        path = reports / "progress.json"
        dump_json(path, {"task_id": task["task_id"], "generated_at": now_iso(), "rows": rows})
        written["json"] = str(path)
    if "xlsx" in formats:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        path = reports / "progress.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "收集进度"
        columns = ["employee_id", "name", "status", "submission_version", "submitted_at", "reviewed_at", "late", "missing_fields", "conflict_fields", "attachment_count", "consent_confirmed", "latest_version", "approved_version", "pending_count", "approval_method", "retryable", "received_at", "recognition_sources"]
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row_index, row in enumerate(rows, start=2):
            for column_index, key in enumerate(columns, start=1):
                value = json.dumps(row[key], ensure_ascii=False) if key == "recognition_sources" else ",".join(row[key]) if isinstance(row[key], list) else row[key]
                cell = sheet.cell(row_index, column_index, value)
                if isinstance(value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        buffer = io.BytesIO()
        workbook.save(buffer)
        atomic_write(path, buffer.getvalue())
        written["xlsx"] = str(path)
    return written


def write_reports(task_dir, formats):
    root, _ = load_task(task_dir)
    with task_lock(root):
        result = _write_reports(root, formats)
        with closing(connect_db(root)) as db, db:
            audit(db, "report")
        return result


def cmd_report(args) -> dict[str, Any]:
    formats = [part.strip().lower() for value in args.formats for part in value.split(",") if part.strip()]
    return write_reports(args.task_dir, sorted(set(formats)))


def current_submission(db: sqlite3.Connection, invite_id: str) -> sqlite3.Row:
    row = db.execute(
        "SELECT s.*,i.employee_id,i.name,i.token_hash FROM invites i JOIN submissions s ON s.id=i.current_submission_id AND s.invite_id=i.invite_id WHERE i.invite_id=?",
        (invite_id,),
    ).fetchone()
    if not row:
        raise ValueError("该邀请没有可查看的有效提交")
    return row


def cmd_reveal(args) -> dict[str, Any] | None:
    require_tty("reveal")
    root, task = load_task(args.task_dir)
    require_current_format(task, "查看明文")
    if task_expired(task):
        raise RuntimeError("任务已超过保存期限，禁止查看明文")
    password = getpass.getpass("任务密码: ")
    try:
        private_key = unlock_private_key(root, password)
    except Exception as exc:
        raise RuntimeError("任务密码错误或私钥损坏") from exc
    with closing(connect_db(root)) as db, db:
        row = current_submission(db, args.invite_id)
    payload = decrypt_envelope(root, task, root / row["path"], private_key, args.invite_id)
    validate_payload(task, row, payload)
    terminal_text = lambda value: re.sub(r"[\x00-\x1f\x7f-\x9f]", " ", str(value))
    print("\n=== 授权明文查看 ===")
    print(f"employee_id: {terminal_text(row['employee_id'])}\nname: {terminal_text(row['name'])}\ninvite_id: {args.invite_id}")
    for field in task["fields"]:
        if field["type"] not in ATTACHMENT_TYPES:
            print(f"{terminal_text(field['label'])}: {terminal_text(payload.get('values', {}).get(field['id'], ''))}")
    print("附件:")
    for field_id, items in payload.get("attachments", {}).items():
        for item in items:
            print(f"- {terminal_text(field_id)}: {terminal_text(item.get('name'))} ({terminal_text(item.get('type'))}, {terminal_text(item.get('size'))} bytes)")
    input("\n查看完毕后按 Enter 清空当前屏幕与终端回滚缓冲区…")
    print("\033[2J\033[3J\033[H", end="", flush=True)
    return None


def parse_export_fields(task: dict[str, Any], raw_fields: str | None) -> list[str]:
    if not raw_fields or not raw_fields.strip():
        raise ValueError("必须用 --fields 显式声明导出字段白名单（逗号分隔的字段 id）")
    field_defs = {field["id"]: field for field in task["fields"]}
    field_ids = []
    for part in raw_fields.split(","):
        field_id = part.strip()
        if not field_id or field_id in {"employee_id", "name"}:
            continue  # employee_id/name 默认附带，无需在白名单中重复声明
        field = field_defs.get(field_id)
        if field is None:
            raise ValueError(f"导出字段不在任务字段中: {field_id}")
        if field["type"] in ATTACHMENT_TYPES:
            raise ValueError(f"附件字段不支持明文导出: {field_id}")
        if field_id in field_ids:
            raise ValueError(f"导出字段重复: {field_id}")
        field_ids.append(field_id)
    if not field_ids:
        raise ValueError("导出字段白名单为空（employee_id/name 默认附带，无需声明）")
    return field_ids


def parse_mask_specs(specs: list[str] | None) -> dict[str, str]:
    masks = {}
    for spec in specs or []:
        field_id, sep, rule = spec.partition("=")
        field_id, rule = field_id.strip(), rule.strip()
        if not sep or not FIELD_ID_RE.fullmatch(field_id) or rule not in MASK_RULES:
            raise ValueError(f"脱敏规则无效（应为 字段id=last4|mid4）: {spec}")
        masks[field_id] = rule
    return masks


def mask_value(rule: str, value: Any) -> str:
    text = str(value)
    if rule == "last4":
        return "*" * (len(text) - 4) + text[-4:] if len(text) > 4 else "*" * len(text)
    if rule == "mid4":
        return text[:3] + "****" + text[-4:] if len(text) > 7 else "*" * len(text)
    raise ValueError(f"不支持的脱敏规则: {rule}")


def collect_clear_rows(root: Path, task: dict[str, Any], private_key, field_ids: list[str], masks: dict[str, str]) -> list[dict[str, str]]:
    """只解密每人最新且已通过的提交，重新做硬性校验并只保留白名单字段。"""
    field_defs = {field["id"]: field for field in task["fields"]}
    with closing(connect_db(root)) as db, db:
        rows = db.execute(
            "SELECT s.*,i.employee_id,i.name,i.token_hash FROM submissions s JOIN invites i ON i.invite_id=s.invite_id WHERE s.id=(SELECT MAX(n.id) FROM submissions n WHERE n.invite_id=s.invite_id) AND s.status IN ('verified','verified_manual') ORDER BY i.employee_id"
        ).fetchall()
    output = []
    for row in rows:
        payload = stored_payload(root, task, row, private_key)
        missing, conflicts, _ = validate_payload(task, row, payload)
        if missing or conflicts:
            raise ValueError("EXPORT_VALIDATION_FAILED: 已通过记录的业务校验失败，停止导出")
        values = payload.get("values", {})
        record = {}
        identities = ("employee_id", "name")
        for column in identities:
            if task_mode(task) == "open":
                definition = field_defs.get(column)
                raw = normalize_value(definition["type"], values.get(column, "")) if definition else ""
            else:
                raw = row[column]
            rule = masks.get(column)
            record[column] = mask_value(rule, raw) if rule and raw else raw
        for field_id in field_ids:
            value = normalize_value(field_defs[field_id]["type"], values.get(field_id, ""))
            rule = masks.get(field_id)
            record[field_id] = mask_value(rule, value) if rule and value else value
        output.append(record)
    return output


def collect_open_rows(root: Path, task: dict[str, Any], private_key, attachment_stage: Path | None, attachment_dir_name: str) -> tuple[list[dict[str, str]], int]:
    field_defs = {field["id"]: field for field in task["fields"]}
    with closing(connect_db(root)) as db:
        db_rows = db.execute(
            "SELECT s.*,i.employee_id,i.name,i.token_hash FROM submissions s JOIN invites i ON i.invite_id=s.invite_id WHERE s.id=(SELECT MAX(n.id) FROM submissions n WHERE n.invite_id=s.invite_id) AND s.status IN ('verified','verified_manual')"
        ).fetchall()
    output, attachment_count = [], 0
    for row in db_rows:
        payload = stored_payload(root, task, row, private_key)
        missing, conflicts, attachments = validate_payload(task, row, payload)
        if missing or conflicts:
            raise ValueError("EXPORT_VALIDATION_FAILED: 已通过记录的业务校验失败，停止导出")
        values = payload["values"]
        record = {
            field_id: normalize_value(field_defs[field_id]["type"], values.get(field_id, ""))
            for field_id in field_defs
            if field_defs[field_id]["type"] not in ATTACHMENT_TYPES
        }
        by_field: dict[str, list[dict[str, Any]]] = {}
        for item in attachments:
            by_field.setdefault(item["field_id"], []).append(item)
        person_dir = f"{safe_filename_component(record.get('name'), 'reply')}-{row['invite_id'][-16:]}"
        for field_id, field in field_defs.items():
            if field["type"] not in ATTACHMENT_TYPES:
                continue
            paths = []
            for index, item in enumerate(by_field.get(field_id, []), 1):
                suffix = {"image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp", "application/pdf": ".pdf"}[item["type"]]
                filename = f"{safe_filename_component(field['label'], field_id)}-{field_id}-{index}{suffix}"
                if attachment_stage is None:
                    raise RuntimeError("ATTACHMENT_EXPORT_UNAVAILABLE: 附件暂存目录未创建")
                secure_io.atomic_write(attachment_stage / person_dir / filename, item["data"], overwrite=False)
                paths.append(f"{attachment_dir_name}/{person_dir}/{filename}")
                attachment_count += 1
            record[field_id] = "; ".join(paths)
        output.append(record)
    output.sort(key=lambda row: (row.get("name", ""), row.get("employee_id", "")))
    return output, attachment_count


def collect_open(task_dir: str | Path, submissions_dir: str | Path, out: str | Path) -> dict[str, Any]:
    root, task = load_task(task_dir)
    if task_mode(task) != "open":
        raise ValueError("collect 仅用于无需名单的开放请求任务；旧任务继续使用 ingest/review/export-clear")
    ingested = ingest_task(root, submissions_dir)
    reviewed = review_task(root, None)
    field_ids = [field["id"] for field in task["fields"] if field["id"] not in {"employee_id", "name"}]
    out_path = secure_io.checked_path(out).with_suffix(".xlsx")
    attachment_fields = [field for field in task["fields"] if field["type"] in ATTACHMENT_TYPES]
    attachment_target = out_path.with_name(out_path.stem + "-attachments")
    if attachment_fields and attachment_target.exists():
        raise FileExistsError("OUTPUT_EXISTS: 附件导出目录已存在，请选择新文件名")
    attachment_stage = None
    written: dict[str, str] = {}
    try:
        if attachment_fields:
            out_path.parent.mkdir(parents=True, exist_ok=True)
            attachment_stage = Path(tempfile.mkdtemp(prefix=".yintian-attachments-", dir=out_path.parent))
        with task_lock(root):
            private_key = unlock_private_key(root, None)
            rows, attachment_count = collect_open_rows(root, task, private_key, attachment_stage, attachment_target.name)
            excluded = len(report_rows(root)) - len(rows)
            written = write_clear_export(root, task, rows, field_ids, {}, out_path, ["xlsx"], task["purpose"], task["contact"])
            if attachment_stage is not None:
                if attachment_count:
                    attachment_stage.rename(attachment_target)
                    written["attachments_dir"] = str(attachment_target)
                    attachment_stage = None
                else:
                    shutil.rmtree(attachment_stage)
                    attachment_stage = None
            with closing(connect_db(root)) as db, db:
                audit(db, "collect_open", reason="agent_requested")
    except BaseException:
        if attachment_stage is not None:
            shutil.rmtree(attachment_stage, ignore_errors=True)
        if written.get("xlsx"):
            Path(written["xlsx"]).unlink(missing_ok=True)
        if written.get("attachments_dir"):
            shutil.rmtree(written["attachments_dir"], ignore_errors=True)
        raise
    return {"task_id": task["task_id"], "rows": len(rows), "excluded": excluded, "attachments": attachment_count, "ingest": ingested, "review": reviewed, **written}


def cmd_collect_open(args) -> dict[str, Any]:
    return collect_open(args.task_dir, args.submissions_dir, args.out)


def write_clear_export(root: Path, task: dict[str, Any], rows: list[dict[str, str]], field_ids: list[str], masks: dict[str, str],
                       out: str | Path, formats: list[str], purpose: str, recipient: str) -> dict[str, str]:
    """把已解密的明文行落盘为 XLSX/JSON（0600，字符串单元格防公式注入）；不产出任何加密出口信封。"""
    unknown = set(formats) - {"xlsx", "json"}
    if unknown:
        raise ValueError(f"不支持的导出格式: {sorted(unknown)}")
    out_path = secure_io.checked_path(out)
    if out_path == root or root in out_path.parents:
        raise ValueError("导出路径不能位于任务目录内")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    for fmt in formats:
        target = secure_io.checked_path(out_path.with_suffix('.' + fmt))
        if target.exists():
            raise FileExistsError("OUTPUT_EXISTS: 导出文件已存在，请选择新文件名")
    identity_columns = ["employee_id", "name"] if task_mode(task) != "open" else [key for key in ("employee_id", "name") if any(field["id"] == key for field in task["fields"])]
    columns = identity_columns + field_ids
    written = {}
    buffers = {}
    if "json" in formats:
        path = out_path.with_suffix(".json")
        buffers[path] = canonical({"task_id": task["task_id"], "exported_at": now_iso(), "purpose": purpose, "recipient": recipient, "fields": columns, "masks": masks, "rows": rows})
        written["json"] = str(path)
    if "xlsx" in formats:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        path = out_path.with_suffix(".xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "明文导出"
        labels = {field["id"]: field["label"] for field in task["fields"]}
        sheet.append([labels.get(key, key) if task_mode(task) == "open" else key for key in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.data_type = "s"
        for row_index, record in enumerate(rows, start=2):
            for column_index, key in enumerate(columns, start=1):
                cell = sheet.cell(row_index, column_index, str(record.get(key, "")))
                cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffers[path] = buffer.getvalue()
        written["xlsx"] = str(path)
    created = []
    try:
        for path, data in buffers.items():
            secure_io.atomic_write(path, data, overwrite=False)
            created.append(path)
    except BaseException:
        for path in created:
            path.unlink(missing_ok=True)
        raise
    return written


def export_snapshot(root):
    with closing(connect_db(root)) as db:
        rows = [tuple(row) for row in db.execute("SELECT id,sha256,status FROM submissions ORDER BY id")]
    return sha256_bytes(canonical([load_task(root)[1], rows]))


def cmd_export_clear(args) -> dict[str, Any]:
    require_tty("export-clear")
    root, task = load_task(args.task_dir)
    require_current_format(task, "明文导出")
    if task_expired(task):
        raise RuntimeError("任务已超过保存期限，禁止明文导出；请执行清理")
    purpose, recipient = (getattr(args, "purpose", "") or "").strip(), (getattr(args, "recipient", "") or "").strip()
    if not purpose or not recipient:
        raise ValueError("EXPORT_PURPOSE_REQUIRED: 请填写 --purpose 和 --recipient")
    field_ids = parse_export_fields(task, getattr(args, "fields", None))
    masks = parse_mask_specs(getattr(args, "mask", None))
    unknown_mask = set(masks) - ({"employee_id", "name"} | set(field_ids))
    if unknown_mask:
        raise ValueError(f"脱敏规则指向未导出的字段: {sorted(unknown_mask)}")
    formats = sorted({part.strip().lower() for value in (args.formats or ["xlsx"]) for part in value.split(",") if part.strip()})
    password = getpass.getpass("任务密码: ")
    try:
        private_key = unlock_private_key(root, password)
    except Exception as exc:
        raise RuntimeError("任务密码错误或私钥损坏") from exc
    with task_lock(root):
        rows = collect_clear_rows(root, task, private_key, field_ids, masks)
        snapshot = export_snapshot(root)
        excluded = len(report_rows(root)) - len(rows)
    print("\n=== 明文导出确认 ===")
    print(f"任务: {task['task_id']}（{task['title']}）")
    print(f"导出字段: {', '.join(['employee_id', 'name'] + field_ids)}")
    print(f"脱敏规则: {', '.join(f'{key}={rule}' for key, rule in sorted(masks.items())) or '无（白名单字段全部明文）'}")
    print(f"数据行数: {len(rows)}")
    print(f"未纳入人数: {excluded}（最新版本未通过或尚未提交；详见进度报告）")
    print(f"用途: {purpose or '（未填写）'}")
    print(f"接收方: {recipient or '（未填写）'}")
    typed = input(f"输入任务 ID {task['task_id']} 以确认导出明文: ").strip()
    if typed != task["task_id"]:
        raise RuntimeError("任务 ID 不匹配，已取消导出")
    with task_lock(root):
        if task_expired(load_task(root)[1]) or snapshot != export_snapshot(root):
            raise RuntimeError("STATE_CHANGED: 任务已变化或到期，请重新确认")
        # Re-read authenticated bytes after the human confirmation gap.
        rows = collect_clear_rows(root, task, private_key, field_ids, masks)
        written = write_clear_export(root, task, rows, field_ids, masks, args.out, formats, purpose, recipient)
        with closing(connect_db(root)) as db, db:
            audit(db, "export_clear", reason="approved_latest_only")
    print("警告：明文已落盘，用后请删除，系统无法管控后续传播。", file=sys.stderr)
    return {"task_id": task["task_id"], "rows": len(rows), "excluded": excluded, **written}


def build_task_package(task_dir, *, include_open_secret: bool = False):
    root, task = load_task(task_dir)
    if task_expired(task):
        raise RuntimeError("TASK_EXPIRED: 任务已超过保存期限")
    if task_mode(task) == "open" and not include_open_secret:
        raise RuntimeError("OPEN_PACKAGE_ENCRYPTION_REQUIRED: 开放任务只能通过 export_task 生成加密交接包")
    with task_lock(root):
        for path in root.rglob("*"):
            secure_io.checked_path(path, root)
        candidates = [root / name for name in ("task.json", "public.pem", "private.pem.enc", "roster.csv", "invite-index.csv", "state.sqlite3")]
        if any(not path.is_file() for path in candidates):
            raise ValueError("任务目录缺少必要文件")
        candidates += list((root / "invites").glob("*.yintian-form"))
        candidates += list((root / "credentials").glob("*.yintian-credential"))
        for public_artifact in (root / "REQUEST.yintian-request", root / "FORM.yintian-form"):
            if public_artifact.exists():
                candidates.append(public_artifact)
        candidates += list((root / "submissions").glob("*/*.yintian"))
        candidates += [p for p in (root / "reports").glob("progress.*") if p.suffix in {".json", ".xlsx"}]
        if len(candidates) + (task_mode(task) == "open") > MAX_PACKAGE_FILES:
            raise ValueError("任务内容超过交接包安全上限")
        data, total = {}, 0
        for path in sorted(candidates):
            if path.name == "state.sqlite3":
                with closing(connect_db(root)) as source, closing(sqlite3.connect(":memory:")) as snapshot:
                    source.backup(snapshot)
                    raw = snapshot.serialize()
            else:
                raw = secure_io.read_bytes(path, MAX_PACKAGE_MEMBER_BYTES)
            total += len(raw)
            if len(raw) > MAX_PACKAGE_MEMBER_BYTES or total > MAX_PACKAGE_BYTES:
                raise ValueError("任务内容超过交接包安全上限")
            data[path.relative_to(root).as_posix()] = raw
        if task_mode(task) == "open":
            raw = load_local_task_secret(root, task["task_id"]).encode("ascii")
            total += len(raw)
            if total > MAX_PACKAGE_BYTES:
                raise ValueError("任务内容超过交接包安全上限")
            data["local-open-key"] = raw
        metadata = {"package_version": TASK_PACKAGE_VERSION, "task_id": task["task_id"], "created_at": now_iso(),
                    "manifest": {name: sha256_bytes(raw) for name, raw in data.items()}}
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("package.json", canonical(metadata))
            for name, raw in data.items():
                archive.writestr(task["task_id"] + "/" + name, raw)
        return buffer.getvalue(), {"task_id": task["task_id"], "files": len(data)}


def task_package_aad(task_id: str) -> bytes:
    return f"{ENCRYPTED_TASK_PACKAGE_VERSION}:{task_id}".encode("utf-8")


def export_task(task_dir: str | Path, out: str | Path, handoff_password: str | None = None) -> dict[str, Any]:
    root, task = load_task(task_dir)
    zip_bytes, info = build_task_package(root, include_open_secret=True)
    password = handoff_password if handoff_password is not None else generate_password()
    envelope = aes_gcm_seal(ENCRYPTED_TASK_PACKAGE_VERSION, zip_bytes, password, aad=task_package_aad(task["task_id"]))
    envelope["task_id"] = task["task_id"]
    out_path = secure_io.checked_path(out)
    if out_path == root or root in out_path.parents:
        raise ValueError("导出路径不能位于任务目录内")
    secure_io.atomic_write(out_path, canonical(envelope), overwrite=False)
    return {**info, "out": str(out_path), "handoff_password": password}


def cmd_export(args) -> dict[str, Any]:
    require_tty("export-task")
    result = export_task(args.task_dir, args.out)
    password = result.pop("handoff_password")
    print_once_secret("交接密码（仅显示一次，丢失不可恢复）：", password, "请通过另一独立安全渠道告知接收方；不要与交接包同渠道发送，也不要写入任务目录或聊天提示词。")
    return result


def read_package_member(archive: zipfile.ZipFile, name: str, remaining: int) -> bytes:
    """流式读取任务包成员并按实际字节计数，超过单项或累计上限立即失败。"""
    data = bytearray()
    with archive.open(name) as stream:
        while True:
            chunk = stream.read(1024 * 1024)
            if not chunk:
                return bytes(data)
            data += chunk
            if len(data) > MAX_PACKAGE_MEMBER_BYTES or len(data) > remaining:
                raise ValueError("任务包解压后超过安全上限")


def open_task_package(package_path: Path, handoff_password: str | None = None) -> tuple[zipfile.ZipFile, bool]:
    """打开交接包，返回 (archive, encrypted)。

    yintian-task/3 加密信封先按 stat 尺寸拒绝超限文件，再解密认证；旧版明文 ZIP 打印警告后直接打开。
    嗅探窗口内全为空白且尺寸未超限的内容按 v3 信封处理，避免大段前导空白被误判为明文包。
    """
    size = package_path.stat().st_size
    with package_path.open("rb") as stream:
        prefix = stream.read(V3_SNIFF_BYTES).lstrip()
    if prefix.startswith(b"{") or not prefix:
        if size > MAX_V3_ENVELOPE_BYTES:
            raise ValueError("交接包信封超过 200MB 安全上限，拒绝读取")
        try:
            envelope = json.loads(package_path.read_bytes())
        except (json.JSONDecodeError, UnicodeDecodeError) as exc:
            raise ValueError("交接包信封无效") from exc
        if not isinstance(envelope, dict) or envelope.get("format") != ENCRYPTED_TASK_PACKAGE_VERSION:
            raise ValueError("交接包格式不受支持")
        task_id = envelope.get("task_id")
        if not isinstance(task_id, str) or not TASK_ID_RE.fullmatch(task_id):
            raise ValueError("交接包 task_id 无效")
        password = handoff_password if handoff_password is not None else getpass.getpass("交接密码: ")
        try:
            zip_bytes = aes_gcm_open(envelope, password, task_package_aad(task_id))
        except Exception as exc:
            raise ValueError("交接密码错误或交接包已被篡改，拒绝导入") from exc
        return zipfile.ZipFile(io.BytesIO(zip_bytes)), True
    print("警告：导入的是未加密的明文交接包（yintian-task/1 或 yintian-task/2），不能认证来源；仅在信任渠道下导入。", file=sys.stderr)
    return zipfile.ZipFile(package_path), False


def import_task(package: str | Path, out_parent: str | Path, handoff_password: str | None = None, confirm_plaintext=None) -> dict[str, Any]:
    package_path, parent = secure_io.checked_path(package), secure_io.checked_path(out_parent)
    if package_path.stat().st_size > MAX_PACKAGE_BYTES:
        raise ValueError("任务包超过安全上限")
    archive, encrypted = open_task_package(package_path, handoff_password)
    with archive:
        package_info = archive.getinfo("package.json")
        if package_info.file_size > 1024 * 1024:
            raise ValueError("任务包清单过大")
        metadata = json.loads(read_package_member(archive, "package.json", 1024 * 1024))
        task_id = metadata.get("task_id")
        package_version = metadata.get("package_version")
        if package_version not in {LEGACY_TASK_PACKAGE_VERSION, TASK_PACKAGE_VERSION} or not isinstance(task_id, str) or not TASK_ID_RE.fullmatch(task_id):
            raise ValueError("任务包格式无效")
        if not encrypted and confirm_plaintext is not None and not confirm_plaintext(task_id):
            raise RuntimeError("未加密的明文交接包未经人工确认，已取消导入")
        raw_manifest = metadata.get("manifest")
        if not isinstance(raw_manifest, dict) or len(raw_manifest) > MAX_PACKAGE_FILES:
            raise ValueError("任务包清单无效或文件过多")
        required = {"task.json", "public.pem", "private.pem.enc", "roster.csv", "invite-index.csv", "state.sqlite3"}
        manifest: dict[str, str] = {}
        for raw_rel, expected_hash in raw_manifest.items():
            if not isinstance(raw_rel, str) or not raw_rel or not isinstance(expected_hash, str) or ":" in raw_rel:
                raise ValueError("任务包包含不安全路径")
            if package_version == TASK_PACKAGE_VERSION and "\\" in raw_rel:
                raise ValueError("任务包包含不安全路径")
            rel = raw_rel.replace("\\", "/")
            rel_path = PurePosixPath(rel)
            if rel_path.is_absolute() or ".." in rel_path.parts:
                raise ValueError("任务包包含不安全路径")
            parts = rel_path.parts
            allowed = (
                rel in required
                or rel in {"REQUEST.yintian-request", "FORM.yintian-form", "local-open-key"}
                or (len(parts) == 2 and parts[0] == "invites" and parts[1].endswith(".yintian-form") and INVITE_ID_RE.fullmatch(parts[1][:-13]))
                or (len(parts) == 2 and parts[0] == "credentials" and parts[1].endswith(".yintian-credential") and valid_invite_identifier(parts[1][:-19]))
                or (len(parts) == 3 and parts[0] == "submissions" and valid_invite_identifier(parts[1]) and re.fullmatch(r"v\d{4}_[0-9a-f]{12}\.yintian", parts[2]))
                or rel in {"reports/progress.json", "reports/progress.xlsx"}
            )
            if not allowed or rel in manifest:
                raise ValueError("任务包包含不允许或重复的文件")
            manifest[rel] = expected_hash
        if not required.issubset(manifest):
            raise ValueError("任务包缺少必要文件")
        target = parent / task_id
        if target.exists():
            raise FileExistsError(f"目标任务已存在: {target}")
        expected_prefix = task_id + "/"
        archive_names: dict[str, str] = {}
        for info in archive.infolist():
            if package_version == TASK_PACKAGE_VERSION and "\\" in info.filename:
                raise ValueError("任务包包含不安全路径")
            normalized = info.filename.replace("\\", "/")
            if normalized in archive_names:
                raise ValueError("任务包包含重复条目")
            archive_names[normalized] = info.filename
        names = set(archive_names)
        expected_names = {expected_prefix + rel for rel in manifest}
        if names != expected_names | {"package.json"}:
            raise ValueError("任务包文件不完整或包含额外条目")
        infos = [archive.getinfo(archive_names[name]) for name in expected_names]
        if any(info.file_size > MAX_PACKAGE_MEMBER_BYTES for info in infos) or sum(info.file_size for info in infos) > MAX_PACKAGE_BYTES:
            raise ValueError("任务包解压后超过安全上限")
        target.mkdir(parents=True, mode=0o700)
        imported_secret_path = None
        open_secret = None
        try:
            total_size = 0
            for rel, expected_hash in manifest.items():
                destination = (target.joinpath(*PurePosixPath(rel).parts)).resolve()
                if target != destination and target not in destination.parents:
                    raise ValueError("任务包包含越界路径")
                data = read_package_member(archive, archive_names[expected_prefix + rel], MAX_PACKAGE_BYTES - total_size)
                total_size += len(data)
                if sha256_bytes(data) != expected_hash:
                    raise ValueError(f"任务包哈希不匹配: {rel}")
                if rel == "local-open-key":
                    open_secret = data
                else:
                    atomic_write(destination, data)
            _, imported_task = load_task(target)
            if imported_task["task_id"] != task_id:
                raise ValueError("任务包内外 task_id 不一致")
            if task_mode(imported_task) == "open":
                if not encrypted or open_secret is None:
                    raise ValueError("OPEN_KEY_PACKAGE_INVALID: 开放任务必须通过加密交接包携带本地密钥")
                try:
                    secret = open_secret.decode("ascii")
                except UnicodeDecodeError as exc:
                    raise ValueError("OPEN_KEY_PACKAGE_INVALID: 本地密钥格式无效") from exc
                if not re.fullmatch(r"[A-Za-z0-9_.~-]{28,128}", secret):
                    raise ValueError("OPEN_KEY_PACKAGE_INVALID: 本地密钥格式无效")
                imported_secret_path = save_local_task_secret(target, task_id, secret)
            elif open_secret is not None:
                raise ValueError("任务包包含意外的开放任务密钥")
            from cryptography.hazmat.primitives import serialization

            public_key = serialization.load_pem_public_key((target / "public.pem").read_bytes())
            public_der = public_key.public_bytes(serialization.Encoding.DER, serialization.PublicFormat.SubjectPublicKeyInfo)
            if hashlib.sha256(public_der).hexdigest()[:24] != imported_task["key_id"] or not valid_private_key_blob((target / "private.pem.enc").read_bytes()):
                raise ValueError("任务包密钥材料无效")
            if task_mode(imported_task) == "open":
                private_der = unlock_private_key(target, None).public_key().public_bytes(serialization.Encoding.DER, serialization.PublicFormat.SubjectPublicKeyInfo)
                if not secrets.compare_digest(private_der, public_der):
                    raise ValueError("任务包公私钥不匹配")
            roster = read_roster(target / "roster.csv", allow_empty=task_mode(imported_task) == "open")
            with (target / "invite-index.csv").open("r", encoding="utf-8-sig", newline="") as stream:
                index = list(csv.DictReader(stream))
            if len(index) != len(roster) or (task_mode(imported_task) != "open" and not index) or any(not valid_invite_identifier(row.get("invite_id", "")) for row in index):
                raise ValueError("任务包名单或邀请索引无效")
            if task_mode(imported_task) != "open" and status_task(target)["total"] != len(roster):
                raise ValueError("任务包状态数据库与名单不一致")
            if imported_task["format_version"] in {FORMAT_VERSION, GROUP_FORMAT_VERSION}:
                with closing(connect_db(target)) as db, db:
                    columns = {row[1] for row in db.execute("PRAGMA table_info(invites)")}
                if "token_hash" not in columns:
                    raise ValueError("v2 任务缺少邀请认证数据")
        except Exception:
            shutil.rmtree(target, ignore_errors=True)
            if imported_secret_path is not None:
                imported_secret_path.unlink(missing_ok=True)
            raise
    if os.name != "nt":
        for directory in [target, *(path for path in target.rglob("*") if path.is_dir())]:
            os.chmod(directory, 0o700)
    return {"task_id": task_id, "task_dir": str(target)}


def confirm_plaintext_import(task_id: str) -> bool:
    typed = input(f"该交接包未加密且无法认证来源；输入任务 ID {task_id} 确认导入: ").strip()
    return typed == task_id


def cmd_import(args) -> dict[str, Any]:
    require_tty("import-task")
    return import_task(args.package, args.out, confirm_plaintext=confirm_plaintext_import)


def cmd_purge(args) -> dict[str, Any]:
    require_tty("purge")
    root, task = load_task(args.task_dir)
    if not task_expired(task) and not args.allow_early:
        raise RuntimeError("任务尚未到保存期限；如确需提前删除，请增加 --allow-early")
    typed = input(f"输入任务 ID {task['task_id']} 以确认删除任务目录: ").strip()
    if typed != task["task_id"]:
        raise RuntimeError("任务 ID 不匹配，已取消")
    parent = root.parent
    local_secret = local_task_secret_path(root, task["task_id"]) if task_mode(task) == "open" else None
    with task_lock(root):
        status = status_task(root)
        summary = {"task_id": task["task_id"], "purged_at": now_iso(), "aggregate_counts": status["counts"], "total": status["total"]}
        if os.name == 'nt':
            # Windows cannot delete an open CRT lock file. Remove task data under
            # the lock; after release only the inert lock/empty directory remain.
            for path in root.rglob('*'):
                secure_io.checked_path(path, root)
            for path in root.iterdir():
                if path.name != '.write.lock':
                    shutil.rmtree(path) if path.is_dir() else path.unlink()
        else:
            shutil.rmtree(root)
    if os.name == 'nt':
        if {path.name for path in root.iterdir()} - {'.write.lock'}:
            raise RuntimeError('STATE_CHANGED: 删除结束时目录出现新文件，停止清理')
        (root / '.write.lock').unlink(missing_ok=True)
        root.rmdir()
    summary_path = parent / f"{task['task_id']}.purged.json"
    dump_json(summary_path, summary)
    if local_secret is not None:
        local_secret.unlink(missing_ok=True)
    return {"summary": str(summary_path)}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="SafeFill：端到端加密的私密信息收集管理")
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("init-config", help="生成基础身份信息收集配置")
    p.add_argument("--out", required=True); p.add_argument("--force", action="store_true"); p.add_argument('--mode', choices=['directed', 'group', 'open'], default='open'); p.set_defaults(func=cmd_init_config)
    p = sub.add_parser("create-request", aliases=["create-open"], help="生成供 safefill-fill Agent 读取的机器请求包；不生成 HTML 表单")
    p.add_argument("--config", required=True); p.add_argument("--out", required=True); p.set_defaults(func=cmd_create_open)
    p = sub.add_parser("create", help="创建任务并批量生成离线邀请")
    p.add_argument("--roster", required=True); p.add_argument("--config", required=True); p.add_argument("--out", required=True)
    p.add_argument("--mode", choices=["directed", "group"], default="directed", help="directed 个人邀请；group 公共模板加私下发放的个人凭据")
    p.set_defaults(func=cmd_create)
    p = sub.add_parser("ingest", help="接收 .yintian 密文提交")
    p.add_argument("task_dir"); p.add_argument("submissions_dir"); p.set_defaults(func=cmd_ingest)
    p = sub.add_parser("collect", help="开放请求一键收件、校验、解密并导出 Excel")
    p.add_argument("task_dir"); p.add_argument("submissions_dir"); p.add_argument("--out", required=True); p.set_defaults(func=cmd_collect_open)
    p = sub.add_parser("review", help="本地解密、校验和 OpenVINO OCR 复核")
    p.add_argument("task_dir"); p.add_argument("--retry-needs-review", action="store_true"); p.add_argument("--invite"); p.set_defaults(func=cmd_review)
    p = sub.add_parser("decide", help="本人终端逐项核对 OCR 证据或退回重填")
    p.add_argument("task_dir"); p.add_argument("invite_id"); p.add_argument("--version", required=True, type=int)
    p.add_argument("--action", choices=["confirm", "return"], required=True); p.add_argument("--operator", required=True); p.set_defaults(func=cmd_decide)
    p = sub.add_parser("migrate", help="显式备份并升级数据库；旧无认证群发任务保持只读")
    p.add_argument("task_dir"); p.set_defaults(func=cmd_migrate)
    p = sub.add_parser("doctor", help="检查 Python、核心依赖、OCR 与 Tk，不读取私密材料")
    p.add_argument("--vault", help="仅检查目录权限"); p.set_defaults(func=cmd_doctor)
    p = sub.add_parser("status", help="查看数据最小化统计")
    p.add_argument("task_dir"); p.set_defaults(func=cmd_status)
    p = sub.add_parser("report", help="生成保留姓名/员工号的数据最小化 XLSX/JSON 报告")
    p.add_argument("task_dir"); p.add_argument("--formats", nargs="+", default=["xlsx", "json"]); p.set_defaults(func=cmd_report)
    p = sub.add_parser("reveal", help="人工终端按需查看一份明文")
    p.add_argument("task_dir"); p.add_argument("invite_id"); p.set_defaults(func=cmd_reveal)
    p = sub.add_parser("export-task", help="导出加密的敏感任务交接包（yintian-task/3，一次性交接密码仅显示一次）")
    p.add_argument("task_dir"); p.add_argument("--out", required=True); p.set_defaults(func=cmd_export)
    p = sub.add_parser("import-task", help="导入任务交接包（加密包需交互输入交接密码；v1/v2 明文旧包需输入任务 ID 二次确认）")
    p.add_argument("package"); p.add_argument("--out", required=True); p.set_defaults(func=cmd_import)
    p = sub.add_parser("purge", help="删除任务目录；不保证擦除备份或磁盘残留")
    p.add_argument("task_dir"); p.add_argument("--allow-early", action="store_true"); p.set_defaults(func=cmd_purge)
    p = sub.add_parser(
        "export-clear",
        help="【仅限人类授权操作员在本人终端运行，禁止 Agent 代跑】把白名单字段的当前有效提交解密导出为明文 XLSX/JSON（需任务密码并输入任务 ID 二次确认）",
        description="仅限人类授权操作员在本人交互终端运行，禁止 Agent 代跑或旁观。明文导出是最后一道出口：只导出 --fields 白名单字段（name/employee_id 默认附带），落盘 0600，系统无法管控后续传播。",
    )
    p.add_argument("task_dir")
    p.add_argument("--out", required=True, help="导出文件路径（多格式时按后缀派生同名 .xlsx/.json）")
    p.add_argument("--fields", required=True, help="逗号分隔的导出字段 id 白名单，仅限非附件字段；name/employee_id 默认附带")
    p.add_argument("--mask", action="append", default=[], metavar="字段id=规则", help="脱敏规则 last4（保留后四位）或 mid4（保留前3后4），可重复；未声明字段明文导出")
    p.add_argument("--purpose", default="", help="本次明文导出的用途，回显确认并写入 JSON 元数据")
    p.add_argument("--recipient", default="", help="明文接收方，回显确认并写入 JSON 元数据")
    p.add_argument("--formats", nargs="+", default=["xlsx"], help="导出格式：xlsx、json（默认 xlsx）")
    p.set_defaults(func=cmd_export_clear)
    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()
    try:
        result = args.func(args)
        if result is not None:
            print(json.dumps(result, ensure_ascii=False, indent=2))
    except Exception as exc:
        print(f"错误: {terminal_text(exc)}", file=sys.stderr)
        raise SystemExit(1)


if __name__ == "__main__":
    main()
