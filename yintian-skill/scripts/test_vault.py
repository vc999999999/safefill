"""Synthetic employee profile and real authenticated group roundtrip tests."""
import base64
import contextlib
import io
import json
import os
from pathlib import Path
import tempfile
import types
import unittest
from unittest.mock import patch

import collection
import fill
import vault
import collection
import test_collection as fixtures
from openpyxl import load_workbook


class VaultTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.password = 'Synthetic-vault-password-2026'
        self.profile = {'version': 1, 'types': {'phone': 'phone_cn', 'secret_note': 'text'},
                        'values': {'phone': '13800138000', 'secret_note': 'SYNTHETIC_PRIVATE_DO_NOT_EXPORT'}, 'attachments': {}}

    def test_vault_authenticated_encryption_and_no_plaintext_temps(self):
        path = self.root / 'personal.yintian-vault'
        vault.save_vault(path, self.password, self.profile, create=True)
        self.assertEqual(vault.load_vault(path, self.password), self.profile)
        self.assertEqual(list(self.root.iterdir()), [path])
        self.assertNotIn(self.profile['values']['secret_note'].encode(), path.read_bytes())
        if os.name != 'nt':
            self.assertEqual(path.stat().st_mode & 0o777, 0o600)
        with self.assertRaisesRegex(ValueError, 'VAULT_UNLOCK_FAILED'):
            vault.load_vault(path, 'wrong')
        data = json.loads(path.read_bytes())
        raw = bytearray(base64.b64decode(data['ciphertext'])); raw[-1] ^= 1
        data['ciphertext'] = base64.b64encode(raw).decode()
        path.write_text(json.dumps(data))
        with self.assertRaisesRegex(ValueError, 'VAULT_UNLOCK_FAILED'):
            vault.load_vault(path, self.password)

    def test_only_requested_fields_and_explicit_ambiguity(self):
        form = {'fields': [{'id': 'mobile', 'type': 'phone_cn'}]}
        values, attachments, missing, matches = vault.select_fields(form, self.profile)
        self.assertEqual(values, {'mobile': self.profile['values']['phone']})
        self.assertNotIn('secret_note', matches)
        self.profile['types']['other_phone'] = 'phone_cn'
        self.profile['values']['other_phone'] = '13800138001'
        self.assertEqual(vault.select_fields(form, self.profile)[2], ['mobile'])
        self.assertEqual(vault.select_fields(form, self.profile, {'mobile': 'phone'})[0], values)
        with self.assertRaisesRegex(ValueError, 'MAPPING_INVALID'):
            vault.select_fields(form, self.profile, {'mobile': 'secret_note'})

    def test_no_consent_cannot_produce_ciphertext(self):
        out = self.root / 'reply.yintian'
        with self.assertRaisesRegex(fill.FillError, 'CONSENT_REQUIRED'):
            fill.seal_data({}, {}, {}, out)
        self.assertFalse(out.exists())

    def test_agent_non_tty_cannot_unlock_or_seal(self):
        for command in [['fill', 'fake', '--vault', 'fake', '--out', 'fake'],
                        ['seal', 'fake', '--values', 'fake', '--out', 'fake'],
                        ['vault-init', 'fake', '--vault', 'fake']]:
            with self.subTest(command=command[0]), fixtures.non_tty(), patch.object(vault, 'load_vault', side_effect=AssertionError('private file read')), contextlib.redirect_stderr(io.StringIO()):
                self.assertEqual(fill.main(command), 1)

    def test_profile_to_group_submission_to_excel(self):
        ns = fixtures.make_group_scenario(self.root)
        task = collection.load_json(ns.task_dir / 'task.json')
        profile = {'version': 1, 'types': {f['id']: f['type'] for f in task['fields']},
                   'values': dict(ns.values), 'attachments': {}}
        profile['types']['secret_note'] = 'text'
        profile['values']['secret_note'] = 'SYNTHETIC_PRIVATE_DO_NOT_EXPORT'
        path = self.root / 'personal.yintian-vault'
        vault.save_vault(path, self.password, profile, create=True)
        prompts = iter([task['key_id'], task['task_id']])
        out = ns.incoming / 'reply.yintian'
        with fixtures.fake_tty(input_func=lambda _: next(prompts), getpass_func=lambda _: self.password):
            self.assertEqual(fill.main(['fill', str(ns.task_dir / 'FORM.yintian-form'), '--vault', str(path),
                                        '--credential', str(ns.task_dir / 'credentials/GRP-E001.yintian-credential'), '--out', str(out)]), 0)
        self.assertNotIn(ns.values['phone'].encode(), out.read_bytes())
        self.assertEqual(collection.ingest_task(ns.task_dir, ns.incoming)['accepted'], 1)
        self.assertEqual(collection.review_task(ns.task_dir, fixtures.PASSWORD)['verified'], 1)
        with collection.connect_db(ns.task_dir) as db:
            row = db.execute('SELECT * FROM submissions').fetchone()
        payload = collection.stored_payload(ns.task_dir, task, row, collection.unlock_private_key(ns.task_dir, fixtures.PASSWORD))
        self.assertNotIn('secret_note', payload['values'])
        self.assertEqual(payload['values'], ns.values)
        result_path = self.root / 'result.xlsx'
        with fixtures.fake_tty(input_func=lambda _: ns.task_id, getpass_func=lambda _: fixtures.PASSWORD), contextlib.redirect_stderr(io.StringIO()):
            result = collection.cmd_export_clear(fixtures.export_clear_args(ns, result_path, mask=[]))
        book = load_workbook(result['xlsx'])
        self.assertEqual(book.active.cell(2, 3).value, ns.values['phone'])
        book.close()
        self.assertFalse(list(self.root.glob('*values.json')))
        for file in (ns.task_dir / 'reports').glob('*'):
            self.assertNotIn(b'SYNTHETIC_PRIVATE_DO_NOT_EXPORT', file.read_bytes())

    def test_wrong_credential_and_cancel_write_nothing(self):
        ns = fixtures.make_group_scenario(self.root)
        form = fill.load_form(ns.task_dir / 'FORM.yintian-form')
        with self.assertRaisesRegex(fill.FillError, 'CREDENTIAL_REQUIRED'):
            fill.bind_credential(form)
        bound = fill.bind_credential(form, ns.task_dir / 'credentials/GRP-E002.yintian-credential')
        with self.assertRaisesRegex(fill.FillError, 'NAME_MISMATCH|IDENTITY_MISMATCH'):
            fill.seal_data(bound, ns.values, {}, self.root / 'bad.yintian', confirmed=True)
        bound = fill.bind_credential(form, ns.task_dir / 'credentials/GRP-E001.yintian-credential')
        with fixtures.fake_tty(input_func=lambda _: 'cancel'):
            with self.assertRaisesRegex(fill.FillError, 'KEY_UNCONFIRMED'):
                fill.confirm_submission(bound, ns.values, {})
        self.assertFalse((self.root / 'bad.yintian').exists())

    def test_initial_vault_cli_and_terminal_control_sanitization(self):
        ns = fixtures.make_group_scenario(self.root)
        path = self.root / 'personal.yintian-vault'
        # Use the actual field order rather than the order of a separate dictionary.
        task = collection.load_json(ns.task_dir / 'task.json')
        secrets = iter([self.password, self.password, *[ns.values[field['id']] for field in task['fields']]])
        with fixtures.fake_tty(getpass_func=lambda _: next(secrets), input_func=lambda _: 'SAVE'):
            self.assertEqual(fill.main(['vault-init', str(ns.task_dir / 'FORM.yintian-form'), '--vault', str(path)]), 0)
        self.assertEqual(vault.load_vault(path, self.password)['values'], ns.values)
        form = fill.load_form(ns.task_dir / 'FORM.yintian-form')
        info = fill.inspect_info(form)
        info['title'] = '\x1b]52;c;bad\x07\u202eTitle'
        info['fields'][0]['label'] = '\x1b[31m工号'
        captured = io.StringIO()
        with contextlib.redirect_stdout(captured):
            fill.print_inspect(info)
        self.assertNotIn('\x1b', captured.getvalue())
        self.assertNotIn('\u202e', captured.getvalue())


if __name__ == '__main__':
    unittest.main(verbosity=2)
