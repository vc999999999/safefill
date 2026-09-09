import json
import os
import stat
import sys
import types
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "safefill-fill" / "scripts"
sys.path.insert(0, str(SCRIPTS))

import collection  # noqa: E402
import fill  # noqa: E402
import secure_io  # noqa: E402
import vault  # noqa: E402
import vlm_extract  # noqa: E402


def private_dir(path: Path) -> Path:
    path.mkdir(mode=0o700)
    if os.name != "nt":
        path.chmod(0o700)
    return path


def private_json(path: Path, value) -> Path:
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
    if os.name != "nt":
        path.chmod(0o600)
    return path


@pytest.fixture
def isolated_vault(tmp_path, monkeypatch):
    vault_dir = private_dir(tmp_path / "vault")
    key_dir = private_dir(tmp_path / "keys")
    monkeypatch.setenv("YINTIAN_VAULT_DIR", str(vault_dir))
    monkeypatch.setenv("YINTIAN_VAULT_KEY_DIR", str(key_dir))
    return vault_dir, key_dir


def make_request(tmp_path: Path):
    config = {
        "title": "员工联系方式",
        "purpose": "用于紧急联络",
        "deadline": "2098-12-31T23:59:59Z",
        "retention_until": "2099-01-31T23:59:59Z",
        "contact": "人事部 hr@example.com",
        "correction": "使用本人旧回执更正",
        "template_version": "1.0",
        "fields": [
            {"id": "name", "label": "姓名", "type": "text", "required": True, "sensitive": True},
            {"id": "phone", "label": "手机号", "type": "phone_cn", "required": True, "sensitive": True},
        ],
    }
    config_path = tmp_path / "collection.json"
    collection.dump_json(config_path, config)
    args = collection.build_parser().parse_args(
        ["create-request", "--config", str(config_path), "--out", str(tmp_path / "tasks")])
    created = args.func(args)
    return created, Path(created["request"]), Path(created["task_dir"])


def run(command):
    args = fill.build_parser().parse_args(command)
    return args.func(args)


def stage_profile(tmp_path: Path, entries: dict):
    private = tmp_path / "private"
    if not private.exists():
        private_dir(private)
    answers = private_json(private / "answers.json", {"entries": entries})
    confirmation = private / "change.yintian-confirmation"
    staged = run(["vault-stage", "--answers", str(answers), "--confirmation-out", str(confirmation)])
    assert not answers.exists()
    assert confirmation.is_file()
    if os.name != "nt":
        assert stat.S_IMODE(confirmation.stat().st_mode) == 0o600
    applied = run(["vault-apply", "--confirmation", str(confirmation)])
    assert not confirmation.exists()
    return staged, applied


def test_exact_matching_never_guesses_semantics():
    profile = {"format": vault.FORMAT_V2, "entries": {
        "birth_date": {"type": "date", "label": "出生日期", "value": "1990-01-01", "source": {"kind": "manual"}},
        "emergency_phone": {"type": "phone_cn", "label": "紧急联系人电话", "value": "13900139000", "source": {"kind": "manual"}},
        "registered_address": {"type": "address", "label": "户籍地址", "value": "旧地址", "source": {"kind": "manual"}},
    }}
    form = {"fields": [
        {"id": "hire_date", "label": "入职日期", "type": "date", "required": True},
        {"id": "phone", "label": "本人手机号", "type": "phone_cn", "required": True},
        {"id": "current_address", "label": "现居地址", "type": "address", "required": True},
    ]}
    values, _attachments, missing, matches = vault.select_fields(form, profile)
    assert values == {} and missing == ["hire_date", "phone", "current_address"] and matches == {}
    assert vault.select_fields(form, profile, {
        "hire_date": "birth_date", "phone": "emergency_phone", "current_address": "registered_address",
    })[0] == {"hire_date": "1990-01-01", "phone": "13900139000", "current_address": "旧地址"}


def test_stage_preview_fill_and_excel_roundtrip(tmp_path, isolated_vault):
    created, request, task_dir = make_request(tmp_path)
    staged, applied = stage_profile(tmp_path, {
        "name": {"type": "text", "label": "姓名", "value": "张三"},
        "phone": {"type": "phone_cn", "label": "手机号", "value": "13800138000"},
    })
    assert staged["changes"][0]["new_value"] == "张三"
    assert applied["created"] == ["name", "phone"]
    private = tmp_path / "private"
    confirmation = private / "submit.yintian-confirmation"
    preview = run(["vault-preview", str(request), "--confirmation-out", str(confirmation)])
    assert [field["value"] for field in preview["fields"]] == ["张三", "13800138000"]
    assert all(field["source_entry"] in {"name", "phone"} for field in preview["fields"])
    incoming = private_dir(tmp_path / "incoming")
    result = run(["vault-fill", str(request), "--confirmation", str(confirmation), "--out-dir", str(incoming)])
    reply = Path(result["out"])
    assert reply.name.startswith("张三-") and b"13800138000" not in reply.read_bytes()
    assert not confirmation.exists()
    exported = collection.collect_open(task_dir, incoming, tmp_path / "result.xlsx")
    book = load_workbook(exported["xlsx"], read_only=True)
    rows = list(book.active.values)
    book.close()
    assert rows == [("姓名", "手机号"), ("张三", "13800138000")]
    assert created["request"] == str(request)

    stale_previous = private / "previous-copy.yintian"
    stale_previous.write_bytes(reply.read_bytes())
    previous_confirmation = private / "previous-stale.yintian-confirmation"
    run(["vault-preview", str(request), "--previous", str(stale_previous),
         "--confirmation-out", str(previous_confirmation)])
    stale_previous.write_text(json.dumps(json.loads(stale_previous.read_text(encoding="utf-8")), indent=2), encoding="utf-8")
    with pytest.raises(fill.FillError, match="CONFIRMATION_STALE"):
        run(["vault-fill", str(request), "--previous", str(stale_previous),
             "--confirmation", str(previous_confirmation), "--out-dir", str(incoming)])
    assert not previous_confirmation.exists()

    correction_confirmation = private / "correction.yintian-confirmation"
    correction_dir = private_dir(tmp_path / "correction")
    run(["vault-preview", str(request), "--previous", str(reply),
         "--confirmation-out", str(correction_confirmation)])
    corrected = run(["vault-fill", str(request), "--previous", str(reply),
                     "--confirmation", str(correction_confirmation), "--out-dir", str(correction_dir)])
    assert corrected["invite_id"] == json.loads(reply.read_bytes())["invite_id"]
    with pytest.raises(fill.FillError, match="CONFIRMATION_INVALID"):
        run(["vault-fill", str(request), "--previous", str(reply),
             "--confirmation", str(correction_confirmation), "--out-dir", str(correction_dir)])


def test_confirmation_is_bound_to_vault_and_consumed_on_stale(tmp_path, isolated_vault):
    _created, request, _task_dir = make_request(tmp_path)
    stage_profile(tmp_path, {
        "name": {"type": "text", "value": "张三"},
        "phone": {"type": "phone_cn", "value": "13800138000"},
    })
    confirmation = tmp_path / "private" / "submit.yintian-confirmation"
    run(["vault-preview", str(request), "--confirmation-out", str(confirmation)])
    stage_profile(tmp_path, {"phone": {"type": "phone_cn", "value": "13900139000"}})
    with pytest.raises(fill.FillError, match="CONFIRMATION_STALE"):
        run(["vault-fill", str(request), "--confirmation", str(confirmation), "--out-dir", str(tmp_path / "incoming")])
    assert not confirmation.exists()


def test_confirmation_is_bound_to_vault_and_key_paths(tmp_path, isolated_vault):
    _created, request, _task_dir = make_request(tmp_path)
    stage_profile(tmp_path, {
        "name": {"type": "text", "value": "张三"},
        "phone": {"type": "phone_cn", "value": "13800138000"},
    })
    confirmation = tmp_path / "private" / "submit-path.yintian-confirmation"
    run(["vault-preview", str(request), "--confirmation-out", str(confirmation)])
    copied_vault_dir = private_dir(tmp_path / "copied-vault")
    copied_key_dir = private_dir(tmp_path / "copied-keys")
    copied_vault = copied_vault_dir / vault.VAULT_FILENAME
    copied_key = copied_key_dir / vault.KEY_FILENAME
    copied_vault.write_bytes(vault.default_vault_path().read_bytes())
    copied_key.write_bytes(vault.default_key_path().read_bytes())
    if os.name != "nt":
        copied_vault.chmod(0o600)
        copied_key.chmod(0o600)
    with pytest.raises(fill.FillError, match="CONFIRMATION_STALE"):
        run(["vault-fill", str(request), "--confirmation", str(confirmation),
             "--out-dir", str(tmp_path / "incoming"),
             "--vault", str(copied_vault), "--key-file", str(copied_key)])
    assert not confirmation.exists()


def test_confirmation_is_bound_to_request_bytes(tmp_path, isolated_vault):
    _created, request, _task_dir = make_request(tmp_path)
    stage_profile(tmp_path, {
        "name": {"type": "text", "value": "张三"},
        "phone": {"type": "phone_cn", "value": "13800138000"},
    })
    confirmation = tmp_path / "private" / "submit.yintian-confirmation"
    run(["vault-preview", str(request), "--confirmation-out", str(confirmation)])
    changed = json.loads(request.read_text(encoding="utf-8"))
    changed["purpose"] = "被修改的用途"
    request.write_text(json.dumps(changed, ensure_ascii=False, indent=2), encoding="utf-8")
    with pytest.raises(fill.FillError, match="CONFIRMATION_STALE"):
        run(["vault-fill", str(request), "--confirmation", str(confirmation), "--out-dir", str(tmp_path / "incoming")])
    assert not confirmation.exists()


def test_tampered_and_expired_confirmations_are_rejected(tmp_path, isolated_vault, monkeypatch):
    _created, request, _task_dir = make_request(tmp_path)
    stage_profile(tmp_path, {
        "name": {"type": "text", "value": "张三"},
        "phone": {"type": "phone_cn", "value": "13800138000"},
    })
    private = tmp_path / "private"
    tampered = private / "tampered.yintian-confirmation"
    run(["vault-preview", str(request), "--confirmation-out", str(tampered)])
    blob = bytearray(tampered.read_bytes())
    blob[-2] ^= 1
    tampered.write_bytes(blob)
    with pytest.raises(fill.FillError, match="CONFIRMATION_INVALID"):
        run(["vault-fill", str(request), "--confirmation", str(tampered), "--out-dir", str(tmp_path / "incoming")])
    assert tampered.exists()

    monkeypatch.setattr(vault, "CONFIRMATION_TTL_SECONDS", -1)
    expired = private / "expired.yintian-confirmation"
    run(["vault-preview", str(request), "--confirmation-out", str(expired)])
    with pytest.raises(fill.FillError, match="CONFIRMATION_EXPIRED"):
        run(["vault-fill", str(request), "--confirmation", str(expired), "--out-dir", str(tmp_path / "incoming")])
    assert expired.exists()


def test_cancelled_stage_does_not_create_or_change_vault(tmp_path, isolated_vault):
    private = private_dir(tmp_path / "private")
    answers = private_json(private / "answers.json", {
        "entries": {"name": {"type": "text", "value": "张三"}},
    })
    confirmation = private / "cancel.yintian-confirmation"
    staged = run(["vault-stage", "--answers", str(answers), "--confirmation-out", str(confirmation)])
    assert staged["changes"][0]["action"] == "add"
    assert not answers.exists() and confirmation.exists()
    status = run(["vault-status"])
    assert status["vault"] is False
    confirmation.unlink()


def test_invalid_confirmation_path_never_deletes_vault_or_key(tmp_path, isolated_vault):
    stage_profile(tmp_path, {"name": {"type": "text", "value": "张三"}})
    vault_path = vault.default_vault_path()
    key_path = vault.default_key_path()
    before_vault, before_key = vault_path.read_bytes(), key_path.read_bytes()
    with pytest.raises(fill.FillError, match="VAULT_PATH_COLLISION"):
        run(["vault-apply", "--confirmation", str(vault_path)])
    with pytest.raises(fill.FillError, match="VAULT_PATH_COLLISION"):
        run(["vault-apply", "--confirmation", str(key_path)])
    assert vault_path.read_bytes() == before_vault and key_path.read_bytes() == before_key


def test_invalid_custom_migration_target_still_deletes_password(tmp_path):
    private = private_dir(tmp_path / "private")
    password = private_json(private / "password.json", "temporary-password")
    with pytest.raises(fill.FillError, match="VAULT_STORAGE_INVALID"):
        run(["vault-migrate", "--password-file", str(password),
             "--source-vault", str(private / "missing-v1"),
             "--target-vault", str(private / "custom-vault")])
    assert not password.exists()


def test_v1_migration_is_noninteractive_and_preserves_source(tmp_path, isolated_vault, monkeypatch):
    source_dir = private_dir(tmp_path / "legacy")
    source = source_dir / "old.yintian-vault"
    password = "Synthetic-old-password-2026"
    old = {"version": 1, "types": {"name": "text", "phone": "phone_cn"},
           "values": {"name": "张三", "phone": "13800138000"}, "attachments": {}}
    source.write_bytes(collection.canonical(collection.aes_gcm_seal(vault.FORMAT_V1, collection.canonical(old), password, vault.FORMAT_V1.encode())))
    if os.name != "nt":
        source.chmod(0o600)
    password_file = source_dir / "password.txt"
    password_file.write_text(password, encoding="utf-8")
    if os.name != "nt":
        password_file.chmod(0o600)
    target = tmp_path / "vault" / vault.VAULT_FILENAME
    key = tmp_path / "keys" / vault.KEY_FILENAME
    monkeypatch.setattr(vault, "legacy_vault_path", lambda: source)
    result = run(["vault-migrate", "--password-file", str(password_file)])
    assert result["migrated"] and source.is_file() and not password_file.exists()
    profile = vault.load_vault(target, vault.load_or_create_key(key))
    assert profile["entries"]["phone"]["value"] == "13800138000"
    assert run(["vault-status"])["format"] == vault.FORMAT_V2


def test_v1_status_requires_migration_without_creating_key(tmp_path):
    private = private_dir(tmp_path / "legacy")
    source = private / "old.yintian-vault"
    old = {"version": 1, "types": {"name": "text"}, "values": {"name": "张三"}, "attachments": {}}
    source.write_bytes(collection.canonical(collection.aes_gcm_seal(
        vault.FORMAT_V1, collection.canonical(old), "password", vault.FORMAT_V1.encode())))
    if os.name != "nt":
        source.chmod(0o600)
    key_path = private / "missing.key"
    result = run(["vault-status", "--vault", str(source), "--key-file", str(key_path)])
    assert result["migration_required"] is True
    assert not key_path.exists()


def test_wrong_v1_password_writes_nothing_and_deletes_password(tmp_path, isolated_vault):
    private = private_dir(tmp_path / "legacy")
    source = private / "old.yintian-vault"
    old = {"version": 1, "types": {"name": "text"}, "values": {"name": "张三"}, "attachments": {}}
    source.write_bytes(collection.canonical(collection.aes_gcm_seal(vault.FORMAT_V1, collection.canonical(old), "right", vault.FORMAT_V1.encode())))
    password_file = private_json(private / "password.json", "wrong")
    if os.name != "nt":
        source.chmod(0o600)
    with pytest.raises(fill.FillError, match="VAULT_UNLOCK_FAILED"):
        run(["vault-migrate", "--password-file", str(password_file), "--source-vault", str(source),
             "--target-vault", str(tmp_path / "vault" / vault.VAULT_FILENAME),
             "--target-key", str(tmp_path / "keys" / vault.KEY_FILENAME)])
    assert not password_file.exists()
    assert not (tmp_path / "vault" / vault.VAULT_FILENAME).exists()


@pytest.mark.skipif(os.name == "nt", reason="POSIX permission bits")
def test_custom_parent_permissions_are_never_changed(tmp_path):
    shared = tmp_path / "shared"
    shared.mkdir(mode=0o755)
    shared.chmod(0o755)
    private = private_dir(tmp_path / "private")
    answers = private_json(private / "answers.json", {"entries": {"name": {"type": "text", "value": "张三"}}})
    with pytest.raises(fill.FillError, match="VAULT_PERMISSIONS"):
        run(["vault-stage", "--answers", str(answers), "--confirmation-out", str(private / "change"),
             "--vault", str(shared / "vault.yintian-vault"), "--key-file", str(private / "vault.key")])
    assert not answers.exists() and not (private / "vault.key").exists()
    assert stat.S_IMODE(shared.stat().st_mode) == 0o755


def test_custom_vault_requires_explicit_key(tmp_path):
    with pytest.raises(fill.FillError, match="VAULT_STORAGE_INVALID"):
        run(["vault-status", "--vault", str(tmp_path / "vault")])


def test_missing_required_field_never_creates_confirmation(tmp_path, isolated_vault):
    _created, request, _task_dir = make_request(tmp_path)
    stage_profile(tmp_path, {"name": {"type": "text", "value": "张三"}})
    confirmation = tmp_path / "private" / "submit.yintian-confirmation"
    with pytest.raises(fill.FillError, match="VAULT_FIELDS_MISSING"):
        run(["vault-preview", str(request), "--confirmation-out", str(confirmation)])
    assert not confirmation.exists()


@pytest.mark.skipif(os.name == "nt", reason="symlink behavior differs on Windows")
def test_symlink_and_parent_traversal_are_rejected(tmp_path):
    target = tmp_path / "target"
    target.write_bytes(b"secret")
    link = tmp_path / "link"
    link.symlink_to(target)
    with pytest.raises(ValueError, match="PATH_UNSAFE"):
        secure_io.read_bytes(link)
    with pytest.raises(ValueError, match="PATH_UNSAFE"):
        secure_io.checked_path(tmp_path / "child" / ".." / "target")


def test_legacy_v2_storage_is_copied_and_source_is_kept(tmp_path, monkeypatch):
    source_dir = private_dir(tmp_path / "skill-data")
    target_vault_dir = private_dir(tmp_path / "new-vault")
    target_key_dir = private_dir(tmp_path / "new-keys")
    source_vault = source_dir / vault.VAULT_FILENAME
    source_key = source_dir / vault.KEY_FILENAME
    key = vault.load_or_create_key(source_key, create=True)
    profile = {"format": vault.FORMAT_V2, "entries": {"name": {"type": "text", "label": "", "value": "张三", "source": {"kind": "manual"}}}}
    vault.save_vault(source_vault, key, profile, create=True)
    monkeypatch.setattr(vault, "legacy_vault_path", lambda: source_vault)
    monkeypatch.setattr(vault, "legacy_key_path", lambda: source_key)
    monkeypatch.setenv("YINTIAN_VAULT_DIR", str(target_vault_dir))
    monkeypatch.setenv("YINTIAN_VAULT_KEY_DIR", str(target_key_dir))
    target_vault, target_key, migration = vault.resolve_default_storage()
    assert migration["migrated"] and source_vault.is_file() and source_key.is_file()
    assert vault.load_vault(target_vault, vault.load_or_create_key(target_key))["entries"]["name"]["value"] == "张三"
    target_vault, target_key, migration = vault.resolve_default_storage()
    assert migration is None and target_vault.is_file() and target_key.is_file()
    marker = target_vault.parent / vault.MIGRATION_MARKER
    marker.write_text("{}", encoding="utf-8")
    if os.name != "nt":
        marker.chmod(0o600)
    with pytest.raises(RuntimeError, match="VAULT_MIGRATION_CONFLICT"):
        vault.resolve_default_storage()


def test_legacy_v2_conflict_never_prefers_different_target(tmp_path, monkeypatch):
    source_dir = private_dir(tmp_path / "skill-data")
    target_vault_dir = private_dir(tmp_path / "new-vault")
    target_key_dir = private_dir(tmp_path / "new-keys")
    source_vault = source_dir / vault.VAULT_FILENAME
    source_key = source_dir / vault.KEY_FILENAME
    source_secret = vault.load_or_create_key(source_key, create=True)
    vault.save_vault(source_vault, source_secret, {
        "format": vault.FORMAT_V2,
        "entries": {"name": {"type": "text", "label": "", "value": "旧数据", "source": {"kind": "manual"}}},
    }, create=True)
    target_vault = target_vault_dir / vault.VAULT_FILENAME
    target_key = target_key_dir / vault.KEY_FILENAME
    target_secret = vault.load_or_create_key(target_key, create=True)
    vault.save_vault(target_vault, target_secret, {
        "format": vault.FORMAT_V2,
        "entries": {"name": {"type": "text", "label": "", "value": "新数据", "source": {"kind": "manual"}}},
    }, create=True)
    monkeypatch.setattr(vault, "legacy_vault_path", lambda: source_vault)
    monkeypatch.setattr(vault, "legacy_key_path", lambda: source_key)
    monkeypatch.setenv("YINTIAN_VAULT_DIR", str(target_vault_dir))
    monkeypatch.setenv("YINTIAN_VAULT_KEY_DIR", str(target_key_dir))
    with pytest.raises(RuntimeError, match="VAULT_MIGRATION_CONFLICT"):
        vault.resolve_default_storage()
    assert vault.load_vault(target_vault, target_secret)["entries"]["name"]["value"] == "新数据"


def test_v1_migration_conflict_never_overwrites_target(tmp_path):
    private = private_dir(tmp_path / "private")
    source = private / "old.yintian-vault"
    old = {"version": 1, "types": {"name": "text"}, "values": {"name": "张三"}, "attachments": {}}
    source.write_bytes(collection.canonical(collection.aes_gcm_seal(
        vault.FORMAT_V1, collection.canonical(old), "right", vault.FORMAT_V1.encode())))
    password = private / "password.txt"
    password.write_text("right", encoding="utf-8")
    target = private / "existing.yintian-vault"
    target.write_bytes(b"do-not-overwrite")
    key = private / "existing.key"
    key.write_bytes(b"do-not-overwrite")
    if os.name != "nt":
        for path in (source, password, target, key):
            path.chmod(0o600)
    with pytest.raises(fill.FillError, match="VAULT_MIGRATION_CONFLICT"):
        run(["vault-migrate", "--password-file", str(password), "--source-vault", str(source),
             "--target-vault", str(target), "--target-key", str(key)])
    assert target.read_bytes() == b"do-not-overwrite" and key.read_bytes() == b"do-not-overwrite"
    assert source.is_file() and not password.exists()


def test_v1_migration_failure_rolls_back_new_files(tmp_path, monkeypatch):
    source_dir = private_dir(tmp_path / "source")
    target_dir = private_dir(tmp_path / "target")
    key_dir = private_dir(tmp_path / "keys")
    source = source_dir / "old.yintian-vault"
    old = {"version": 1, "types": {"name": "text"}, "values": {"name": "张三"}, "attachments": {}}
    source.write_bytes(collection.canonical(collection.aes_gcm_seal(
        vault.FORMAT_V1, collection.canonical(old), "right", vault.FORMAT_V1.encode())))
    if os.name != "nt":
        source.chmod(0o600)
    def fail_save(*_args, **_kwargs):
        raise OSError("synthetic failure")

    monkeypatch.setattr(vault, "save_vault", fail_save)
    target = target_dir / vault.VAULT_FILENAME
    key = key_dir / vault.KEY_FILENAME
    with pytest.raises(OSError, match="synthetic failure"):
        vault.migrate_v1(source, "right", target, key)
    assert source.is_file() and not target.exists() and not key.exists()


def test_vlm_setup_uses_user_selected_model(tmp_path, monkeypatch):
    root = private_dir(tmp_path / "models")
    monkeypatch.setattr(vlm_extract, "model_root", lambda: root)
    calls = []

    def fake_download(**kwargs):
        calls.append(kwargs)
        Path(kwargs["local_dir"], "openvino_model.bin").write_bytes(b"synthetic-model")

    monkeypatch.setitem(sys.modules, "huggingface_hub", types.SimpleNamespace(snapshot_download=fake_download))
    installed = vlm_extract.setup("vendor/compatible-model", "user-revision")
    assert len(calls) == 1
    assert calls[0]["repo_id"] == "vendor/compatible-model"
    assert calls[0]["revision"] == "user-revision"
    assert Path(calls[0]["local_dir"]).parent == root
    assert Path(calls[0]["local_dir"]).name.startswith(".safefill-vlm-")
    assert vlm_extract.verify_model(
        "vendor/compatible-model", "user-revision", Path(installed["path"]))["revision"] == "user-revision"
    Path(installed["path"], "poison.py").write_text("raise SystemExit", encoding="utf-8")
    with pytest.raises(vlm_extract.VlmUnavailable, match="完整性"):
        vlm_extract.setup("vendor/compatible-model", "user-revision")
    assert len(calls) == 1


def test_vlm_inference_is_manifest_checked_and_forced_offline(tmp_path, monkeypatch):
    model = private_dir(tmp_path / "model")
    (model / "openvino_model.bin").write_bytes(b"synthetic-model")
    (model / vlm_extract.MANIFEST_NAME).write_bytes(collection.canonical(
        vlm_extract._manifest(model, "vendor/compatible-model", None)))
    if os.name != "nt":
        (model / vlm_extract.MANIFEST_NAME).chmod(0o600)
    image = tmp_path / "id.png"
    Image.new("RGB", (2, 2), "white").save(image)
    monkeypatch.setattr(vlm_extract, "model_dir", lambda *_args: model)

    class FakePipeline:
        def __init__(self, path, device):
            assert Path(path) == model and device == "CPU"
            assert os.environ["HF_HUB_OFFLINE"] == "1" and os.environ["TRANSFORMERS_OFFLINE"] == "1"

        def generate(self, prompt, **kwargs):
            assert prompt == vlm_extract.PROMPT and "image" in kwargs
            return '{"name":"张三","id_number":"","phone":"","address":""}'

    class FakeArray:
        def __getitem__(self, _item):
            return self

    monkeypatch.setitem(sys.modules, "numpy", types.SimpleNamespace(
        uint8="uint8", asarray=lambda *_args, **_kwargs: FakeArray()))
    monkeypatch.setitem(sys.modules, "openvino", types.SimpleNamespace(Tensor=lambda value: value))
    monkeypatch.setitem(sys.modules, "openvino_genai", types.SimpleNamespace(VLMPipeline=FakePipeline))
    result = vlm_extract.extract_fields(image, "vendor/compatible-model")
    assert result["fields"] == {"name": "张三"}


def test_legacy_open_request_and_safe_filename(tmp_path):
    _created, request, _task_dir = make_request(tmp_path)
    legacy = json.loads(request.read_text(encoding="utf-8"))
    legacy["format"] = collection.LEGACY_OPEN_FORM_FORMAT_VERSION
    legacy.pop("kind")
    legacy.pop("target_skill")
    legacy_path = tmp_path / "legacy.yintian-form"
    collection.dump_json(legacy_path, legacy)
    assert fill.load_form(legacy_path)["mode"] == "open"
    assert fill.reply_filename(" 张/三:*? ").startswith("张_三-")


def test_duplicate_names_and_corrupt_ciphertext_are_isolated(tmp_path, isolated_vault):
    _created, request, task_dir = make_request(tmp_path)
    stage_profile(tmp_path, {
        "name": {"type": "text", "value": "同名员工"},
        "phone": {"type": "phone_cn", "value": "13800138000"},
    })
    private = tmp_path / "private"
    incoming = private_dir(tmp_path / "incoming")
    replies = []
    for index in range(2):
        confirmation = private / f"submit-{index}.yintian-confirmation"
        run(["vault-preview", str(request), "--confirmation-out", str(confirmation)])
        replies.append(Path(run(["vault-fill", str(request), "--confirmation", str(confirmation),
                                 "--out-dir", str(incoming)])["out"]))
    assert replies[0].name != replies[1].name
    (incoming / "损坏-ABC123.yintian").write_bytes(b"not-json")
    result = collection.collect_open(task_dir, incoming, tmp_path / "duplicates.xlsx")
    assert result["rows"] == 2 and result["ingest"]["rejected"] == 1
    book = load_workbook(result["xlsx"], read_only=True)
    rows = list(book.active.values)
    book.close()
    assert [row[0] for row in rows[1:]] == ["同名员工", "同名员工"]


def test_open_flow_exports_encrypted_attachment(tmp_path, isolated_vault):
    config = {
        "title": "附件收集", "purpose": "归档证件", "deadline": "2098-12-31T23:59:59Z",
        "retention_until": "2099-01-31T23:59:59Z", "contact": "hr@example.com",
        "correction": "联系 HR", "template_version": "1.0",
        "fields": [
            {"id": "name", "label": "姓名", "type": "text", "required": True, "sensitive": True},
            {"id": "id_front", "label": "证件正面", "type": "image_attachment", "required": True,
             "sensitive": True, "multiple": False, "ocr_fields": [], "ocr_backend": "manual"},
        ],
    }
    config_path = tmp_path / "attachment-config.json"
    collection.dump_json(config_path, config)
    args = collection.build_parser().parse_args(
        ["create-request", "--config", str(config_path), "--out", str(tmp_path / "tasks")])
    created = args.func(args)
    request, task_dir = Path(created["request"]), Path(created["task_dir"])
    image = tmp_path / "synthetic-id.png"
    Image.new("RGB", (2, 2), "white").save(image)
    stage_profile(tmp_path, {
        "name": {"type": "text", "value": "张三"},
        "id_front": {"type": "image_attachment", "paths": [str(image)]},
    })
    private = tmp_path / "private"
    confirmation = private / "attachment.yintian-confirmation"
    preview = run(["vault-preview", str(request), "--confirmation-out", str(confirmation)])
    attachment = next(field["value"] for field in preview["fields"] if field["id"] == "id_front")
    assert attachment[0]["name"] == image.name and "data_b64" not in attachment[0]
    incoming = private_dir(tmp_path / "incoming-attachments")
    reply = Path(run(["vault-fill", str(request), "--confirmation", str(confirmation),
                      "--out-dir", str(incoming)])["out"])
    assert image.read_bytes() not in reply.read_bytes()
    exported = collection.collect_open(task_dir, incoming, tmp_path / "attachments.xlsx")
    exported_dir = Path(exported["attachments_dir"])
    extracted = list(exported_dir.rglob("*.png"))
    assert len(extracted) == 1 and extracted[0].read_bytes() == image.read_bytes()


@pytest.mark.parametrize("mode", ["group", "directed"])
def test_authenticated_legacy_modes_still_roundtrip(tmp_path, isolated_vault, mode):
    fields = [
        {"id": "name", "label": "姓名", "type": "text", "required": True, "sensitive": True},
        {"id": "phone", "label": "手机号", "type": "phone_cn", "required": True, "sensitive": True},
    ]
    if mode == "group":
        fields.insert(0, {"id": "employee_id", "label": "工号", "type": "text", "required": True, "sensitive": False})
    config = {
        "title": "兼容流程", "purpose": "兼容性测试", "deadline": "2098-12-31T23:59:59Z",
        "retention_until": "2099-01-31T23:59:59Z", "contact": "hr@example.com",
        "correction": "联系 HR", "template_version": "1.0", "fields": fields,
    }
    config_path = tmp_path / "legacy-config.json"
    collection.dump_json(config_path, config)
    roster = tmp_path / "roster.csv"
    roster.write_text("employee_id,name\nE001,张三\n", encoding="utf-8")
    password = "Synthetic-task-password-2026"
    created = collection.create_task(roster, config_path, tmp_path / "tasks", password, require_terminal=False, mode=mode)
    task_dir = Path(created["task_dir"])
    if mode == "group":
        form = Path(created["form"])
        credential = task_dir / "credentials" / "GRP-E001.yintian-credential"
    else:
        form = next((task_dir / "invites").glob("*.yintian-form"))
        credential = None
    stage_profile(tmp_path, {"phone": {"type": "phone_cn", "value": "13800138000"}})
    private = tmp_path / "private"
    if credential:
        stale_confirmation = private / "credential-stale.yintian-confirmation"
        run(["vault-preview", str(form), "--credential", str(credential),
             "--confirmation-out", str(stale_confirmation)])
        original_credential = credential.read_bytes()
        credential.write_text(json.dumps(json.loads(original_credential), ensure_ascii=False, indent=2), encoding="utf-8")
        with pytest.raises(fill.FillError, match="CONFIRMATION_STALE"):
            run(["vault-fill", str(form), "--credential", str(credential),
                 "--confirmation", str(stale_confirmation), "--out-dir", str(private / "stale")])
        credential.write_bytes(original_credential)
    confirmation = private / f"{mode}.yintian-confirmation"
    preview_args = ["vault-preview", str(form), "--confirmation-out", str(confirmation)]
    if credential:
        preview_args += ["--credential", str(credential)]
    preview = run(preview_args)
    assert {item["id"]: item["value"] for item in preview["fields"]}["name"] == "张三"
    incoming = private_dir(tmp_path / "incoming")
    fill_args = ["vault-fill", str(form), "--confirmation", str(confirmation), "--out-dir", str(incoming)]
    if credential:
        fill_args += ["--credential", str(credential)]
    run(fill_args)
    assert collection.ingest_task(task_dir, incoming)["accepted"] == 1
    assert collection.review_task(task_dir, password)["verified"] == 1
