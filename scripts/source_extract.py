"""隐填 · 本地源文件文本提取

仅供本地复核组件把授权材料转成文本，不提供命令行或 MCP 明文输出。
支持常见文本、JSON/CSV、PDF、Word、Excel 和图片。图片会走本地 OCR,
不会上传。
"""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

import privacy

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".log"}
JSON_SUFFIXES = {".json"}
CSV_SUFFIXES = {".csv", ".tsv"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}


def _read_text(path: Path, max_chars: int) -> str:
    return path.read_text(encoding="utf-8", errors="replace")[:max_chars]


def _read_json(path: Path, max_chars: int) -> str:
    data = json.loads(path.read_text(encoding="utf-8"))
    return json.dumps(data, ensure_ascii=False, indent=2)[:max_chars]


def _read_csv(path: Path, max_chars: int) -> str:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    rows: list[str] = []
    with path.open("r", encoding="utf-8", errors="replace", newline="") as fh:
        reader = csv.reader(fh, delimiter=delimiter)
        for row in reader:
            rows.append(" | ".join(row))
            if sum(len(r) for r in rows) >= max_chars:
                break
    return "\n".join(rows)[:max_chars]


def _read_pdf(path: Path, max_chars: int) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise RuntimeError("读取 PDF 需要安装 pypdf: pip install pypdf") from exc

    reader = PdfReader(str(path))
    parts: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"[page {i}]\n{text}")
        if sum(len(p) for p in parts) >= max_chars:
            break
    return "\n\n".join(parts)[:max_chars]


def _read_docx(path: Path, max_chars: int) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise RuntimeError("读取 DOCX 需要安装 python-docx: pip install python-docx") from exc

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)[:max_chars]


def _read_xlsx(path: Path, max_chars: int) -> str:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise RuntimeError("读取 XLSX 需要安装 openpyxl: pip install openpyxl") from exc

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"[sheet {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]
            if any(vals):
                parts.append(" | ".join(vals))
            if sum(len(p) for p in parts) >= max_chars:
                return "\n".join(parts)[:max_chars]
    return "\n".join(parts)[:max_chars]


def _read_image(path: Path, max_chars: int) -> dict[str, Any]:
    import ocr_engine

    result = ocr_engine.scan(str(path))
    return {
        "text": result.get("full_text", "")[:max_chars],
        "ocr_blocks": result.get("blocks", []),
        "ocr_runtime": {
            "requested_device": result.get("requested_device"),
            "execution_devices": result.get("execution_devices", []),
            "elapsed_ms": result.get("elapsed_ms"),
        },
    }


def extract_text(file_path: str, max_chars: int = 40000) -> dict[str, Any]:
    """Extract local text from one source file."""
    path = Path(file_path).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {file_path}")

    privacy.assert_in_vault(str(path))
    suffix = path.suffix.lower()

    if suffix in TEXT_SUFFIXES:
        text = _read_text(path, max_chars)
        kind = "text"
    elif suffix in JSON_SUFFIXES:
        text = _read_json(path, max_chars)
        kind = "json"
    elif suffix in CSV_SUFFIXES:
        text = _read_csv(path, max_chars)
        kind = "table"
    elif suffix == ".pdf":
        text = _read_pdf(path, max_chars)
        kind = "pdf"
    elif suffix == ".docx":
        text = _read_docx(path, max_chars)
        kind = "docx"
    elif suffix == ".xlsx":
        text = _read_xlsx(path, max_chars)
        kind = "xlsx"
    elif suffix in IMAGE_SUFFIXES:
        image_result = _read_image(path, max_chars)
        text = image_result["text"]
        kind = "image_ocr"
    else:
        raise ValueError(f"暂不支持的文件类型: {suffix or '(无后缀)'}")

    output = {
        "path": str(path),
        "kind": kind,
        "chars": len(text),
        "text": text,
    }
    if kind == "image_ocr":
        output.update({"ocr_blocks": image_result["ocr_blocks"], "ocr_runtime": image_result["ocr_runtime"]})
    return output


def extract_files(file_paths: list[str], max_chars_per_file: int = 40000) -> dict[str, Any]:
    # 单个文件失败不拖垮整批:文本材料照常处理,失败项(如缺 OCR 后端的图片)记为 error。
    files: list[dict[str, Any]] = []
    for p in file_paths:
        try:
            files.append(extract_text(p, max_chars=max_chars_per_file))
        except Exception as exc:
            files.append({"path": p, "kind": "error", "chars": 0, "text": "", "error": str(exc)})
    return {
        "count": len(files),
        "files": files,
        "combined_text": "\n\n".join(
            f"===== {item['path']} ({item['kind']}) =====\n{item['text']}"
            for item in files
        ),
    }
