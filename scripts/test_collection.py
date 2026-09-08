"""端到端收集工作流回归检查（pytest 可收集，也可直接运行）。

每个 test_ 函数用独立的 tmp_path 隔离；文件底部的 main() 依次调用全部
test_ 函数，使 `python scripts/test_collection.py` 仍可脱离 pytest 运行。
"""
from __future__ import annotations

import base64
import builtins
import contextlib
import csv
import inspect
import io
import json
import os
import stat
import struct
import sys
import tempfile
import types
import warnings
import zipfile
from pathlib import Path

import collection


class TtyCapture(io.StringIO):
    def isatty(self) -> bool:
        return True

PASSWORD = "CorrectHorseBatteryStaple-2026"
VALID_ID = "110105198503071230"


def invite_token(task_dir: Path, invite_id: str) -> str:
    html = (task_dir / "invites" / f"{invite_id}.html").read_text(encoding="utf-8")
    marker = '<script id="cfg" type="application/json">'
    return json.loads(html.split(marker, 1)[1].split("</script>", 1)[0])["invite_token"]


def wrap_payload(task_dir: Path, invite_id: str, payload: dict) -> dict:
    """把载荷按 yintian-submission/2 加密成信封（RSA-OAEP 包裹随机 AES-256 密钥，AAD 绑定信封头）。"""
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import padding
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    task = collection.load_json(task_dir / "task.json")
    aes = AESGCM.generate_key(bit_length=256)
    iv = bytes(range(12))
    envelope = {
        "format_version": task["format_version"], "task_id": task["task_id"], "invite_id": invite_id,
        "schema_hash": task["schema_hash"], "key_id": task["key_id"],
        "algorithms": {"content": "AES-256-GCM", "key_wrap": "RSA-OAEP-3072-SHA256"},
    }
    public = serialization.load_pem_public_key((task_dir / "public.pem").read_bytes())
    wrapped = public.encrypt(aes, padding.OAEP(mgf=padding.MGF1(hashes.SHA256()), algorithm=hashes.SHA256(), label=None))
    ciphertext = AESGCM(aes).encrypt(iv, json.dumps(payload, ensure_ascii=False).encode(), collection.aad_for(envelope))
    envelope.update(encrypted_key_b64=base64.b64encode(wrapped).decode(), iv_b64=base64.b64encode(iv).decode(), ciphertext_b64=base64.b64encode(ciphertext).decode())
    return envelope


def make_envelope(task_dir: Path, invite_id: str, values: dict, attachments: dict | None = None, overrides: dict | None = None) -> dict:
    task = collection.load_json(task_dir / "task.json")
    payload = {
        "format_version": task["format_version"],
        "task_id": task["task_id"],
        "invite_id": invite_id,
        "invite_token": invite_token(task_dir, invite_id),
        "schema_hash": task["schema_hash"],
        "notice_hash": task["notice_hash"],
        "template_version": task["template_version"],
        "submitted_at": collection.now_iso(),
        "consent_confirmed": True,
        "values": values,
        "attachments": attachments or {"id_front": [], "id_back": []},
    }
    if overrides:
        payload.update(overrides)
    return wrap_payload(task_dir, invite_id, payload)


def make_scenario(tmp: Path) -> types.SimpleNamespace:
    """创建一个任务目录（含一名员工与一份邀请），返回常用路径与常量。"""
    config = collection.default_config()
    config["purpose"] = "为依法办理员工商业保险收集必要身份资料"
    config["contact"] = "人事部王老师，内线 8001"
    config["correction"] = "在截止日前联系人事部王老师撤回原提交并重新提交"
    config["deadline"] = "2098-12-31T23:59:59Z"
    config["retention_until"] = "2099-12-31T23:59:59Z"
    config["fields"].append({"id": "effective_date", "label": "生效日期", "type": "date", "required": False})
    config_path = tmp / "config.json"
    collection.dump_json(config_path, config)
    roster = tmp / "roster.csv"
    with roster.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream); writer.writerow(["employee_id", "name"]); writer.writerow(["=1+1", "张三"])
    created = collection.create_task(roster, config_path, tmp / "tasks", PASSWORD, require_terminal=False)
    task_dir = Path(created["task_dir"])
    invite_id = next(csv.DictReader((task_dir / "invite-index.csv").open(encoding="utf-8-sig")))["invite_id"]
    values = {"name": "张三", "phone": "13800138000", "id_number": VALID_ID, "address": "北京市朝阳区", "effective_date": "2024-02-30"}
    incoming = tmp / "incoming"; incoming.mkdir()
    return types.SimpleNamespace(
        tmp=tmp, config_path=config_path, roster=roster, created=created,
        task_dir=task_dir, task_id=created["task_id"], invite_id=invite_id,
        values=values, incoming=incoming,
    )


def submit(ns: types.SimpleNamespace, filename: str, **kwargs) -> None:
    collection.dump_json(ns.incoming / filename, make_envelope(ns.task_dir, ns.invite_id, ns.values, **kwargs))


def submit_accepted(ns: types.SimpleNamespace, filename: str = "one.yintian", **kwargs) -> None:
    submit(ns, filename, **kwargs)
    assert collection.ingest_task(ns.task_dir, ns.incoming)["accepted"] == 1


def make_package(ns: types.SimpleNamespace) -> Path:
    """生成明文 yintian-task/2 交接包，供兼容导入与 ZIP 结构攻击测试复用。"""
    package = ns.tmp / "task.yintian-task"
    zip_bytes, _ = collection.build_task_package(ns.task_dir)
    package.write_bytes(zip_bytes)
    os.chmod(package, 0o600)
    return package


def package_parts(package: Path) -> tuple[dict, dict]:
    with zipfile.ZipFile(package) as source:
        metadata = json.loads(source.read("package.json"))
        members = {name: source.read(name) for name in source.namelist() if name != "package.json"}
    return metadata, members


def write_package(path: Path, metadata: dict, members: dict, extra: dict | None = None) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("package.json", json.dumps(metadata))
        for name, data in members.items():
            archive.writestr(name, data)
        for name, data in (extra or {}).items():
            archive.writestr(name, data)


def expire_task(task_dir: Path) -> None:
    doc = collection.load_json(task_dir / "task.json")
    doc["deadline"] = "2020-01-01T00:00:00Z"
    doc["retention_until"] = "2020-02-01T00:00:00Z"
    doc["notice_hash"] = collection.sha256_bytes(collection.canonical({key: doc[key] for key in ("title", "purpose", "deadline", "retention_until", "contact", "correction")}))
    collection.dump_json(task_dir / "task.json", doc)


def assert_raises(expected, action, message: str) -> None:
    try:
        action()
    except expected:
        return
    raise AssertionError(message)


def forbidden_prompt(*args, **kwargs):
    raise AssertionError("被拒绝的操作不应出现交互提示")


@contextlib.contextmanager
def fake_tty(input_func=None, getpass_func=None, capture=None):
    real_stdin, real_stdout = sys.stdin, sys.stdout
    real_getpass, real_input = collection.getpass.getpass, builtins.input
    capture = capture if capture is not None else TtyCapture()
    sys.stdin, sys.stdout = types.SimpleNamespace(isatty=lambda: True), capture
    collection.getpass.getpass = getpass_func if getpass_func is not None else real_getpass
    builtins.input = input_func if input_func is not None else real_input
    try:
        yield capture
    finally:
        sys.stdin, sys.stdout = real_stdin, real_stdout
        collection.getpass.getpass, builtins.input = real_getpass, real_input


@contextlib.contextmanager
def non_tty():
    real_stdin, real_stdout = sys.stdin, sys.stdout
    real_getpass, real_input = collection.getpass.getpass, builtins.input
    sys.stdin, sys.stdout = types.SimpleNamespace(isatty=lambda: False), io.StringIO()
    collection.getpass.getpass = forbidden_prompt
    builtins.input = forbidden_prompt
    try:
        yield
    finally:
        sys.stdin, sys.stdout = real_stdin, real_stdout
        collection.getpass.getpass, builtins.input = real_getpass, real_input


def test_placeholder_config_rejected(tmp_path: Path) -> None:
    assert_raises(ValueError, lambda: collection.validate_config(collection.default_config()), "占位配置不应创建正式任务")


def test_create_task_layout(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    assert len(list((ns.task_dir / "invites").glob("*.html"))) == 1


def test_ingest_and_dedup(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit(ns, "one.yintian")
    first = collection.ingest_task(ns.task_dir, ns.incoming)
    assert first["accepted"] == 1
    assert collection.ingest_task(ns.task_dir, ns.incoming)["duplicates"] == 1


def test_review_flags_missing_attachments(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    reviewed = collection.review_task(ns.task_dir, PASSWORD)
    assert reviewed["needs_review"] == 1  # 必填证件附件缺失
    rows = collection.report_rows(ns.task_dir)
    assert rows[0]["status"] == "needs_review" and "id_front" in rows[0]["missing_fields"]
    assert "effective_date" in rows[0]["conflict_fields"]


def test_reports_redacted_and_formula_safe(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    collection.review_task(ns.task_dir, PASSWORD)
    reports = collection.write_reports(ns.task_dir, ["json", "xlsx"])
    report_text = Path(reports["json"]).read_text(encoding="utf-8")
    assert VALID_ID not in report_text and "13800138000" not in report_text and "北京市朝阳区" not in report_text
    sheet = load_workbook(reports["xlsx"], data_only=False).active
    assert sheet["A2"].data_type == "s" and sheet["A2"].value == "=1+1"


def test_unsafe_submitted_at_neutralized(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    collection.review_task(ns.task_dir, PASSWORD)
    unsafe_time = '=HYPERLINK("https://example.invalid")'
    submit(ns, "unsafe-time.yintian", overrides={"submitted_at": unsafe_time})
    assert collection.ingest_task(ns.task_dir, ns.incoming)["accepted"] == 1
    assert collection.review_task(ns.task_dir, PASSWORD)["needs_review"] == 1
    reports = collection.write_reports(ns.task_dir, ["json", "xlsx"])
    assert unsafe_time not in Path(reports["json"]).read_text(encoding="utf-8")
    assert load_workbook(reports["xlsx"], data_only=False).active["E2"].value is None


def test_wrong_token_marked_invalid(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit(ns, "wrong-token.yintian", overrides={"invite_token": "wrong"})
    assert collection.ingest_task(ns.task_dir, ns.incoming)["accepted"] == 1
    assert collection.review_task(ns.task_dir, PASSWORD)["invalid"] == 1


def test_fake_image_attachment_marked_invalid(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    fake_image = b"not-a-jpeg"
    bad_attachments = {
        "id_front": [{"name": "front.jpg", "type": "image/jpeg", "size": len(fake_image), "sha256": collection.sha256_bytes(fake_image), "data_b64": base64.b64encode(fake_image).decode()}],
        "id_back": [],
    }
    submit(ns, "fake-image.yintian", attachments=bad_attachments)
    assert collection.ingest_task(ns.task_dir, ns.incoming)["accepted"] == 1
    assert collection.review_task(ns.task_dir, PASSWORD)["invalid"] == 1


def test_compare_ocr_field_matching(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    original_ocr = collection.ocr_attachment
    collection.ocr_attachment = lambda item: ("", ["ocr_no_text"])
    try:
        task_fields = collection.load_json(ns.task_dir / "task.json")["fields"]
        assert set(collection.compare_ocr({"values": ns.values}, [{"field_id": "id_front"}], task_fields)) == {"ocr_no_identity_fields", "ocr_no_text"}
        assert collection.compare_ocr({"values": ns.values}, [{"field_id": "profile_photo"}], task_fields) == []
        custom_fields = [
            {"id": "name", "label": "姓名", "type": "text"},
            {"id": "mobile", "label": "手机号", "type": "phone_cn"},
            {"id": "credential", "label": "证件号", "type": "cn_id"},
            {"id": "residence", "label": "住址", "type": "address"},
            {"id": "card_front", "label": "证件正面", "type": "image_attachment"},
            {"id": "card_back", "label": "证件反面", "type": "image_attachment"},
        ]
        collection.ocr_attachment = lambda item: (f"姓名 张三\n住址 北京市海淀区\n公民身份号码 {VALID_ID}\n联系电话 13900139000\n", [])
        custom_values = {"name": "张三", "mobile": "13800138000", "credential": VALID_ID, "residence": "北京市朝阳区"}
        assert set(collection.compare_ocr({"values": custom_values}, [{"field_id": "card_front"}], custom_fields)) == {"mobile", "residence"}
    finally:
        collection.ocr_attachment = original_ocr


def test_tampered_ciphertext_does_not_overwrite(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns, "one.yintian")
    submit_accepted(ns, "second.yintian")
    collection.review_task(ns.task_dir, PASSWORD)
    tampered = make_envelope(ns.task_dir, ns.invite_id, ns.values)
    raw = bytearray(base64.b64decode(tampered["ciphertext_b64"])); raw[-1] ^= 1
    tampered["ciphertext_b64"] = base64.b64encode(raw).decode()
    collection.dump_json(ns.incoming / "two.yintian", tampered)
    assert collection.ingest_task(ns.task_dir, ns.incoming)["accepted"] == 1
    assert collection.review_task(ns.task_dir, PASSWORD)["invalid"] == 1
    assert collection.report_rows(ns.task_dir)[0]["submission_version"] == 2  # 无效新版本不覆盖旧有效版本


def test_ingest_rejects_unreadable_files(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    (ns.incoming / "bad.yintian").mkdir()
    (ns.incoming / "13800138000.yintian").write_text("{", encoding="utf-8")
    with (ns.incoming / "large.yintian").open("wb") as stream:
        stream.truncate(collection.MAX_ENVELOPE_BYTES + 1)
    original_hash = collection.sha256_file

    def failing_hash(path: Path) -> str:
        if path.name == "13800138000.yintian":
            raise OSError(f"cannot read {path}")
        return original_hash(path)

    collection.sha256_file = failing_hash
    try:
        rejected = collection.ingest_task(ns.task_dir, ns.incoming)
    finally:
        collection.sha256_file = original_hash
    assert rejected["rejected"] == 3 and "13800138000" not in json.dumps(rejected)


def test_export_blocks_symlink_escape(tmp_path: Path) -> None:
    if os.name == "nt":
        return
    ns = make_scenario(tmp_path)
    outside_link = ns.task_dir / "outside-link"
    outside_link.symlink_to(ns.tmp / "roster.csv")
    try:
        assert_raises(ValueError, lambda: collection.export_task(ns.task_dir, ns.tmp / "task.yintian-task"), "任务目录内的符号链接必须阻止导出")
    finally:
        outside_link.unlink()


def test_export_refuses_overwriting_task_files(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    private_before = (ns.task_dir / "private.pem.enc").read_bytes()
    assert_raises(ValueError, lambda: collection.export_task(ns.task_dir, ns.task_dir / "private.pem.enc"), "导出不应覆盖任务自身文件")
    assert (ns.task_dir / "private.pem.enc").read_bytes() == private_before


def test_export_respects_import_file_limit(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    old_file_limit = collection.MAX_PACKAGE_FILES
    collection.MAX_PACKAGE_FILES = 0
    try:
        assert_raises(ValueError, lambda: collection.export_task(ns.task_dir, ns.tmp / "over-limit.yintian-task"), "导出必须遵守导入端文件数上限")
    finally:
        collection.MAX_PACKAGE_FILES = old_file_limit


def test_export_import_roundtrip(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_package(ns)
    with zipfile.ZipFile(package) as archive:
        assert not any(name.endswith("outside-link") for name in archive.namelist())
    imported = collection.import_task(package, ns.tmp / "imported")
    assert collection.status_task(imported["task_dir"])["total"] == 1
    assert collection.status_task(imported["task_dir"])["task_status"] == "active"


def test_import_legacy_windows_package(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_package(ns)
    legacy_windows = ns.tmp / "legacy-windows.yintian-task"
    with zipfile.ZipFile(package) as source, zipfile.ZipFile(legacy_windows, "w") as archive:
        metadata = json.loads(source.read("package.json"))
        metadata["package_version"] = collection.LEGACY_TASK_PACKAGE_VERSION
        metadata["manifest"] = {rel.replace("/", "\\"): digest for rel, digest in metadata["manifest"].items()}
        archive.writestr("package.json", json.dumps(metadata))
        for raw_rel in metadata["manifest"]:
            rel = raw_rel.replace("\\", "/")
            archive.writestr(f"{ns.task_id}\\{raw_rel}", source.read(f"{ns.task_id}/{rel}"))
    legacy_imported = collection.import_task(legacy_windows, ns.tmp / "legacy-windows-import")
    assert collection.status_task(legacy_imported["task_dir"])["total"] == 1


def test_import_incomplete_package_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    incomplete = ns.tmp / "incomplete.yintian-task"
    task_bytes = (ns.task_dir / "task.json").read_bytes()
    incomplete_metadata = {"package_version": collection.TASK_PACKAGE_VERSION, "task_id": ns.task_id, "manifest": {"task.json": collection.sha256_bytes(task_bytes)}}
    with zipfile.ZipFile(incomplete, "w") as archive:
        archive.writestr("package.json", json.dumps(incomplete_metadata))
        archive.writestr(f"{ns.task_id}/task.json", task_bytes)
    incomplete_parent = ns.tmp / "incomplete-import"
    assert_raises(ValueError, lambda: collection.import_task(incomplete, incomplete_parent), "缺少必要文件的任务包不应导入")
    assert not (incomplete_parent / ns.task_id).exists()


def test_import_task_id_mismatch_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    mismatch = ns.tmp / "mismatch.yintian-task"
    inner_task = collection.load_json(ns.task_dir / "task.json")
    inner_task["task_id"] = "YT-20990101-ABCDEF"
    required_payloads = {
        "task.json": json.dumps(inner_task, ensure_ascii=False).encode(),
        **{name: (ns.task_dir / name).read_bytes() for name in ("public.pem", "private.pem.enc", "roster.csv", "invite-index.csv", "state.sqlite3")},
    }
    mismatch_metadata = {"package_version": collection.TASK_PACKAGE_VERSION, "task_id": ns.task_id, "manifest": {name: collection.sha256_bytes(data) for name, data in required_payloads.items()}}
    with zipfile.ZipFile(mismatch, "w") as archive:
        archive.writestr("package.json", json.dumps(mismatch_metadata))
        for name, data in required_payloads.items():
            archive.writestr(f"{ns.task_id}/{name}", data)
    mismatch_parent = ns.tmp / "mismatch-import"
    assert_raises(ValueError, lambda: collection.import_task(mismatch, mismatch_parent), "任务包内外 task_id 不一致时不应导入")
    assert not (mismatch_parent / ns.task_id).exists()


def test_import_invalid_content_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    invalid_content = ns.tmp / "invalid-content.yintian-task"
    invalid_payloads = {
        "task.json": json.dumps({"format_version": collection.FORMAT_VERSION, "task_id": ns.task_id}).encode(),
        **{name: (ns.task_dir / name).read_bytes() for name in ("public.pem", "private.pem.enc", "roster.csv", "invite-index.csv", "state.sqlite3")},
    }
    invalid_metadata = {"package_version": collection.TASK_PACKAGE_VERSION, "task_id": ns.task_id, "manifest": {name: collection.sha256_bytes(data) for name, data in invalid_payloads.items()}}
    with zipfile.ZipFile(invalid_content, "w") as archive:
        archive.writestr("package.json", json.dumps(invalid_metadata))
        for name, data in invalid_payloads.items():
            archive.writestr(f"{ns.task_id}/{name}", data)
    invalid_parent = ns.tmp / "invalid-content-import"
    assert_raises(ValueError, lambda: collection.import_task(invalid_content, invalid_parent), "内容残缺的任务包不应导入")
    assert not (invalid_parent / ns.task_id).exists()


def test_legacy_format_version_readonly(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    task_document = collection.load_json(ns.task_dir / "task.json")
    task_document["format_version"] = collection.LEGACY_FORMAT_VERSION
    collection.dump_json(ns.task_dir / "task.json", task_document)
    try:
        assert collection.status_task(ns.task_dir)["total"] == 1
        assert_raises(RuntimeError, lambda: collection.ingest_task(ns.task_dir, ns.incoming), "旧版任务不应继续接收提交")
    finally:
        task_document["format_version"] = collection.FORMAT_VERSION
        collection.dump_json(ns.task_dir / "task.json", task_document)


def test_wrong_password_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    assert_raises(Exception, lambda: collection.unlock_private_key(ns.task_dir, "wrong-password"), "错误密码不应解锁")


def test_file_permissions(tmp_path: Path) -> None:
    if os.name == "nt":
        return
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    collection.review_task(ns.task_dir, PASSWORD)
    reports = collection.write_reports(ns.task_dir, ["json", "xlsx"])
    package = make_package(ns)
    assert stat.S_IMODE(ns.task_dir.stat().st_mode) == 0o700
    for path in (ns.task_dir / "roster.csv", ns.task_dir / "invite-index.csv", ns.task_dir / "state.sqlite3", Path(reports["xlsx"]), package):
        assert stat.S_IMODE(path.stat().st_mode) == 0o600


def test_algorithm_mismatch_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns, "one.yintian")
    submit_accepted(ns, "second.yintian")
    collection.review_task(ns.task_dir, PASSWORD)
    bad_algo_dir = ns.tmp / "bad-algorithms"; bad_algo_dir.mkdir()
    wrong_algo = make_envelope(ns.task_dir, ns.invite_id, ns.values)
    wrong_algo["algorithms"] = {"content": "AES-256-CBC", "key_wrap": "RSA-OAEP-3072-SHA256"}
    collection.dump_json(bad_algo_dir / "wrong-algo.yintian", wrong_algo)
    missing_algo = make_envelope(ns.task_dir, ns.invite_id, ns.values)
    del missing_algo["algorithms"]
    collection.dump_json(bad_algo_dir / "missing-algo.yintian", missing_algo)
    algo_result = collection.ingest_task(ns.task_dir, bad_algo_dir)
    assert algo_result["rejected"] == 2 and algo_result["accepted"] == 0
    current_row = collection.report_rows(ns.task_dir)[0]
    assert current_row["submission_version"] == 2 and current_row["status"] == "needs_review"  # 算法不匹配的提交不覆盖已有有效版本


def test_import_unsafe_paths_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_package(ns)
    base_metadata, base_members = package_parts(package)
    for label, attack_rel in (("traversal", "../evil.txt"), ("absolute", "/etc/passwd")):
        attack_metadata = json.loads(json.dumps(base_metadata))
        attack_metadata["manifest"][attack_rel] = collection.sha256_bytes(b"evil")
        attack_package = ns.tmp / f"{label}.yintian-task"
        write_package(attack_package, attack_metadata, base_members, extra={f"{ns.task_id}/{attack_rel}": b"evil"})
        assert_raises(ValueError, lambda: collection.import_task(attack_package, ns.tmp / f"{label}-import"), f"包含不安全路径的任务包不应导入: {label}")
        assert not (ns.tmp / f"{label}-import" / ns.task_id).exists()


def test_import_duplicate_entries_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_package(ns)
    base_metadata, base_members = package_parts(package)
    duplicate = ns.tmp / "duplicate.yintian-task"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        write_package(duplicate, base_metadata, base_members, extra={f"{ns.task_id}/roster.csv": base_members[f"{ns.task_id}/roster.csv"]})
    assert_raises(ValueError, lambda: collection.import_task(duplicate, ns.tmp / "duplicate-import"), "包含重复条目的任务包不应导入")
    assert not (ns.tmp / "duplicate-import" / ns.task_id).exists()


def test_import_oversized_member_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_package(ns)
    old_member_limit = collection.MAX_PACKAGE_MEMBER_BYTES
    collection.MAX_PACKAGE_MEMBER_BYTES = 4096
    try:
        assert_raises(ValueError, lambda: collection.import_task(package, ns.tmp / "oversized-import"), "超过单项上限的任务包不应导入")
        assert not (ns.tmp / "oversized-import" / ns.task_id).exists()
    finally:
        collection.MAX_PACKAGE_MEMBER_BYTES = old_member_limit


def test_import_faked_member_size_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_package(ns)
    raw_package = bytearray(package.read_bytes())
    pos = 0
    while True:
        pos = raw_package.find(b"PK\x01\x02", pos)
        if pos < 0:
            raise AssertionError("未找到目标中央目录条目")
        name_length, extra_length, comment_length = struct.unpack_from("<HHH", raw_package, pos + 28)
        if bytes(raw_package[pos + 46 : pos + 46 + name_length]).decode() == f"{ns.task_id}/state.sqlite3":
            struct.pack_into("<I", raw_package, pos + 24, 1)
            break
        pos += 46 + name_length + extra_length + comment_length
    fake_size = ns.tmp / "fake-size.yintian-task"
    fake_size.write_bytes(bytes(raw_package))
    assert_raises((ValueError, zipfile.BadZipFile), lambda: collection.import_task(fake_size, ns.tmp / "fake-size-import"), "声明大小作假的任务包不应导入")
    assert not (ns.tmp / "fake-size-import" / ns.task_id).exists()


def test_read_package_member_quota(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_package(ns)
    with zipfile.ZipFile(package) as archive:
        assert_raises(ValueError, lambda: collection.read_package_member(archive, f"{ns.task_id}/state.sqlite3", 10), "流式计数必须拒绝超过剩余配额的成员")
        assert collection.read_package_member(archive, f"{ns.task_id}/roster.csv", collection.MAX_PACKAGE_BYTES) == (ns.task_dir / "roster.csv").read_bytes()


def test_reveal_sanitizes_control_chars(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    collection.review_task(ns.task_dir, PASSWORD)
    capture = TtyCapture()
    real_decrypt, real_validate = collection.decrypt_envelope, collection.validate_payload
    collection.decrypt_envelope = lambda *args, **kwargs: {"values": {"name": "张三"}, "attachments": {"id_front": [{"name": "front.jpg", "type": "image/jpeg", "size": "5\x1b[31mred"}]}}
    collection.validate_payload = lambda *args, **kwargs: ([], [], [])
    try:
        with fake_tty(input_func=lambda prompt="": "", getpass_func=lambda prompt="": PASSWORD, capture=capture):
            assert collection.cmd_reveal(types.SimpleNamespace(task_dir=str(ns.task_dir), invite_id=ns.invite_id)) is None
    finally:
        collection.decrypt_envelope, collection.validate_payload = real_decrypt, real_validate
    shown = capture.getvalue()
    assert "\x1b[31m" not in shown and "5 [31mred bytes" in shown


def test_expired_task_rejects_all_actions(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    expire_task(ns.task_dir)
    for action in (
        lambda: collection.ingest_task(ns.task_dir, ns.incoming),
        lambda: collection.review_task(ns.task_dir, PASSWORD),
        lambda: collection.write_reports(ns.task_dir, ["json"]),
        lambda: collection.export_task(ns.task_dir, ns.tmp / "expired.yintian-task"),
    ):
        assert_raises(RuntimeError, action, "过期任务必须拒绝接收、复核、报告与导出")
    assert collection.status_task(ns.task_dir)["task_status"] == "expired"


def test_expired_task_rejects_reveal_without_prompt(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    collection.review_task(ns.task_dir, PASSWORD)
    expire_task(ns.task_dir)
    with fake_tty(input_func=forbidden_prompt, getpass_func=forbidden_prompt):
        assert_raises(RuntimeError, lambda: collection.cmd_reveal(types.SimpleNamespace(task_dir=str(ns.task_dir), invite_id=ns.invite_id)), "过期任务必须拒绝查看明文")


def test_purge_refuses_unexpired_task_without_allow_early(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    with fake_tty(input_func=forbidden_prompt, getpass_func=forbidden_prompt):
        assert_raises(RuntimeError, lambda: collection.cmd_purge(types.SimpleNamespace(task_dir=str(ns.task_dir), allow_early=False)), "任务未过期且不带 --allow-early 时 purge 必须拒绝")
    assert ns.task_dir.is_dir()
    assert not (ns.task_dir.parent / f"{ns.task_id}.purged.json").exists()


def test_purge_requires_matching_task_id(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    expire_task(ns.task_dir)
    with fake_tty(input_func=lambda prompt="": "WRONG-TASK-ID"):
        assert_raises(RuntimeError, lambda: collection.cmd_purge(types.SimpleNamespace(task_dir=str(ns.task_dir), allow_early=False)), "任务 ID 不匹配时必须取消删除")
        assert ns.task_dir.is_dir()
    with fake_tty(input_func=lambda prompt="": ns.task_id):
        purged = collection.cmd_purge(types.SimpleNamespace(task_dir=str(ns.task_dir), allow_early=False))
    assert not ns.task_dir.exists() and Path(purged["summary"]).is_file()


def test_non_tty_sensitive_operations_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    with non_tty():
        for action in (
            lambda: collection.cmd_create(types.SimpleNamespace(roster=str(ns.roster), config=str(ns.config_path), out=str(ns.tmp / "blocked"))),
            lambda: collection.cmd_review(types.SimpleNamespace(task_dir=str(ns.task_dir))),
            lambda: collection.cmd_reveal(types.SimpleNamespace(task_dir=str(ns.task_dir), invite_id=ns.invite_id)),
            lambda: collection.cmd_purge(types.SimpleNamespace(task_dir=str(ns.task_dir), allow_early=True)),
            lambda: collection.cmd_export(types.SimpleNamespace(task_dir=str(ns.task_dir), out=str(ns.tmp / "blocked.yintian-task"))),
            lambda: collection.cmd_import(types.SimpleNamespace(package=str(ns.tmp / "missing.yintian-task"), out=str(ns.tmp / "blocked-import"))),
        ):
            assert_raises(RuntimeError, action, "非 TTY 环境必须拒绝敏感操作")


HANDOFF_PASSWORD = "HandoffChannelSecret-2026"


def make_encrypted_package(ns: types.SimpleNamespace, password: str = HANDOFF_PASSWORD) -> Path:
    package = ns.tmp / "task-v3.yintian-task"
    result = collection.export_task(ns.task_dir, package, handoff_password=password)
    assert result["handoff_password"] == password
    return package


def read_envelope(package: Path) -> dict:
    return json.loads(package.read_text(encoding="utf-8"))


def test_export_produces_encrypted_v3_envelope(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_encrypted_package(ns)
    envelope = read_envelope(package)
    assert envelope["format"] == collection.ENCRYPTED_TASK_PACKAGE_VERSION == "yintian-task/3"
    assert envelope["task_id"] == ns.task_id
    assert envelope["kdf"]["name"] == "scrypt" and envelope["kdf"]["n"] == 32768 and envelope["kdf"]["r"] == 8 and envelope["kdf"]["p"] == 1
    assert len(base64.b64decode(envelope["kdf"]["salt"], validate=True)) == 16
    assert envelope["cipher"] == "AES-256-GCM" and len(base64.b64decode(envelope["nonce"], validate=True)) == 12
    assert not zipfile.is_zipfile(package)  # 不再是明文 ZIP
    plaintext_markers = collection.build_task_package(ns.task_dir)[0]
    assert plaintext_markers[:4] == b"PK\x03\x04" and plaintext_markers not in package.read_bytes()


def test_export_import_v3_roundtrip(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    collection.review_task(ns.task_dir, PASSWORD)
    package = make_encrypted_package(ns)
    imported = collection.import_task(package, ns.tmp / "v3-import", handoff_password=HANDOFF_PASSWORD)
    status = collection.status_task(imported["task_dir"])
    assert status["total"] == 1 and status["task_status"] == "active"
    reviewed = collection.review_task(imported["task_dir"], PASSWORD)  # 任务密码在新环境中仍可复核
    assert reviewed["needs_review"] + reviewed["verified"] + reviewed["invalid"] == 0  # 已复核过，无待处理提交
    rows = collection.report_rows(Path(imported["task_dir"]))
    assert rows[0]["status"] == "needs_review" and rows[0]["submission_version"] == 1


def test_import_v3_wrong_password_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_encrypted_package(ns)
    target_parent = ns.tmp / "wrong-pass-import"
    with fake_tty(getpass_func=forbidden_prompt):
        assert_raises(ValueError, lambda: collection.import_task(package, target_parent, handoff_password="wrong-password"), "错误交接密码必须拒绝导入")
    assert not target_parent.exists()  # 拒绝时不产生任何落盘


def test_import_v3_tampered_ciphertext_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_encrypted_package(ns)
    envelope = read_envelope(package)
    raw = bytearray(base64.b64decode(envelope["ciphertext"]))
    raw[-1] ^= 1
    envelope["ciphertext"] = base64.b64encode(raw).decode()
    tampered = ns.tmp / "tampered.yintian-task"
    tampered.write_text(json.dumps(envelope), encoding="utf-8")
    target_parent = ns.tmp / "tampered-import"
    assert_raises(ValueError, lambda: collection.import_task(tampered, target_parent, handoff_password=HANDOFF_PASSWORD), "篡改密文必须被 GCM 认证拒绝")
    assert not target_parent.exists()


def test_import_plaintext_v2_package_still_works(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_package(ns)
    captured = io.StringIO()
    with contextlib.redirect_stderr(captured):
        imported = collection.import_task(package, ns.tmp / "v2-import")
    assert captured.getvalue().strip()  # 明文包必须有显著警告（不断言具体文案）
    assert collection.status_task(imported["task_dir"])["total"] == 1


def test_private_key_is_yintian_key_envelope(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    blob = (ns.task_dir / "private.pem.enc").read_bytes()
    assert not blob.startswith(b"-----BEGIN")
    envelope = json.loads(blob)
    assert envelope["format"] == "yintian-key/1" and envelope["cipher"] == "AES-256-GCM"
    assert envelope["kdf"]["name"] == "scrypt" and envelope["kdf"]["n"] == 32768


def test_wrong_password_review_fails(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    assert_raises(RuntimeError, lambda: collection.review_task(ns.task_dir, "wrong-password"), "错误任务密码必须导致 review 失败")


def test_legacy_pem_private_key_still_reviewable(tmp_path: Path) -> None:
    from cryptography.hazmat.primitives import serialization

    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    private_key = collection.unlock_private_key(ns.task_dir, PASSWORD)
    legacy_pem = private_key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.BestAvailableEncryption(PASSWORD.encode()))
    (ns.task_dir / "private.pem.enc").write_bytes(legacy_pem)  # 模拟存量任务的旧格式私钥文件
    reviewed = collection.review_task(ns.task_dir, PASSWORD)
    assert reviewed["needs_review"] == 1  # 旧格式私钥仍能解密复核


def test_ingest_versions_per_invite_limited(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns, "first.yintian")
    collection.review_task(ns.task_dir, PASSWORD)
    for index in range(2, collection.MAX_VERSIONS_PER_INVITE + 2):
        submit(ns, f"extra-{index:02d}.yintian")
    result = collection.ingest_task(ns.task_dir, ns.incoming)
    assert result["accepted"] == collection.MAX_VERSIONS_PER_INVITE - 1 and result["rejected"] == 1
    row = collection.report_rows(ns.task_dir)[0]
    assert row["submission_version"] == 1 and row["status"] == "needs_review"  # 已有有效版本不受影响
    with closing_db(ns.task_dir) as db:
        versions = db.execute("SELECT COUNT(*) FROM submissions WHERE invite_id=?", (ns.invite_id,)).fetchone()[0]
    assert versions == collection.MAX_VERSIONS_PER_INVITE
    again = collection.ingest_task(ns.task_dir, ns.incoming)
    assert again["accepted"] == 0 and again["rejected"] == 1  # 超限提交始终拒绝且不覆盖


def closing_db(task_dir: Path):
    return contextlib.closing(collection.connect_db(task_dir))


def test_positive_int_env_fallback() -> None:
    key = "YINTIAN_MAX_VERSIONS_PER_INVITE"
    original = os.environ.get(key)
    try:
        os.environ[key] = "abc"
        assert collection.positive_int_env(key, 10) == 10
        os.environ[key] = "-3"
        assert collection.positive_int_env(key, 10) == 10
        os.environ[key] = "5"
        assert collection.positive_int_env(key, 10) == 5
    finally:
        if original is None:
            os.environ.pop(key, None)
        else:
            os.environ[key] = original


def test_inflated_kdf_params_rejected_before_derivation(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_encrypted_package(ns)  # 在禁用派生前先造好合法 v3 包
    key_envelope = json.loads((ns.task_dir / "private.pem.enc").read_bytes())
    key_envelope["kdf"]["n"] = 2**20
    package_envelope = read_envelope(package)
    package_envelope["kdf"]["r"] = 64
    evil_package = ns.tmp / "evil-kdf.yintian-task"
    evil_package.write_text(json.dumps(package_envelope), encoding="utf-8")
    real_derive = collection.derive_scrypt_key

    def forbidden_derive(*args, **kwargs):
        raise AssertionError("畸形 KDF 参数不得在派生密钥前放行")

    collection.derive_scrypt_key = forbidden_derive
    try:
        assert_raises(ValueError, lambda: collection.decrypt_private_key(key_envelope, PASSWORD), "私钥信封的恶意 KDF 参数必须在派生前拒绝")
        assert_raises(ValueError, lambda: collection.import_task(evil_package, ns.tmp / "evil-kdf-import", handoff_password=HANDOFF_PASSWORD), "v3 交接包的恶意 KDF 参数必须在派生前拒绝")
        assert not (ns.tmp / "evil-kdf-import").exists()
    finally:
        collection.derive_scrypt_key = real_derive
    valid_salt = base64.b64encode(b"0" * 16).decode()
    salt, n, r, p = collection.validate_kdf_params({"name": "scrypt", "salt": valid_salt, "n": 32768, "r": 8, "p": 1})
    assert (len(salt), n, r, p) == (16, 32768, 8, 1)  # 合法包参数不受影响
    for bad in ({"n": 2**17, "r": 8, "p": 1}, {"n": 2**9, "r": 8, "p": 1}, {"n": 2**15, "r": 9, "p": 1}, {"n": 2**15, "r": 8, "p": 3}):
        kdf = {"name": "scrypt", "salt": valid_salt, **bad}
        assert_raises(ValueError, lambda: collection.validate_kdf_params(kdf), f"KDF 参数越界必须拒绝: {bad}")


class NonClosingStringIO(io.StringIO):
    def close(self) -> None:
        pass


def test_one_time_password_written_to_control_terminal(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    real_opener = collection.open_control_terminal
    try:
        terminal = NonClosingStringIO()
        collection.open_control_terminal = lambda: terminal
        capture = TtyCapture()
        with fake_tty(capture=capture):
            collection.cmd_create(types.SimpleNamespace(roster=str(ns.roster), config=str(ns.config_path), out=str(ns.tmp / "tty-task")))
        shown = terminal.getvalue()
        assert capture.getvalue() == "" and "任务密码" in shown
        password_line = [line for line in shown.splitlines() if line and not line.startswith(("任务密码", "请通过"))][0]
        assert password_line not in capture.getvalue()

        terminal2 = NonClosingStringIO()
        collection.open_control_terminal = lambda: terminal2
        capture2 = TtyCapture()
        with fake_tty(capture=capture2):
            collection.cmd_export(types.SimpleNamespace(task_dir=str(ns.task_dir), out=str(ns.tmp / "tty.yintian-task")))
        assert capture2.getvalue() == "" and "交接密码" in terminal2.getvalue()

        collection.open_control_terminal = lambda: None  # 控制终端不可用时回退 stderr，绝不回退 stdout
        capture3 = TtyCapture()
        err = io.StringIO()
        with fake_tty(capture=capture3), contextlib.redirect_stderr(err):
            collection.cmd_export(types.SimpleNamespace(task_dir=str(ns.task_dir), out=str(ns.tmp / "tty-fallback.yintian-task")))
        assert capture3.getvalue() == "" and "交接密码" in err.getvalue()
    finally:
        collection.open_control_terminal = real_opener


def test_v3_envelope_size_limit_and_whitespace_sniffing(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_encrypted_package(ns)
    padded = ns.tmp / "padded.yintian-task"
    padded.write_text(" " * 5000 + package.read_text(encoding="utf-8"), encoding="utf-8")
    imported = collection.import_task(padded, ns.tmp / "padded-import", handoff_password=HANDOFF_PASSWORD)
    assert collection.status_task(imported["task_dir"])["total"] == 1  # >4KB 前导空白不再误判为明文包
    real_limit = collection.MAX_V3_ENVELOPE_BYTES
    collection.MAX_V3_ENVELOPE_BYTES = 1024
    try:
        assert_raises(ValueError, lambda: collection.import_task(package, ns.tmp / "oversize-import", handoff_password=HANDOFF_PASSWORD), "超限 v3 信封必须在完整读取前拒绝")
        assert not (ns.tmp / "oversize-import").exists()
    finally:
        collection.MAX_V3_ENVELOPE_BYTES = real_limit


def test_compare_ocr_front_back_word_boundary(tmp_path: Path) -> None:
    for field_id, expected in (("front", True), ("back", True), ("id_front", True), ("id_card_back", True), ("front_side", True),
                               ("backdrop", False), ("feedback_scan", False), ("upfront", False), ("back2", False)):
        assert collection.is_identity_attachment_field(field_id) is expected, field_id
    original_ocr = collection.ocr_attachment
    calls = []
    collection.ocr_attachment = lambda item: (calls.append(item["field_id"]), ("", ["ocr_no_text"]))[1]
    try:
        fields = [
            {"id": "name", "label": "姓名", "type": "text"},
            {"id": "backdrop", "label": "背景图", "type": "image_attachment"},
            {"id": "feedback_scan", "label": "反馈扫描件", "type": "image_attachment"},
            {"id": "id_front", "label": "证件正面", "type": "image_attachment"},
            {"id": "id_card_back", "label": "证件背面", "type": "image_attachment"},
        ]
        attachments = [{"field_id": field_id} for field_id in ("backdrop", "feedback_scan", "id_front", "id_card_back")]
        collection.compare_ocr({"values": {"name": "张三"}}, attachments, fields)
        assert sorted(calls) == ["id_card_back", "id_front"]  # backdrop 类 id 不触发证件 OCR 比对
    finally:
        collection.ocr_attachment = original_ocr


def test_unlock_rejects_unencrypted_or_unknown_key_blob(tmp_path: Path) -> None:
    from cryptography.hazmat.primitives import serialization

    ns = make_scenario(tmp_path)
    private_key = collection.unlock_private_key(ns.task_dir, PASSWORD)
    plain_pem = private_key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption())
    original = (ns.task_dir / "private.pem.enc").read_bytes()
    try:
        (ns.task_dir / "private.pem.enc").write_bytes(plain_pem)
        assert_raises(ValueError, lambda: collection.unlock_private_key(ns.task_dir, PASSWORD), "未加密 PKCS8 PEM 必须明确拒绝")
        (ns.task_dir / "private.pem.enc").write_bytes(b"random garbage")
        assert_raises(ValueError, lambda: collection.unlock_private_key(ns.task_dir, PASSWORD), "无法识别的私钥数据必须明确拒绝")
    finally:
        (ns.task_dir / "private.pem.enc").write_bytes(original)


def test_import_plaintext_package_requires_confirmation(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    package = make_package(ns)
    blocked_parent = ns.tmp / "confirm-blocked"
    with fake_tty(input_func=lambda prompt="": "WRONG-TASK-ID"):
        assert_raises(RuntimeError, lambda: collection.cmd_import(types.SimpleNamespace(package=str(package), out=str(blocked_parent))), "确认不符必须拒绝导入明文包")
    assert not (blocked_parent / ns.task_id).exists()
    with fake_tty(input_func=lambda prompt="": ns.task_id):
        imported = collection.cmd_import(types.SimpleNamespace(package=str(package), out=str(ns.tmp / "confirm-ok")))
    assert collection.status_task(imported["task_dir"])["total"] == 1


def test_imported_task_dir_permissions(tmp_path: Path) -> None:
    if os.name == "nt":
        return
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    collection.review_task(ns.task_dir, PASSWORD)
    collection.write_reports(ns.task_dir, ["json"])
    package = make_package(ns)
    imported = collection.import_task(package, ns.tmp / "perm-import")
    root = Path(imported["task_dir"])
    assert stat.S_IMODE(root.stat().st_mode) == 0o700
    for sub in ("invites", "submissions", "reports", f"submissions/{ns.invite_id}"):
        assert stat.S_IMODE((root / sub).stat().st_mode) == 0o700, sub


def make_group_scenario(tmp: Path) -> types.SimpleNamespace:
    """创建 group 模式任务（无令牌、单份 FORM.yintian-form、E001/E002 两名员工），返回常用路径与常量。"""
    config = collection.default_config()
    config["purpose"] = "为依法办理员工商业保险收集必要身份资料"
    config["contact"] = "人事部王老师，内线 8001"
    config["correction"] = "在截止日前联系人事部王老师撤回原提交并重新提交"
    config["deadline"] = "2098-12-31T23:59:59Z"
    config["retention_until"] = "2099-12-31T23:59:59Z"
    config["fields"] = [
        {"id": "name", "label": "姓名", "type": "text", "required": True, "sensitive": False},
        {"id": "employee_id", "label": "工号", "type": "text", "required": True, "sensitive": False},
        {"id": "phone", "label": "手机号", "type": "phone_cn", "required": True, "sensitive": True},
        {"id": "id_number", "label": "身份证号", "type": "cn_id", "required": True, "sensitive": True},
        {"id": "address", "label": "住址", "type": "address", "required": True, "sensitive": True},
    ]
    config_path = tmp / "group-config.json"
    collection.dump_json(config_path, config)
    roster = tmp / "group-roster.csv"
    with roster.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream); writer.writerow(["employee_id", "name"]); writer.writerow(["E001", "张三"]); writer.writerow(["E002", "李四"])
    warning = io.StringIO()
    with contextlib.redirect_stderr(warning):
        created = collection.create_task(roster, config_path, tmp / "group-tasks", PASSWORD, require_terminal=False, mode="group")
    assert "冒名" in warning.getvalue()  # group 创建必须打印可冒名风险提示
    task_dir = Path(created["task_dir"])
    values = {"name": "张三", "employee_id": "E001", "phone": "13800138000", "id_number": VALID_ID, "address": "北京市朝阳区"}
    incoming = tmp / "group-incoming"; incoming.mkdir()
    return types.SimpleNamespace(tmp=tmp, config_path=config_path, roster=roster, created=created, task_dir=task_dir, task_id=created["task_id"], values=values, incoming=incoming)


def make_group_envelope(task_dir: Path, employee_id: str, values: dict, overrides: dict | None = None) -> dict:
    """构造 group 模式提交：invite_id=GRP-<employee_id>，载荷不含 invite_token。"""
    task = collection.load_json(task_dir / "task.json")
    invite_id = collection.group_invite_id(employee_id)
    payload = {
        "format_version": task["format_version"],
        "task_id": task["task_id"],
        "invite_id": invite_id,
        "schema_hash": task["schema_hash"],
        "notice_hash": task["notice_hash"],
        "template_version": task["template_version"],
        "submitted_at": collection.now_iso(),
        "consent_confirmed": True,
        "values": values,
        "attachments": {},
    }
    if overrides:
        payload.update(overrides)
    return wrap_payload(task_dir, invite_id, payload)


def submit_group(ns: types.SimpleNamespace, employee_id: str, filename: str, values: dict | None = None, **kwargs) -> None:
    collection.dump_json(ns.incoming / filename, make_group_envelope(ns.task_dir, employee_id, values if values is not None else ns.values, **kwargs))


def test_directed_invite_form_matches_html_config(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    form = collection.load_json(ns.task_dir / "invites" / f"{ns.invite_id}.yintian-form")
    assert form["format"] == collection.FORM_FORMAT_VERSION == "yintian-form/1"
    assert form["mode"] == "directed"
    html = (ns.task_dir / "invites" / f"{ns.invite_id}.html").read_text(encoding="utf-8")
    marker = '<script id="cfg" type="application/json">'
    embedded = json.loads(html.split(marker, 1)[1].split("</script>", 1)[0])
    for key in ("invite_id", "invite_token", "key_id", "schema_hash", "notice_hash", "task_id", "public_key_pem"):
        assert form[key] == embedded[key], key  # .yintian-form 与 HTML 内嵌配置是同一份数据
    assert form["invite_id"] == ns.invite_id and form["invite_token"] == invite_token(ns.task_dir, ns.invite_id)
    assert "yintian-form/1 的人类可读渲染版" in html  # 真伪核对区已标注渲染关系


def test_group_create_layout(tmp_path: Path) -> None:
    ns = make_group_scenario(tmp_path)
    task = collection.load_json(ns.task_dir / "task.json")
    assert task["mode"] == "group"
    form = collection.load_json(ns.task_dir / "FORM.yintian-form")
    assert form["format"] == "yintian-form/1" and form["mode"] == "group"
    assert "invite_id" not in form and "invite_token" not in form and "name" not in form  # 单份群发，无任何个人标识与令牌
    assert form["key_id"] == task["key_id"] and form["public_key_pem"] == (ns.task_dir / "public.pem").read_text(encoding="utf-8")
    assert [field["id"] for field in form["fields"]] == [field["id"] for field in task["fields"]]
    assert not list((ns.task_dir / "invites").glob("*"))  # group 模式不生成个人邀请文件
    index_rows = list(csv.DictReader((ns.task_dir / "invite-index.csv").open(encoding="utf-8-sig")))
    assert [row["invite_id"] for row in index_rows] == ["GRP-E001", "GRP-E002"]
    with closing_db(ns.task_dir) as db:
        stored = [row[0] for row in db.execute("SELECT invite_id FROM invites ORDER BY employee_id")]
    assert stored == ["GRP-E001", "GRP-E002"]  # 名单照常入库，invite_id 采用 GRP- 约定


def test_group_config_requires_employee_id_field(tmp_path: Path) -> None:
    config = collection.default_config()
    config["purpose"] = "为依法办理员工商业保险收集必要身份资料"
    config["contact"] = "人事部王老师，内线 8001"
    config["correction"] = "在截止日前联系人事部王老师撤回原提交并重新提交"
    config["deadline"] = "2098-12-31T23:59:59Z"
    config["retention_until"] = "2099-12-31T23:59:59Z"
    config_path = tmp_path / "config.json"
    collection.dump_json(config_path, config)
    roster = tmp_path / "roster.csv"
    roster.write_text("employee_id,name\nE001,张三\n", encoding="utf-8")
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        assert_raises(ValueError, lambda: collection.create_task(roster, config_path, tmp_path / "blocked", PASSWORD, require_terminal=False, mode="group"), "group 模式缺少 employee_id 字段必须拒绝创建")
    assert "冒名" in err.getvalue()  # 拒绝前也已如实提示风险
    assert not (tmp_path / "blocked").exists()
    created = collection.create_task(roster, config_path, tmp_path / "ok", PASSWORD, require_terminal=False)  # 默认 directed 不受影响
    assert collection.load_json(Path(created["task_dir"]) / "task.json")["mode"] == "directed"


def test_group_ingest_review_and_export_clear(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    ns = make_group_scenario(tmp_path)
    submit_group(ns, "E001", "one.yintian")
    assert collection.ingest_task(ns.task_dir, ns.incoming)["accepted"] == 1
    assert collection.review_task(ns.task_dir, PASSWORD) == {"verified": 1, "needs_review": 0, "invalid": 0}
    rows = collection.report_rows(ns.task_dir)
    assert rows[0]["status"] == "verified" and rows[0]["submission_version"] == 1
    out = ns.tmp / "group-result.xlsx"
    args = types.SimpleNamespace(task_dir=str(ns.task_dir), out=str(out), fields="phone,id_number", mask=["id_number=last4"], purpose="保险办理", recipient="保险公司对接人", formats=["xlsx"])
    with fake_tty(input_func=lambda prompt="": ns.task_id, getpass_func=lambda prompt="": PASSWORD):
        result = collection.cmd_export_clear(args)
    assert result["rows"] == 1
    sheet = load_workbook(result["xlsx"]).active
    assert [cell.value for cell in sheet[1]] == ["employee_id", "name", "phone", "id_number"]
    assert [sheet.cell(2, index).value for index in range(1, 5)] == ["E001", "张三", "13800138000", "**************1230"]


def test_group_identity_mismatch_flagged(tmp_path: Path) -> None:
    ns = make_group_scenario(tmp_path)
    submit_group(ns, "E001", "01-wrong-emp.yintian", values=dict(ns.values, employee_id="E002"))  # 冒用他人工号
    submit_group(ns, "E001", "02-wrong-name.yintian", values=dict(ns.values, name="李四"))  # 错姓名
    assert collection.ingest_task(ns.task_dir, ns.incoming)["accepted"] == 2
    assert collection.review_task(ns.task_dir, PASSWORD)["needs_review"] == 2
    with closing_db(ns.task_dir) as db:
        conflicts = [json.loads(row[0]) for row in db.execute("SELECT conflict_fields FROM submissions WHERE invite_id='GRP-E001' ORDER BY version")]
    assert "employee_id_roster_mismatch" in conflicts[0] and "name_roster_mismatch" not in conflicts[0]
    assert "name_roster_mismatch" in conflicts[1] and "employee_id_roster_mismatch" not in conflicts[1]


def test_group_version_cap_per_employee(tmp_path: Path) -> None:
    ns = make_group_scenario(tmp_path)
    for index in range(collection.MAX_VERSIONS_PER_INVITE):
        submit_group(ns, "E001", f"e001-{index:02d}.yintian")
    assert collection.ingest_task(ns.task_dir, ns.incoming)["accepted"] == collection.MAX_VERSIONS_PER_INVITE
    submit_group(ns, "E001", "e001-overflow.yintian")
    overflow = collection.ingest_task(ns.task_dir, ns.incoming)
    assert overflow["accepted"] == 0 and overflow["rejected"] == 1  # 超限拒绝且不覆盖已有版本
    other_values = {"name": "李四", "employee_id": "E002", "phone": "13900139000", "id_number": VALID_ID, "address": "上海市黄浦区"}
    submit_group(ns, "E002", "e002-01.yintian", values=other_values)
    assert collection.ingest_task(ns.task_dir, ns.incoming)["accepted"] == 1  # 版本上限按工号维度计数，他人不受影响
    with closing_db(ns.task_dir) as db:
        assert db.execute("SELECT COUNT(*) FROM submissions WHERE invite_id='GRP-E001'").fetchone()[0] == collection.MAX_VERSIONS_PER_INVITE
        assert db.execute("SELECT COUNT(*) FROM submissions WHERE invite_id='GRP-E002'").fetchone()[0] == 1


def test_envelope_mode_isolation(tmp_path: Path) -> None:
    ns = make_group_scenario(tmp_path)
    directed_dir = tmp_path / "directed"; directed_dir.mkdir()
    directed_ns = make_scenario(directed_dir)
    assert collection.task_mode({}) == "directed"  # 旧任务没有 mode 键，默认 directed
    collection.dump_json(ns.incoming / "inv-style.yintian", make_envelope(directed_ns.task_dir, directed_ns.invite_id, directed_ns.values))
    assert collection.ingest_task(ns.task_dir, ns.incoming)["rejected"] == 1  # group 任务只接受 GRP- 标识
    collection.dump_json(directed_ns.incoming / "grp-style.yintian", make_group_envelope(directed_ns.task_dir, "E001", directed_ns.values))
    assert collection.ingest_task(directed_ns.task_dir, directed_ns.incoming)["rejected"] == 1  # directed 任务只接受 INV- 标识
    submit(directed_ns, "no-token.yintian", overrides={"invite_token": None})
    assert collection.ingest_task(directed_ns.task_dir, directed_ns.incoming)["accepted"] == 1
    assert collection.review_task(directed_ns.task_dir, PASSWORD)["invalid"] == 1  # directed 模式令牌缺失仍判 invalid，校验链未松动


def export_clear_args(ns: types.SimpleNamespace, out: Path, **overrides) -> types.SimpleNamespace:
    base = dict(task_dir=str(ns.task_dir), out=str(out), fields="phone,id_number,address", mask=["id_number=last4", "phone=mid4"],
                purpose="办理员工商业保险", recipient="保险公司对接人", formats=["xlsx", "json"])
    base.update(overrides)
    return types.SimpleNamespace(**base)


def test_export_clear_end_to_end(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    assert collection.review_task(ns.task_dir, PASSWORD)["needs_review"] == 1  # 缺证件附件；needs_review 仍属 current+valid
    capture = TtyCapture()
    warning = io.StringIO()
    with fake_tty(input_func=lambda prompt="": ns.task_id, getpass_func=lambda prompt="": PASSWORD, capture=capture), contextlib.redirect_stderr(warning):
        result = collection.cmd_export_clear(export_clear_args(ns, ns.tmp / "result.xlsx"))
    assert result["rows"] == 1 and set(result) >= {"xlsx", "json"}
    shown = capture.getvalue()
    for expected in (ns.task_id, "phone", "id_number=last4", "phone=mid4", "办理员工商业保险", "保险公司对接人", "数据行数: 1"):
        assert expected in shown, expected  # 导出前完整回显任务/字段/脱敏/行数/用途/接收方
    assert "系统无法管控后续传播" in warning.getvalue()
    sheet = load_workbook(result["xlsx"]).active
    assert [cell.value for cell in sheet[1]] == ["employee_id", "name", "phone", "id_number", "address"]
    row = [sheet.cell(2, index).value for index in range(1, 6)]
    assert row == ["=1+1", "张三", "138****8000", "**************1230", "北京市朝阳区"]
    assert all(sheet.cell(2, index).data_type == "s" for index in range(1, 6))  # 字符串单元格防公式注入
    document = collection.load_json(Path(result["json"]))
    assert document["rows"] == [dict(zip(["employee_id", "name", "phone", "id_number", "address"], row))]
    assert document["purpose"] == "办理员工商业保险" and document["recipient"] == "保险公司对接人"
    json_text = Path(result["json"]).read_text(encoding="utf-8")
    assert VALID_ID not in json_text and "13800138000" not in json_text and "北京市朝阳区" in json_text  # 脱敏生效、未声明字段明文
    assert "effective_date" not in json_text  # 白名单外字段不导出
    if os.name != "nt":
        assert stat.S_IMODE(Path(result["xlsx"]).stat().st_mode) == 0o600
        assert stat.S_IMODE(Path(result["json"]).stat().st_mode) == 0o600


def test_export_clear_non_tty_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    collection.review_task(ns.task_dir, PASSWORD)
    with non_tty():
        assert_raises(RuntimeError, lambda: collection.cmd_export_clear(export_clear_args(ns, ns.tmp / "blocked.xlsx")), "非 TTY 环境必须拒绝明文导出")
    assert not (ns.tmp / "blocked.xlsx").exists()


def test_export_clear_wrong_password_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    collection.review_task(ns.task_dir, PASSWORD)
    with fake_tty(input_func=forbidden_prompt, getpass_func=lambda prompt="": "wrong-password"):
        assert_raises(RuntimeError, lambda: collection.cmd_export_clear(export_clear_args(ns, ns.tmp / "wrong-pass.xlsx")), "错误任务密码必须拒绝明文导出")
    assert not (ns.tmp / "wrong-pass.xlsx").exists()


def test_export_clear_expired_rejected(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    collection.review_task(ns.task_dir, PASSWORD)
    expire_task(ns.task_dir)
    with fake_tty(input_func=forbidden_prompt, getpass_func=forbidden_prompt):
        assert_raises(RuntimeError, lambda: collection.cmd_export_clear(export_clear_args(ns, ns.tmp / "expired.xlsx")), "过期任务必须拒绝明文导出且不出现交互提示")
    assert not (ns.tmp / "expired.xlsx").exists()


def test_export_clear_confirm_mismatch_cancels(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    submit_accepted(ns)
    collection.review_task(ns.task_dir, PASSWORD)
    with fake_tty(input_func=lambda prompt="": "YT-00000000-WRONG", getpass_func=lambda prompt="": PASSWORD):
        assert_raises(RuntimeError, lambda: collection.cmd_export_clear(export_clear_args(ns, ns.tmp / "cancelled.xlsx")), "任务 ID 不匹配必须取消导出")
    assert not (ns.tmp / "cancelled.xlsx").exists() and not (ns.tmp / "cancelled.json").exists()


def test_export_clear_fields_and_mask_validation(tmp_path: Path) -> None:
    ns = make_scenario(tmp_path)
    parser = collection.build_parser()
    assert_raises(SystemExit, lambda: parser.parse_args(["export-clear", str(ns.task_dir), "--out", "x.xlsx"]), "缺 --fields 必须拒绝")
    task = collection.load_task(ns.task_dir)[1]
    assert_raises(ValueError, lambda: collection.parse_export_fields(task, ""), "空白字段名单必须拒绝")
    assert_raises(ValueError, lambda: collection.parse_export_fields(task, "no_such_field"), "未知字段必须拒绝")
    assert_raises(ValueError, lambda: collection.parse_export_fields(task, "id_front"), "附件字段必须拒绝")
    assert_raises(ValueError, lambda: collection.parse_export_fields(task, "name"), "只含默认附带字段的名单必须拒绝")
    assert collection.parse_export_fields(task, "phone,name,employee_id") == ["phone"]  # 身份列默认附带，重复声明被归并
    assert_raises(ValueError, lambda: collection.parse_mask_specs(["phone=bad"]), "未知脱敏规则必须拒绝")
    assert_raises(ValueError, lambda: collection.parse_mask_specs(["phone"]), "缺规则的脱敏声明必须拒绝")
    assert collection.mask_value("last4", VALID_ID) == "**************1230"
    assert collection.mask_value("mid4", "13800138000") == "138****8000"
    assert collection.mask_value("last4", "123") == "***" and collection.mask_value("mid4", "1234567") == "*******"
    with fake_tty(input_func=forbidden_prompt, getpass_func=forbidden_prompt):
        assert_raises(ValueError, lambda: collection.cmd_export_clear(export_clear_args(ns, ns.tmp / "m.xlsx", mask=["effective_date=last4"])), "脱敏指向未导出字段必须在交互前拒绝")


def main() -> None:
    tests = [(name, fn) for name, fn in sorted(globals().items()) if name.startswith("test_") and callable(fn)]
    for name, fn in tests:
        if "tmp_path" in inspect.signature(fn).parameters:
            with tempfile.TemporaryDirectory() as tmp:
                fn(Path(tmp))
        else:
            fn()
    print(f"PASS: encrypted collection workflow ({len(tests)} checks)")


if __name__ == "__main__":
    main()
