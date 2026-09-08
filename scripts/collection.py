"""隐填 · 端到端加密的私密信息收集管理 CLI。"""
from __future__ import annotations

import argparse
import base64
import csv
import getpass
import hashlib
import io
import json
import os
import re
import secrets
import shutil
import sqlite3
import string
import sys
import tempfile
import zipfile
from contextlib import closing, contextmanager
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any

FORMAT_VERSION = "yintian-submission/2"
LEGACY_FORMAT_VERSION = "yintian-submission/1"
LEGACY_TASK_PACKAGE_VERSION = "yintian-task/1"
TASK_PACKAGE_VERSION = "yintian-task/2"
ENCRYPTED_TASK_PACKAGE_VERSION = "yintian-task/3"
KEY_ENVELOPE_VERSION = "yintian-key/1"
FORM_FORMAT_VERSION = "yintian-form/1"
GROUP_INVITE_PREFIX = "GRP-"
TASK_MODES = {"directed", "group"}
MASK_RULES = {"last4", "mid4"}
LEGACY_KEY_PEM_PREFIX = b"-----BEGIN ENCRYPTED PRIVATE KEY-----"
SCRYPT_N = 2**15
SCRYPT_R = 8
SCRYPT_P = 1
ALLOWED_TYPES = {"text", "phone_cn", "cn_id", "date", "address", "single_choice", "image_attachment", "pdf_attachment"}
ATTACHMENT_TYPES = {"image_attachment", "pdf_attachment"}
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TOTAL_BYTES = 15 * 1024 * 1024
MAX_PDF_PAGES = 20
PDF_RENDER_SCALE = 2
MAX_ENVELOPE_BYTES = 32 * 1024 * 1024
MAX_V3_ENVELOPE_BYTES = 200 * 1024 * 1024
V3_SNIFF_BYTES = 1024 * 1024
MAX_PACKAGE_FILES = 20_000
MAX_PACKAGE_BYTES = 2 * 1024 * 1024 * 1024
MAX_PACKAGE_MEMBER_BYTES = 128 * 1024 * 1024
MAX_IMAGE_PIXELS = int(os.getenv("YINTIAN_MAX_IMAGE_PIXELS", "40000000"))


def positive_int_env(name: str, default: int) -> int:
    try:
        value = int(os.getenv(name, ""))
    except ValueError:
        return default
    return value if value > 0 else default


MAX_VERSIONS_PER_INVITE = positive_int_env("YINTIAN_MAX_VERSIONS_PER_INVITE", 10)
TASK_ID_RE = re.compile(r"^YT-[0-9]{8}-[A-Z0-9]{6}$")
INVITE_ID_RE = re.compile(r"^INV-[A-Z0-9]{10}$")
FIELD_ID_RE = re.compile(r"^[a-z][a-z0-9_]{1,63}$")


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def parse_time(value: str) -> datetime:
    value = value.strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        value += "T23:59:59+00:00"
    elif value.endswith("Z"):
        value = value[:-1] + "+00:00"
    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def normalize_submitted_at(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    try:
        parsed = parse_time(value)
    except (TypeError, ValueError):
        return None
    return parsed.replace(microsecond=0).isoformat().replace("+00:00", "Z")


def valid_date(value: str) -> bool:
    try:
        return datetime.strptime(value, "%Y-%m-%d").strftime("%Y-%m-%d") == value
    except (TypeError, ValueError):
        return False


def canonical(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def atomic_write(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, delete=False) as stream:
            temp = Path(stream.name)
            stream.write(data)
        temp.replace(path)
    except BaseException:
        if temp is not None:
            try:
                temp.unlink()
            except FileNotFoundError:
                pass
        raise


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def dump_json(path: Path, value: Any) -> None:
    atomic_write(path, json.dumps(value, ensure_ascii=False, indent=2).encode("utf-8") + b"\n")


def random_id(prefix: str, length: int) -> str:
    alphabet = string.ascii_uppercase + string.digits
    return prefix + "".join(secrets.choice(alphabet) for _ in range(length))


def generate_password(length: int = 28) -> str:
    alphabet = string.ascii_letters + string.digits + "-_.~"
    while True:
        value = "".join(secrets.choice(alphabet) for _ in range(length))
        if any(c.islower() for c in value) and any(c.isupper() for c in value) and any(c.isdigit() for c in value):
            return value


def open_control_terminal():
    """打开控制终端用于显示一次性密码（POSIX 用 /dev/tty，Windows 用 CON）；打不开时返回 None。"""
    try:
        return open("CON" if os.name == "nt" else "/dev/tty", "w", encoding="utf-8", errors="replace")
    except OSError:
        return None


def print_once_secret(heading: str, secret: str, footnote: str) -> None:
    """一次性密码只写控制终端；控制终端不可用时回退 stderr，绝不写可能被管道采集的 stdout。"""
    stream = open_control_terminal()
    try:
        target = stream if stream is not None else sys.stderr
        print(f"\n{heading}", file=target)
        print(secret, file=target)
        print(footnote + "\n", file=target)
    finally:
        if stream is not None:
            stream.close()


def require_tty(action: str) -> None:
    if not (sys.stdin.isatty() and sys.stdout.isatty()):
        raise RuntimeError(f"{action} 必须在交互终端中运行；TTY 检查不能证明该终端未被 Agent 控制。")


def task_expired(task: dict[str, Any]) -> bool:
    return datetime.now(timezone.utc) > parse_time(task["retention_until"])


def task_late(task: dict[str, Any], received_at: str) -> bool:
    return parse_time(received_at) > parse_time(task["deadline"])


def load_task(task_dir: str | Path) -> tuple[Path, dict[str, Any]]:
    root = Path(task_dir).expanduser().resolve()
    task_path = root / "task.json"
    if not task_path.is_file():
        raise FileNotFoundError(f"不是有效任务目录: {root}")
    task = load_json(task_path)
    if not TASK_ID_RE.fullmatch(task.get("task_id", "")):
        raise ValueError("task.json 中的 task_id 无效")
    if task.get("format_version") not in {LEGACY_FORMAT_VERSION, FORMAT_VERSION}:
        raise ValueError("任务格式版本不受支持")
    if task.get("mode", "directed") not in TASK_MODES:
        raise ValueError("task.json 中的 mode 无效")
    required = ("key_id", "title", "purpose", "deadline", "retention_until", "contact", "correction", "template_version", "schema_hash", "notice_hash", "fields")
    if any(key not in task for key in required) or not isinstance(task["fields"], list):
        raise ValueError("task.json 缺少必要配置")
    if parse_time(task["retention_until"]) <= parse_time(task["deadline"]):
        raise ValueError("task.json 保存期限无效")
    if sha256_bytes(canonical(task["fields"])) != task["schema_hash"]:
        raise ValueError("task.json 字段模板哈希不匹配")
    notice = {key: task[key] for key in ("title", "purpose", "deadline", "retention_until", "contact", "correction")}
    if sha256_bytes(canonical(notice)) != task["notice_hash"]:
        raise ValueError("task.json 告知内容哈希不匹配")
    return root, task


def require_current_format(task: dict[str, Any], action: str) -> None:
    if task["format_version"] != FORMAT_VERSION:
        raise RuntimeError(f"旧版任务不支持{action}；请新建 v2 任务继续收集")


def task_mode(task: dict[str, Any]) -> str:
    """任务模式只从本地 task.json 读取（load_task 已校验取值），提交信封与载荷无法伪造。"""
    return task.get("mode", "directed")


def group_invite_id(employee_id: str) -> str:
    return GROUP_INVITE_PREFIX + employee_id


def parse_group_employee_id(invite_id: str) -> str:
    if not invite_id.startswith(GROUP_INVITE_PREFIX):
        raise ValueError("group 提交的 invite_id 必须以 GRP- 开头")
    employee_id = invite_id[len(GROUP_INVITE_PREFIX):]
    if not employee_id or employee_id != employee_id.strip() or len(employee_id) > 128 or any(ord(char) < 0x20 for char in employee_id):
        raise ValueError("group 提交的 invite_id 格式无效")
    return employee_id


@contextmanager
def task_lock(root: Path):
    lock = root / ".write.lock"
    try:
        descriptor = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError as exc:
        raise RuntimeError(f"任务正由另一进程写入；如确认上次异常退出，请删除 {lock}") from exc
    try:
        os.write(descriptor, f"pid={os.getpid()} time={now_iso()}\n".encode())
        os.close(descriptor)
        yield
    finally:
        try:
            lock.unlink()
        except FileNotFoundError:
            pass


def connect_db(root: Path) -> sqlite3.Connection:
    db = sqlite3.connect(root / "state.sqlite3", timeout=5.0)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA busy_timeout=5000")
    db.execute("PRAGMA foreign_keys=ON")
    return db


def init_db(root: Path, rows: list[dict[str, str]]) -> None:
    db_path = root / "state.sqlite3"
    if os.name != "nt":
        os.close(os.open(db_path, os.O_CREAT | os.O_WRONLY, 0o600))
    with closing(connect_db(root)) as db, db:
        db.executescript(
            """
            CREATE TABLE invites (
                invite_id TEXT PRIMARY KEY,
                employee_id TEXT NOT NULL UNIQUE,
                name TEXT NOT NULL,
                token_hash TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'invited',
                current_submission_id INTEGER,
                created_at TEXT NOT NULL
            );
            CREATE TABLE submissions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                invite_id TEXT NOT NULL REFERENCES invites(invite_id),
                version INTEGER NOT NULL,
                sha256 TEXT NOT NULL UNIQUE,
                path TEXT NOT NULL,
                received_at TEXT NOT NULL,
                submitted_at TEXT,
                reviewed_at TEXT,
                late INTEGER NOT NULL DEFAULT 0,
                status TEXT NOT NULL,
                missing_fields TEXT NOT NULL DEFAULT '[]',
                conflict_fields TEXT NOT NULL DEFAULT '[]',
                attachment_count INTEGER NOT NULL DEFAULT 0,
                consent_confirmed INTEGER NOT NULL DEFAULT 0,
                UNIQUE(invite_id, version)
            );
            """
        )
        db.executemany(
            "INSERT INTO invites(invite_id,employee_id,name,token_hash,status,created_at) VALUES(:invite_id,:employee_id,:name,:token_hash,'invited',:created_at)",
            rows,
        )
    os.chmod(root / "state.sqlite3", 0o600)


def default_config() -> dict[str, Any]:
    return {
        "title": "基础身份信息收集",
        "purpose": "请填写本次工作所需的基础身份信息。请将用途修改为具体、必要的业务目的。",
        "deadline": "2026-12-31T23:59:59+08:00",
        "retention_until": "2027-01-31T23:59:59+08:00",
        "contact": "请填写联系人和联系方式",
        "correction": "如需更正，请联系任务发起人并使用原邀请重新提交。",
        "template_version": "1.0",
        "fields": [
            {"id": "name", "label": "姓名", "type": "text", "required": True, "sensitive": False},
            {"id": "phone", "label": "手机号", "type": "phone_cn", "required": True, "sensitive": True},
            {"id": "id_number", "label": "身份证号", "type": "cn_id", "required": True, "sensitive": True},
            {"id": "address", "label": "住址", "type": "address", "required": True, "sensitive": True},
            {"id": "id_front", "label": "身份证正面", "type": "image_attachment", "required": True, "sensitive": True},
            {"id": "id_back", "label": "身份证反面", "type": "image_attachment", "required": True, "sensitive": True},
        ],
    }


def validate_config(config: dict[str, Any], mode: str = "directed") -> dict[str, Any]:
    if mode not in TASK_MODES:
        raise ValueError("mode 必须是 directed 或 group")
    for key in ("title", "purpose", "deadline", "retention_until", "contact", "correction", "template_version", "fields"):
        if not config.get(key):
            raise ValueError(f"配置缺少必填项: {key}")
    deadline, retention = parse_time(config["deadline"]), parse_time(config["retention_until"])
    if deadline <= datetime.now(timezone.utc):
        raise ValueError("deadline 必须晚于当前时间")
    if retention <= deadline:
        raise ValueError("retention_until 必须晚于 deadline")
    defaults = default_config()
    for key in ("purpose", "contact", "correction"):
        if config[key] == defaults[key]:
            raise ValueError(f"请先把 {key} 的占位内容改为真实、具体的信息")
    seen = set()
    fields = []
    for raw in config["fields"]:
        field = dict(raw)
        field_id, field_type = field.get("id", ""), field.get("type", "")
        if not FIELD_ID_RE.fullmatch(field_id) or field_id in seen:
            raise ValueError(f"字段 id 无效或重复: {field_id}")
        if field_type not in ALLOWED_TYPES:
            raise ValueError(f"不支持的字段类型: {field_type}")
        if not field.get("label"):
            raise ValueError(f"字段缺少 label: {field_id}")
        if field_type == "single_choice" and not field.get("options"):
            raise ValueError(f"单选字段缺少 options: {field_id}")
        field["required"] = bool(field.get("required"))
        field["sensitive"] = bool(field.get("sensitive", field_type in {"phone_cn", "cn_id", "address"} or field_type in ATTACHMENT_TYPES))
        if field_type in ATTACHMENT_TYPES:
            field["multiple"] = bool(field.get("multiple", False))
        fields.append(field)
        seen.add(field_id)
    if "name" not in seen:
        raise ValueError("任务字段必须包含 id=name 的姓名字段")
    if mode == "group" and "employee_id" not in seen:
        raise ValueError("group 模式任务字段必须包含 id=employee_id 的工号字段（无令牌模式下身份靠工号+姓名与名单比对）")
    result = dict(config)
    result["fields"] = fields
    return result


def read_roster(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames or not {"employee_id", "name"}.issubset(reader.fieldnames):
            raise ValueError("名单 CSV 必须包含 employee_id,name 表头")
        rows = []
        ids = set()
        for line, raw in enumerate(reader, 2):
            employee_id, name = (raw.get("employee_id") or "").strip(), (raw.get("name") or "").strip()
            if not employee_id or not name:
                raise ValueError(f"名单第 {line} 行缺少 employee_id 或 name")
            if employee_id in ids:
                raise ValueError(f"名单 employee_id 重复: {employee_id}")
            ids.add(employee_id)
            rows.append({"employee_id": employee_id, "name": name})
    if not rows:
        raise ValueError("名单为空")
    return rows


def derive_scrypt_key(password: str, salt: bytes, n: int = SCRYPT_N, r: int = SCRYPT_R, p: int = SCRYPT_P) -> bytes:
    from cryptography.hazmat.primitives.kdf.scrypt import Scrypt

    return Scrypt(salt=salt, length=32, n=n, r=r, p=p).derive(password.encode("utf-8"))


def validate_kdf_params(kdf: Any) -> tuple[bytes, int, int, int]:
    if not isinstance(kdf, dict) or kdf.get("name") != "scrypt":
        raise ValueError("KDF 参数不受支持")
    try:
        salt = base64.b64decode(kdf.get("salt", ""), validate=True)
        n, r, p = int(kdf["n"]), int(kdf["r"]), int(kdf["p"])
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError("KDF 参数无效") from exc
    if len(salt) != 16 or not (2**10 <= n <= 2**16 and n & (n - 1) == 0) or not (1 <= r <= 8 and 1 <= p <= 2):
        raise ValueError("KDF 参数超出允许范围")
    return salt, n, r, p


def aes_gcm_seal(format_name: str, plaintext: bytes, password: str, aad: bytes | None = None) -> dict[str, Any]:
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    salt, nonce = secrets.token_bytes(16), secrets.token_bytes(12)
    key = derive_scrypt_key(password, salt)
    ciphertext = AESGCM(key).encrypt(nonce, plaintext, aad if aad is not None else format_name.encode("utf-8"))
    return {
        "format": format_name,
        "kdf": {"name": "scrypt", "salt": base64.b64encode(salt).decode("ascii"), "n": SCRYPT_N, "r": SCRYPT_R, "p": SCRYPT_P},
        "cipher": "AES-256-GCM",
        "nonce": base64.b64encode(nonce).decode("ascii"),
        "ciphertext": base64.b64encode(ciphertext).decode("ascii"),
    }


def aes_gcm_open(envelope: dict[str, Any], password: str, aad: bytes) -> bytes:
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    if not isinstance(envelope, dict) or envelope.get("cipher") != "AES-256-GCM":
        raise ValueError("加密信封格式不受支持")
    salt, n, r, p = validate_kdf_params(envelope.get("kdf"))
    try:
        nonce = base64.b64decode(envelope["nonce"], validate=True)
        ciphertext = base64.b64decode(envelope["ciphertext"], validate=True)
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError("加密信封缺少 nonce 或密文") from exc
    if len(nonce) != 12:
        raise ValueError("加密信封 nonce 无效")
    key = derive_scrypt_key(password, salt, n, r, p)
    return AESGCM(key).decrypt(nonce, ciphertext, aad)


def encrypt_private_key(private_der: bytes, password: str) -> bytes:
    envelope = aes_gcm_seal(KEY_ENVELOPE_VERSION, private_der, password)
    return json.dumps(envelope, ensure_ascii=False, indent=2).encode("utf-8") + b"\n"


def decrypt_private_key(envelope: dict[str, Any], password: str) -> bytes:
    if envelope.get("format") != KEY_ENVELOPE_VERSION:
        raise ValueError("私钥信封格式不受支持")
    return aes_gcm_open(envelope, password, KEY_ENVELOPE_VERSION.encode("utf-8"))


def valid_private_key_blob(data: bytes) -> bool:
    blob = data.lstrip()
    if blob.startswith(LEGACY_KEY_PEM_PREFIX):
        return True
    if blob.startswith(b"{"):
        try:
            envelope = json.loads(blob)
        except (json.JSONDecodeError, UnicodeDecodeError):
            return False
        return isinstance(envelope, dict) and envelope.get("format") == KEY_ENVELOPE_VERSION
    return False


def generate_keys(password: str) -> tuple[bytes, bytes, str]:
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa

    private = rsa.generate_private_key(public_exponent=65537, key_size=3072)
    public = private.public_key()
    public_pem = public.public_bytes(serialization.Encoding.PEM, serialization.PublicFormat.SubjectPublicKeyInfo)
    public_der = public.public_bytes(serialization.Encoding.DER, serialization.PublicFormat.SubjectPublicKeyInfo)
    private_der = private.private_bytes(serialization.Encoding.DER, serialization.PrivateFormat.PKCS8, serialization.NoEncryption())
    return public_pem, encrypt_private_key(private_der, password), hashlib.sha256(public_der).hexdigest()[:24]


def safe_json_for_html(value: Any) -> str:
    return json.dumps(value, ensure_ascii=True, separators=(",", ":")).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")


def csv_text(value: str) -> str:
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def render_invite(template: str, config: dict[str, Any]) -> str:
    return template.replace("__YINTIAN_CONFIG__", safe_json_for_html(config))


def cmd_init_config(args) -> dict[str, Any]:
    out = Path(args.out).expanduser()
    if out.exists() and not args.force:
        raise FileExistsError(f"文件已存在: {out}")
    dump_json(out, default_config())
    return {"config": str(out.resolve())}


def create_task(roster_path: Path, config_path: Path, out_parent: Path, password: str, require_terminal: bool = True, mode: str = "directed") -> dict[str, Any]:
    if require_terminal:
        require_tty("create")
    if mode == "group":
        print("风险提示：group 群发模式没有个人认证令牌，任何拿到表单的人都可以冒名为任何工号提交；"
              "身份仅靠复核时名单比对（工号+姓名）标记并由人工兜底，高敏感场景请改用 directed 模式。", file=sys.stderr)
    config = validate_config(load_json(config_path), mode=mode)
    roster = read_roster(roster_path)
    task_id = "YT-" + datetime.now().strftime("%Y%m%d") + "-" + random_id("", 6)
    root = out_parent.expanduser().resolve() / task_id
    if root.exists():
        raise FileExistsError(f"任务目录已存在: {root}")
    root.mkdir(parents=True, mode=0o700)
    (root / "invites").mkdir(mode=0o700)
    (root / "submissions").mkdir(mode=0o700)
    (root / "reports").mkdir(mode=0o700)

    public_pem, private_pem, key_id = generate_keys(password)
    schema_hash = sha256_bytes(canonical(config["fields"]))
    notice = {key: config[key] for key in ("title", "purpose", "deadline", "retention_until", "contact", "correction")}
    notice_hash = sha256_bytes(canonical(notice))
    task = {
        "format_version": FORMAT_VERSION,
        "mode": mode,
        "task_id": task_id,
        "key_id": key_id,
        "title": config["title"],
        "purpose": config["purpose"],
        "deadline": config["deadline"],
        "retention_until": config["retention_until"],
        "contact": config["contact"],
        "correction": config["correction"],
        "template_version": config["template_version"],
        "schema_hash": schema_hash,
        "notice_hash": notice_hash,
        "fields": config["fields"],
        "created_at": now_iso(),
    }
    dump_json(root / "task.json", task)
    atomic_write(root / "public.pem", public_pem)
    atomic_write(root / "private.pem.enc", private_pem)
    atomic_write(root / "roster.csv", roster_path.read_bytes())

    db_rows, index_rows = [], []
    if mode == "group":
        form = {"format": FORM_FORMAT_VERSION, **task, "public_key_pem": public_pem.decode("ascii")}
        dump_json(root / "FORM.yintian-form", form)
        for item in roster:
            invite_id = group_invite_id(item["employee_id"])
            parse_group_employee_id(invite_id)  # 创建期闭环校验：名单工号必须能构成合法的 GRP- 标识
            db_rows.append({**item, "invite_id": invite_id, "token_hash": sha256_bytes(secrets.token_bytes(32)), "created_at": now_iso()})
            index_rows.append({**item, "invite_id": invite_id, "invite_file": "FORM.yintian-form"})
    else:
        template = (Path(__file__).resolve().parents[1] / "assets" / "invite_template.html").read_text(encoding="utf-8")
        for item in roster:
            invite_id = random_id("INV-", 10)
            invite_token = secrets.token_urlsafe(32)
            invite_name = invite_id + ".html"
            invite_config = dict(task)
            invite_config.update({"invite_id": invite_id, "invite_token": invite_token, "name": item["name"], "public_key_pem": public_pem.decode("ascii")})
            atomic_write(root / "invites" / invite_name, render_invite(template, invite_config).encode("utf-8"))
            dump_json(root / "invites" / (invite_id + ".yintian-form"), {"format": FORM_FORMAT_VERSION, **invite_config})
            created_at = now_iso()
            db_rows.append({**item, "invite_id": invite_id, "token_hash": sha256_bytes(invite_token.encode()), "created_at": created_at})
            index_rows.append({**item, "invite_id": invite_id, "invite_file": f"invites/{invite_name}"})
    index_path = root / "invite-index.csv"
    index_buffer = io.StringIO()
    writer = csv.DictWriter(index_buffer, fieldnames=["employee_id", "name", "invite_id", "invite_file"])
    writer.writeheader()
    writer.writerows([{key: csv_text(value) for key, value in row.items()} for row in index_rows])
    atomic_write(index_path, index_buffer.getvalue().encode("utf-8-sig"))
    init_db(root, db_rows)
    return {"task_id": task_id, "task_dir": str(root), "invite_count": len(index_rows)}


def cmd_create(args) -> dict[str, Any]:
    require_tty("create")
    password = generate_password()
    result = create_task(Path(args.roster), Path(args.config), Path(args.out), password, require_terminal=False, mode=getattr(args, "mode", "directed"))
    print_once_secret("任务密码（仅显示一次，丢失不可恢复）：", password, "请通过独立安全渠道交给授权处理人员，不要写入任务目录或聊天提示词。")
    return result


def validate_envelope_header(envelope: dict[str, Any], task: dict[str, Any]) -> str:
    for key in ("format_version", "task_id", "invite_id", "schema_hash", "key_id", "encrypted_key_b64", "iv_b64", "ciphertext_b64"):
        if not isinstance(envelope.get(key), str) or not envelope[key]:
            raise ValueError(f"提交包缺少字段: {key}")
    if envelope["format_version"] != FORMAT_VERSION or envelope["task_id"] != task["task_id"]:
        raise ValueError("提交包属于其他任务或格式版本")
    if envelope.get("algorithms") != {"content": "AES-256-GCM", "key_wrap": "RSA-OAEP-3072-SHA256"}:
        raise ValueError("提交包算法套件不受支持")
    if envelope["schema_hash"] != task["schema_hash"] or envelope["key_id"] != task["key_id"]:
        raise ValueError("提交包字段模板或公钥不匹配")
    invite_id = envelope["invite_id"]
    if task_mode(task) == "group":
        parse_group_employee_id(invite_id)
    elif not INVITE_ID_RE.fullmatch(invite_id):
        raise ValueError("invite_id 格式无效")
    for key in ("encrypted_key_b64", "iv_b64", "ciphertext_b64"):
        base64.b64decode(envelope[key], validate=True)
    return invite_id


def ingest_task(task_dir: str | Path, submissions_dir: str | Path) -> dict[str, Any]:
    root, task = load_task(task_dir)
    require_current_format(task, "接收")
    if task_expired(task):
        raise RuntimeError("任务已超过保存期限，停止接收新提交")
    source = Path(submissions_dir).expanduser().resolve()
    if not source.is_dir():
        raise NotADirectoryError(source)
    summary = {"accepted": 0, "duplicates": 0, "rejected": 0, "errors": []}
    with task_lock(root), closing(connect_db(root)) as db, db:
        for index, path in enumerate(sorted(source.glob("*.yintian")), start=1):
            digest = None
            try:
                if not path.is_file() or path.is_symlink():
                    raise ValueError("提交项不是普通文件")
                if path.stat().st_size > MAX_ENVELOPE_BYTES:
                    raise ValueError("提交包超过 32MB 上限")
                digest = sha256_file(path)
                if db.execute("SELECT 1 FROM submissions WHERE sha256=?", (digest,)).fetchone():
                    summary["duplicates"] += 1
                    continue
                envelope = load_json(path)
                invite_id = validate_envelope_header(envelope, task)
                if not db.execute("SELECT 1 FROM invites WHERE invite_id=?", (invite_id,)).fetchone():
                    raise ValueError("invite_id 不属于当前任务")
                existing_versions = db.execute("SELECT COUNT(*) FROM submissions WHERE invite_id=?", (invite_id,)).fetchone()[0]
                if existing_versions >= MAX_VERSIONS_PER_INVITE:
                    raise ValueError(f"该邀请的提交版本数已达上限 {MAX_VERSIONS_PER_INVITE}，拒绝继续存储")
                version = db.execute("SELECT COALESCE(MAX(version),0)+1 FROM submissions WHERE invite_id=?", (invite_id,)).fetchone()[0]
                target_dir = root / "submissions" / invite_id
                target_dir.mkdir(parents=True, exist_ok=True)
                target = target_dir / f"v{version:04d}_{digest[:12]}.yintian"
                shutil.copyfile(path, target)
                received_at = now_iso()
                db.execute(
                    "INSERT INTO submissions(invite_id,version,sha256,path,received_at,late,status) VALUES(?,?,?,?,?,?,?)",
                    (invite_id, version, digest, target.relative_to(root).as_posix(), received_at, int(task_late(task, received_at)), "submitted"),
                )
                db.execute("UPDATE invites SET status='submitted' WHERE invite_id=? AND status='invited'", (invite_id,))
                summary["accepted"] += 1
            except Exception as exc:
                summary["rejected"] += 1
                summary["errors"].append({"file_ref": digest[:12] if digest else f"entry-{index}", "error": type(exc).__name__})
    return summary


def cmd_ingest(args) -> dict[str, Any]:
    return ingest_task(args.task_dir, args.submissions_dir)


def unlock_private_key(root: Path, password: str):
    from cryptography.hazmat.primitives import serialization

    data = (root / "private.pem.enc").read_bytes()
    blob = data.lstrip()
    if blob.startswith(b"{"):
        private_der = decrypt_private_key(json.loads(data), password)
        return serialization.load_der_private_key(private_der, password=None)
    if not blob.startswith(LEGACY_KEY_PEM_PREFIX):
        raise ValueError("私钥文件既不是 yintian-key/1 信封，也不是加密的 PKCS#8 PEM，拒绝解锁")
    return serialization.load_pem_private_key(data, password=password.encode())


def aad_for(envelope: dict[str, Any]) -> bytes:
    return canonical([envelope["format_version"], envelope["task_id"], envelope["invite_id"], envelope["schema_hash"], envelope["key_id"]])


def decrypt_envelope(root: Path, task: dict[str, Any], path: Path, private_key, expected_invite_id: str) -> dict[str, Any]:
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.asymmetric import padding
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    path = path.resolve()
    submissions_root = (root / "submissions").resolve()
    if submissions_root not in path.parents or not path.is_file():
        raise ValueError("提交路径不在任务 submissions 目录内")
    envelope = load_json(path)
    invite_id = validate_envelope_header(envelope, task)
    if invite_id != expected_invite_id:
        raise ValueError("提交包与数据库邀请不匹配")
    encrypted_key = base64.b64decode(envelope["encrypted_key_b64"], validate=True)
    aes_key = private_key.decrypt(encrypted_key, padding.OAEP(mgf=padding.MGF1(algorithm=hashes.SHA256()), algorithm=hashes.SHA256(), label=None))
    plaintext = AESGCM(aes_key).decrypt(base64.b64decode(envelope["iv_b64"], validate=True), base64.b64decode(envelope["ciphertext_b64"], validate=True), aad_for(envelope))
    payload = json.loads(plaintext)
    for key in ("format_version", "task_id", "invite_id", "schema_hash"):
        if payload.get(key) != envelope[key]:
            raise ValueError(f"解密载荷 {key} 与外层不一致")
    return payload


def normalize_value(field_type: str, value: str) -> str:
    value = str(value or "").strip()
    if field_type == "phone_cn":
        value = re.sub(r"[\s-]", "", value)
        value = re.sub(r"^\+?86", "", value)
    if field_type == "cn_id":
        value = value.upper()
    return value


def validate_payload(task: dict[str, Any], invite: sqlite3.Row, payload: dict[str, Any]) -> tuple[list[str], list[str], list[dict[str, Any]]]:
    from ocr_matcher import validate_chinese_id

    missing, conflicts, attachment_items = [], [], []
    if payload.get("notice_hash") != task["notice_hash"]:
        conflicts.append("notice_hash")
    if payload.get("template_version") != task["template_version"]:
        conflicts.append("template_version")
    if payload.get("consent_confirmed") is not True:
        conflicts.append("consent")
    token = payload.get("invite_token", "")
    if task_mode(task) == "directed" and (not isinstance(token, str) or not secrets.compare_digest(sha256_bytes(token.encode()), invite["token_hash"])):
        raise ValueError("邀请认证令牌无效")
    if normalize_submitted_at(payload.get("submitted_at")) is None:
        conflicts.append("submitted_at")
    values, attachments = payload.get("values", {}), payload.get("attachments", {})
    if not isinstance(values, dict) or not isinstance(attachments, dict):
        raise ValueError("values/attachments 格式无效")
    total_size = 0
    for field in task["fields"]:
        field_id, field_type = field["id"], field["type"]
        if field_type in ATTACHMENT_TYPES:
            items = attachments.get(field_id, [])
            if not isinstance(items, list):
                raise ValueError(f"附件字段格式无效: {field_id}")
            if field.get("required") and not items:
                missing.append(field_id)
            if not field.get("multiple") and len(items) > 1:
                conflicts.append(field_id)
            for item in items:
                raw = base64.b64decode(item.get("data_b64", ""), validate=True)
                mime = item.get("type", "")
                if len(raw) != int(item.get("size", -1)) or len(raw) > MAX_FILE_BYTES or sha256_bytes(raw) != item.get("sha256"):
                    raise ValueError(f"附件大小或哈希无效: {field_id}")
                if field_type == "image_attachment" and mime not in {"image/jpeg", "image/png", "image/webp"}:
                    raise ValueError(f"图片类型不支持: {mime}")
                if field_type == "pdf_attachment" and mime != "application/pdf":
                    raise ValueError("PDF 附件类型无效")
                if field_type == "image_attachment":
                    from PIL import Image

                    with Image.open(io.BytesIO(raw)) as image:
                        expected = {"image/jpeg": "JPEG", "image/png": "PNG", "image/webp": "WEBP"}[mime]
                        if image.format != expected or image.width * image.height > MAX_IMAGE_PIXELS:
                            raise ValueError(f"图片内容或像素尺寸无效: {field_id}")
                        image.verify()
                elif not raw.startswith(b"%PDF-"):
                    raise ValueError("PDF 文件头无效")
                total_size += len(raw)
                attachment_items.append({"field_id": field_id, "name": item.get("name", ""), "type": mime, "data": raw})
        else:
            value = normalize_value(field_type, values.get(field_id, ""))
            if field.get("required") and not value:
                missing.append(field_id)
            if value and field_type == "phone_cn" and not re.fullmatch(r"1[3-9]\d{9}", value):
                conflicts.append(field_id)
            if value and field_type == "cn_id" and not validate_chinese_id(value):
                conflicts.append(field_id)
            if value and field_type == "date" and not valid_date(value):
                conflicts.append(field_id)
            if value and field_type == "single_choice" and value not in field.get("options", []):
                conflicts.append(field_id)
    if total_size > MAX_TOTAL_BYTES:
        raise ValueError("附件总大小超过 15MB")
    if normalize_value("text", values.get("name", "")) != invite["name"]:
        conflicts.append("name_roster_mismatch")
    if task_mode(task) == "group" and normalize_value("text", values.get("employee_id", "")) != invite["employee_id"]:
        conflicts.append("employee_id_roster_mismatch")
    return sorted(set(missing)), sorted(set(conflicts)), attachment_items


def ocr_attachment(item: dict[str, Any]) -> tuple[str, list[str]]:
    try:
        import ocr_engine
    except Exception as exc:
        return "", [f"ocr_unavailable:{type(exc).__name__}"]
    try:
        if item["type"].startswith("image/"):
            result = ocr_engine.scan_bytes(item["data"])
            flags = ["ocr_low_confidence"] if any(block.get("low_confidence") for block in result["blocks"]) else []
            text = result["full_text"]
            if not text.strip():
                flags.append("ocr_no_text")
            return text, flags
        if item["type"] == "application/pdf":
            try:
                import pypdfium2 as pdfium
            except ImportError:
                return "", ["pdf_ocr_unavailable"]
            document = pdfium.PdfDocument(item["data"])
            flags = []
            if len(document) > MAX_PDF_PAGES:
                flags.append("pdf_page_limit")
            texts = []
            for page_index in range(min(len(document), MAX_PDF_PAGES)):
                page = document[page_index]
                width, height = page.get_size()
                rendered_w, rendered_h = width * PDF_RENDER_SCALE, height * PDF_RENDER_SCALE
                if rendered_w * rendered_h > MAX_IMAGE_PIXELS:
                    flags.append("pdf_pixel_limit")
                    continue
                bitmap = page.render(scale=PDF_RENDER_SCALE)
                image = bitmap.to_pil()
                buffer = io.BytesIO()
                image.save(buffer, format="PNG")
                result = ocr_engine.scan_bytes(buffer.getvalue())
                texts.append(result["full_text"])
                if any(block.get("low_confidence") for block in result["blocks"]):
                    flags.append("ocr_low_confidence")
            text = "\n".join(texts)
            if not text.strip():
                flags.append("ocr_no_text")
            return text, sorted(set(flags))
    except Exception as exc:
        return "", [f"ocr_error:{type(exc).__name__}"]
    return "", []


def is_identity_attachment_field(field_id: str) -> bool:
    """附件字段按下划线分词后含整词 front/back 才视为证件正反面（避免误伤 backdrop、feedback_scan 等）。"""
    return any(part in {"front", "back"} for part in field_id.split("_"))


def compare_ocr(payload: dict[str, Any], attachments: list[dict[str, Any]], field_defs: list[dict[str, Any]]) -> list[str]:
    """按字段类型匹配填写值；对 id 分词后含整词 front/back 的附件字段做 OCR 身份要素比对。"""
    from ocr_matcher import extract_from_text as _extract_from_text

    fields: dict[str, list[Any]] = {}
    flags: list[str] = []
    identity_field_ids = set()
    checks = []
    for field in field_defs:
        field_id, field_type = field["id"], field["type"]
        if field_type in ATTACHMENT_TYPES and is_identity_attachment_field(field_id):
            identity_field_ids.add(field_id)
        if field_id == "name":
            checks.append((field_id, "name", "text"))
        elif field_type == "cn_id":
            checks.append((field_id, "id_number", "cn_id"))
        elif field_type == "phone_cn":
            checks.append((field_id, "phone", "phone_cn"))
        elif field_type == "address":
            checks.append((field_id, "address", "text"))
    identity_attachments = [item for item in attachments if item["field_id"] in identity_field_ids]
    for item in identity_attachments:
        text, item_flags = ocr_attachment(item)
        flags.extend(item_flags)
        if text:
            _extract_from_text(text, f"attachment:{item['field_id']}", fields)
    if identity_attachments and not any(fields.get(field) for field in ("name", "id_number", "address")):
        flags.append("ocr_no_identity_fields")
    values = payload.get("values", {})
    for submitted_id, extracted_id, value_type in checks:
        manual = normalize_value(value_type, values.get(submitted_id, ""))
        candidates = {normalize_value(value_type, candidate.value) for candidate in fields.get(extracted_id, []) if candidate.valid is not False}
        if manual and candidates and manual not in candidates:
            flags.append(submitted_id)
    return sorted(set(flags))


def review_task(task_dir: str | Path, password: str) -> dict[str, Any]:
    root, task = load_task(task_dir)
    require_current_format(task, "复核")
    if task_expired(task):
        raise RuntimeError("任务已超过保存期限，停止复核")
    try:
        private_key = unlock_private_key(root, password)
    except Exception as exc:
        raise RuntimeError("任务密码错误或私钥损坏") from exc
    summary = {"verified": 0, "needs_review": 0, "invalid": 0}
    with task_lock(root), closing(connect_db(root)) as db, db:
        rows = db.execute(
            "SELECT s.*,i.employee_id,i.name,i.token_hash FROM submissions s JOIN invites i ON i.invite_id=s.invite_id WHERE s.status='submitted' ORDER BY s.invite_id,s.version"
        ).fetchall()
        for row in rows:
            try:
                payload = decrypt_envelope(root, task, root / row["path"], private_key, row["invite_id"])
                missing, conflicts, attachments = validate_payload(task, row, payload)
                conflicts = sorted(set(conflicts + compare_ocr(payload, attachments, task["fields"])))
                status = "needs_review" if missing or conflicts else "verified"
                reviewed_at = now_iso()
                db.execute(
                    "UPDATE submissions SET submitted_at=?,reviewed_at=?,status=?,missing_fields=?,conflict_fields=?,attachment_count=?,consent_confirmed=? WHERE id=?",
                    (normalize_submitted_at(payload.get("submitted_at")), reviewed_at, status, json.dumps(missing), json.dumps(conflicts), len(attachments), int(payload.get("consent_confirmed") is True), row["id"]),
                )
                old = db.execute("SELECT current_submission_id FROM invites WHERE invite_id=?", (row["invite_id"],)).fetchone()[0]
                old_row = db.execute("SELECT id,status FROM submissions WHERE id=? AND status IN ('verified','needs_review')", (old,)).fetchone() if old else None
                candidate_is_newer = not old_row or row["id"] > old_row["id"]
                if candidate_is_newer:
                    if old_row and old_row["id"] != row["id"]:
                        db.execute("UPDATE submissions SET status='superseded' WHERE id=?", (old_row["id"],))
                    db.execute("UPDATE invites SET current_submission_id=?,status=? WHERE invite_id=?", (row["id"], status, row["invite_id"]))
                elif row["id"] != old:
                    db.execute("UPDATE submissions SET status='superseded' WHERE id=?", (row["id"],))
                summary[status] += 1
            except Exception:
                db.execute("UPDATE submissions SET reviewed_at=?,status='invalid' WHERE id=?", (now_iso(), row["id"]))
                old = db.execute("SELECT current_submission_id FROM invites WHERE invite_id=?", (row["invite_id"],)).fetchone()[0]
                if not old:
                    db.execute("UPDATE invites SET status='invalid' WHERE invite_id=?", (row["invite_id"],))
                summary["invalid"] += 1
    return summary


def cmd_review(args) -> dict[str, Any]:
    require_tty("review")
    password = getpass.getpass("任务密码: ")
    return review_task(args.task_dir, password)


def report_rows(root: Path) -> list[dict[str, Any]]:
    with closing(connect_db(root)) as db, db:
        rows = db.execute(
            """
            SELECT i.employee_id,i.name,i.invite_id,i.status,
                   s.version AS submission_version,s.submitted_at,s.reviewed_at,s.late,
                   s.missing_fields,s.conflict_fields,s.attachment_count,s.consent_confirmed
            FROM invites i LEFT JOIN submissions s ON s.id=i.current_submission_id AND s.invite_id=i.invite_id
            ORDER BY i.employee_id
            """
        ).fetchall()
    output = []
    for row in rows:
        output.append({
            "employee_id": row["employee_id"], "name": row["name"], "status": row["status"],
            "submission_version": row["submission_version"], "submitted_at": row["submitted_at"], "reviewed_at": row["reviewed_at"],
            "late": bool(row["late"]) if row["late"] is not None else False,
            "missing_fields": json.loads(row["missing_fields"] or "[]"), "conflict_fields": json.loads(row["conflict_fields"] or "[]"),
            "attachment_count": row["attachment_count"] or 0, "consent_confirmed": bool(row["consent_confirmed"]) if row["consent_confirmed"] is not None else False,
        })
    return output


def status_task(task_dir: str | Path) -> dict[str, Any]:
    root, task = load_task(task_dir)
    rows = report_rows(root)
    counts: dict[str, int] = {}
    for row in rows:
        counts[row["status"]] = counts.get(row["status"], 0) + 1
    return {"task_id": task["task_id"], "task_status": "expired" if task_expired(task) else "active", "deadline": task["deadline"], "retention_until": task["retention_until"], "counts": counts, "total": len(rows)}


def cmd_status(args) -> dict[str, Any]:
    return status_task(args.task_dir)


def write_reports(task_dir: str | Path, formats: list[str]) -> dict[str, str]:
    root, task = load_task(task_dir)
    if task_expired(task):
        raise RuntimeError("任务已超过保存期限，停止生成报告；请执行清理")
    unknown = set(formats) - {"xlsx", "json"}
    if unknown:
        raise ValueError(f"不支持的报告格式: {sorted(unknown)}")
    rows = report_rows(root)
    reports = root / "reports"
    reports.mkdir(exist_ok=True)
    written = {}
    if "json" in formats:
        path = reports / "progress.json"
        dump_json(path, {"task_id": task["task_id"], "generated_at": now_iso(), "rows": rows})
        written["json"] = str(path)
    if "xlsx" in formats:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        path = reports / "progress.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "收集进度"
        columns = ["employee_id", "name", "status", "submission_version", "submitted_at", "reviewed_at", "late", "missing_fields", "conflict_fields", "attachment_count", "consent_confirmed"]
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row_index, row in enumerate(rows, start=2):
            for column_index, key in enumerate(columns, start=1):
                value = ",".join(row[key]) if isinstance(row[key], list) else row[key]
                cell = sheet.cell(row_index, column_index, value)
                if isinstance(value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(path)
        os.chmod(path, 0o600)
        written["xlsx"] = str(path)
    return written


def cmd_report(args) -> dict[str, Any]:
    formats = [part.strip().lower() for value in args.formats for part in value.split(",") if part.strip()]
    return write_reports(args.task_dir, sorted(set(formats)))


def current_submission(db: sqlite3.Connection, invite_id: str) -> sqlite3.Row:
    row = db.execute(
        "SELECT s.*,i.employee_id,i.name,i.token_hash FROM invites i JOIN submissions s ON s.id=i.current_submission_id AND s.invite_id=i.invite_id WHERE i.invite_id=?",
        (invite_id,),
    ).fetchone()
    if not row:
        raise ValueError("该邀请没有可查看的有效提交")
    return row


def cmd_reveal(args) -> dict[str, Any] | None:
    require_tty("reveal")
    root, task = load_task(args.task_dir)
    require_current_format(task, "查看明文")
    if task_expired(task):
        raise RuntimeError("任务已超过保存期限，禁止查看明文")
    password = getpass.getpass("任务密码: ")
    try:
        private_key = unlock_private_key(root, password)
    except Exception as exc:
        raise RuntimeError("任务密码错误或私钥损坏") from exc
    with closing(connect_db(root)) as db, db:
        row = current_submission(db, args.invite_id)
    payload = decrypt_envelope(root, task, root / row["path"], private_key, args.invite_id)
    validate_payload(task, row, payload)
    terminal_text = lambda value: re.sub(r"[\x00-\x1f\x7f-\x9f]", " ", str(value))
    print("\n=== 授权明文查看 ===")
    print(f"employee_id: {terminal_text(row['employee_id'])}\nname: {terminal_text(row['name'])}\ninvite_id: {args.invite_id}")
    for field in task["fields"]:
        if field["type"] not in ATTACHMENT_TYPES:
            print(f"{terminal_text(field['label'])}: {terminal_text(payload.get('values', {}).get(field['id'], ''))}")
    print("附件:")
    for field_id, items in payload.get("attachments", {}).items():
        for item in items:
            print(f"- {terminal_text(field_id)}: {terminal_text(item.get('name'))} ({terminal_text(item.get('type'))}, {terminal_text(item.get('size'))} bytes)")
    input("\n查看完毕后按 Enter 清空当前屏幕与终端回滚缓冲区…")
    print("\033[2J\033[3J\033[H", end="", flush=True)
    return None


def parse_export_fields(task: dict[str, Any], raw_fields: str | None) -> list[str]:
    if not raw_fields or not raw_fields.strip():
        raise ValueError("必须用 --fields 显式声明导出字段白名单（逗号分隔的字段 id）")
    field_defs = {field["id"]: field for field in task["fields"]}
    field_ids = []
    for part in raw_fields.split(","):
        field_id = part.strip()
        if not field_id or field_id in {"employee_id", "name"}:
            continue  # employee_id/name 默认附带，无需在白名单中重复声明
        field = field_defs.get(field_id)
        if field is None:
            raise ValueError(f"导出字段不在任务字段中: {field_id}")
        if field["type"] in ATTACHMENT_TYPES:
            raise ValueError(f"附件字段不支持明文导出: {field_id}")
        if field_id in field_ids:
            raise ValueError(f"导出字段重复: {field_id}")
        field_ids.append(field_id)
    if not field_ids:
        raise ValueError("导出字段白名单为空（employee_id/name 默认附带，无需声明）")
    return field_ids


def parse_mask_specs(specs: list[str] | None) -> dict[str, str]:
    masks = {}
    for spec in specs or []:
        field_id, sep, rule = spec.partition("=")
        field_id, rule = field_id.strip(), rule.strip()
        if not sep or not FIELD_ID_RE.fullmatch(field_id) or rule not in MASK_RULES:
            raise ValueError(f"脱敏规则无效（应为 字段id=last4|mid4）: {spec}")
        masks[field_id] = rule
    return masks


def mask_value(rule: str, value: Any) -> str:
    text = str(value)
    if rule == "last4":
        return "*" * (len(text) - 4) + text[-4:] if len(text) > 4 else "*" * len(text)
    if rule == "mid4":
        return text[:3] + "****" + text[-4:] if len(text) > 7 else "*" * len(text)
    raise ValueError(f"不支持的脱敏规则: {rule}")


def collect_clear_rows(root: Path, task: dict[str, Any], private_key, field_ids: list[str], masks: dict[str, str]) -> list[dict[str, str]]:
    """内存解密全部 current+valid 提交（与 review 相同的解密/校验路径），只保留白名单字段。"""
    field_defs = {field["id"]: field for field in task["fields"]}
    with closing(connect_db(root)) as db, db:
        rows = db.execute(
            "SELECT s.*,i.employee_id,i.name,i.token_hash FROM submissions s JOIN invites i ON i.invite_id=s.invite_id WHERE s.id=i.current_submission_id AND s.status IN ('verified','needs_review') ORDER BY i.employee_id"
        ).fetchall()
    output = []
    for row in rows:
        payload = decrypt_envelope(root, task, root / row["path"], private_key, row["invite_id"])
        validate_payload(task, row, payload)  # 导出前再跑一次完整校验，密文任何损坏都直接中止
        values = payload.get("values", {})
        record = {}
        for column, raw in (("employee_id", row["employee_id"]), ("name", row["name"])):
            rule = masks.get(column)
            record[column] = mask_value(rule, raw) if rule and raw else raw
        for field_id in field_ids:
            value = normalize_value(field_defs[field_id]["type"], values.get(field_id, ""))
            rule = masks.get(field_id)
            record[field_id] = mask_value(rule, value) if rule and value else value
        output.append(record)
    return output


def write_clear_export(root: Path, task: dict[str, Any], rows: list[dict[str, str]], field_ids: list[str], masks: dict[str, str],
                       out: str | Path, formats: list[str], purpose: str, recipient: str) -> dict[str, str]:
    """把已解密的明文行落盘为 XLSX/JSON（0600，字符串单元格防公式注入）；不产出任何加密出口信封。"""
    unknown = set(formats) - {"xlsx", "json"}
    if unknown:
        raise ValueError(f"不支持的导出格式: {sorted(unknown)}")
    requested_out = Path(out).expanduser()
    if requested_out.is_symlink():
        raise ValueError("导出路径不能是符号链接")
    out_path = requested_out.resolve()
    if out_path == root or root in out_path.parents:
        raise ValueError("导出路径不能位于任务目录内")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    columns = ["employee_id", "name"] + field_ids
    written = {}
    if "json" in formats:
        path = out_path.with_suffix(".json")
        dump_json(path, {"task_id": task["task_id"], "exported_at": now_iso(), "purpose": purpose, "recipient": recipient, "fields": columns, "masks": masks, "rows": rows})
        os.chmod(path, 0o600)
        written["json"] = str(path)
    if "xlsx" in formats:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        path = out_path.with_suffix(".xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "明文导出"
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row_index, record in enumerate(rows, start=2):
            for column_index, key in enumerate(columns, start=1):
                cell = sheet.cell(row_index, column_index, str(record.get(key, "")))
                cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(path)
        os.chmod(path, 0o600)
        written["xlsx"] = str(path)
    return written


def cmd_export_clear(args) -> dict[str, Any]:
    require_tty("export-clear")
    root, task = load_task(args.task_dir)
    require_current_format(task, "明文导出")
    if task_expired(task):
        raise RuntimeError("任务已超过保存期限，禁止明文导出；请执行清理")
    field_ids = parse_export_fields(task, getattr(args, "fields", None))
    masks = parse_mask_specs(getattr(args, "mask", None))
    unknown_mask = set(masks) - ({"employee_id", "name"} | set(field_ids))
    if unknown_mask:
        raise ValueError(f"脱敏规则指向未导出的字段: {sorted(unknown_mask)}")
    formats = sorted({part.strip().lower() for value in (args.formats or ["xlsx"]) for part in value.split(",") if part.strip()})
    password = getpass.getpass("任务密码: ")
    try:
        private_key = unlock_private_key(root, password)
    except Exception as exc:
        raise RuntimeError("任务密码错误或私钥损坏") from exc
    rows = collect_clear_rows(root, task, private_key, field_ids, masks)
    purpose, recipient = (getattr(args, "purpose", "") or "").strip(), (getattr(args, "recipient", "") or "").strip()
    print("\n=== 明文导出确认 ===")
    print(f"任务: {task['task_id']}（{task['title']}）")
    print(f"导出字段: {', '.join(['employee_id', 'name'] + field_ids)}")
    print(f"脱敏规则: {', '.join(f'{key}={rule}' for key, rule in sorted(masks.items())) or '无（白名单字段全部明文）'}")
    print(f"数据行数: {len(rows)}")
    print(f"用途: {purpose or '（未填写）'}")
    print(f"接收方: {recipient or '（未填写）'}")
    typed = input(f"输入任务 ID {task['task_id']} 以确认导出明文: ").strip()
    if typed != task["task_id"]:
        raise RuntimeError("任务 ID 不匹配，已取消导出")
    written = write_clear_export(root, task, rows, field_ids, masks, args.out, formats, purpose, recipient)
    print("警告：明文已落盘，用后请删除，系统无法管控后续传播。", file=sys.stderr)
    return {"task_id": task["task_id"], "rows": len(rows), **written}


def build_task_package(task_dir: str | Path) -> tuple[bytes, dict[str, Any]]:
    """在内存中构建任务交接 ZIP（明文，含名单标识），供加密导出或测试复用。"""
    root, task = load_task(task_dir)
    if task_expired(task):
        raise RuntimeError("任务已超过保存期限，禁止导出；请执行清理")
    if any(path.is_symlink() for path in root.rglob("*")):
        raise ValueError("任务目录包含符号链接，拒绝导出")
    candidates = [root / name for name in ("task.json", "public.pem", "private.pem.enc", "roster.csv", "invite-index.csv", "state.sqlite3")]
    if any(not path.is_file() for path in candidates):
        raise ValueError("任务目录缺少必要文件")
    candidates += list((root / "invites").glob("*.html"))
    candidates += list((root / "submissions").glob("*/*.yintian"))
    candidates += [path for path in (root / "reports").glob("progress.*") if path.suffix in {".json", ".xlsx"}]
    files = [path for path in sorted(candidates) if path.is_file() and not path.is_symlink()]
    sizes = [path.stat().st_size for path in files]
    if len(files) > MAX_PACKAGE_FILES or any(size > MAX_PACKAGE_MEMBER_BYTES for size in sizes) or sum(sizes) > MAX_PACKAGE_BYTES:
        raise ValueError("任务内容超过交接包安全上限")
    manifest = {path.relative_to(root).as_posix(): sha256_file(path) for path in files}
    metadata = {"package_version": TASK_PACKAGE_VERSION, "task_id": task["task_id"], "created_at": now_iso(), "manifest": manifest}
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("package.json", json.dumps(metadata, ensure_ascii=False, indent=2))
        for path in files:
            archive.write(path, task["task_id"] + "/" + path.relative_to(root).as_posix())
    return buffer.getvalue(), {"task_id": task["task_id"], "files": len(files)}


def task_package_aad(task_id: str) -> bytes:
    return f"{ENCRYPTED_TASK_PACKAGE_VERSION}:{task_id}".encode("utf-8")


def export_task(task_dir: str | Path, out: str | Path, handoff_password: str | None = None) -> dict[str, Any]:
    root, task = load_task(task_dir)
    zip_bytes, info = build_task_package(root)
    password = handoff_password if handoff_password is not None else generate_password()
    envelope = aes_gcm_seal(ENCRYPTED_TASK_PACKAGE_VERSION, zip_bytes, password, aad=task_package_aad(task["task_id"]))
    envelope["task_id"] = task["task_id"]
    requested_out = Path(out).expanduser()
    if requested_out.is_symlink():
        raise ValueError("导出路径不能是符号链接")
    out_path = requested_out.resolve()
    if out_path == root or root in out_path.parents:
        raise ValueError("导出路径不能位于任务目录内")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=out_path.parent, delete=False) as stream:
        temp_path = Path(stream.name)
    try:
        temp_path.write_bytes(json.dumps(envelope, ensure_ascii=False, indent=2).encode("utf-8") + b"\n")
        os.chmod(temp_path, 0o600)
        temp_path.replace(out_path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise
    return {"task_id": task["task_id"], "package": str(out_path), "files": info["files"], "handoff_password": password}


def cmd_export(args) -> dict[str, Any]:
    require_tty("export-task")
    result = export_task(args.task_dir, args.out)
    password = result.pop("handoff_password")
    print_once_secret("交接密码（仅显示一次，丢失不可恢复）：", password, "请通过另一独立安全渠道告知接收方；不要与交接包同渠道发送，也不要写入任务目录或聊天提示词。")
    return result


def read_package_member(archive: zipfile.ZipFile, name: str, remaining: int) -> bytes:
    """流式读取任务包成员并按实际字节计数，超过单项或累计上限立即失败。"""
    data = bytearray()
    with archive.open(name) as stream:
        while True:
            chunk = stream.read(1024 * 1024)
            if not chunk:
                return bytes(data)
            data += chunk
            if len(data) > MAX_PACKAGE_MEMBER_BYTES or len(data) > remaining:
                raise ValueError("任务包解压后超过安全上限")


def open_task_package(package_path: Path, handoff_password: str | None = None) -> tuple[zipfile.ZipFile, bool]:
    """打开交接包，返回 (archive, encrypted)。

    yintian-task/3 加密信封先按 stat 尺寸拒绝超限文件，再解密认证；旧版明文 ZIP 打印警告后直接打开。
    嗅探窗口内全为空白且尺寸未超限的内容按 v3 信封处理，避免大段前导空白被误判为明文包。
    """
    size = package_path.stat().st_size
    with package_path.open("rb") as stream:
        prefix = stream.read(V3_SNIFF_BYTES).lstrip()
    if prefix.startswith(b"{") or not prefix:
        if size > MAX_V3_ENVELOPE_BYTES:
            raise ValueError("交接包信封超过 200MB 安全上限，拒绝读取")
        try:
            envelope = json.loads(package_path.read_bytes())
        except (json.JSONDecodeError, UnicodeDecodeError) as exc:
            raise ValueError("交接包信封无效") from exc
        if not isinstance(envelope, dict) or envelope.get("format") != ENCRYPTED_TASK_PACKAGE_VERSION:
            raise ValueError("交接包格式不受支持")
        task_id = envelope.get("task_id")
        if not isinstance(task_id, str) or not TASK_ID_RE.fullmatch(task_id):
            raise ValueError("交接包 task_id 无效")
        password = handoff_password if handoff_password is not None else getpass.getpass("交接密码: ")
        try:
            zip_bytes = aes_gcm_open(envelope, password, task_package_aad(task_id))
        except Exception as exc:
            raise ValueError("交接密码错误或交接包已被篡改，拒绝导入") from exc
        return zipfile.ZipFile(io.BytesIO(zip_bytes)), True
    print("警告：导入的是未加密的明文交接包（yintian-task/1 或 yintian-task/2），不能认证来源；仅在信任渠道下导入。", file=sys.stderr)
    return zipfile.ZipFile(package_path), False


def import_task(package: str | Path, out_parent: str | Path, handoff_password: str | None = None, confirm_plaintext=None) -> dict[str, Any]:
    package_path, parent = Path(package).expanduser().resolve(), Path(out_parent).expanduser().resolve()
    if package_path.stat().st_size > MAX_PACKAGE_BYTES:
        raise ValueError("任务包超过安全上限")
    archive, encrypted = open_task_package(package_path, handoff_password)
    with archive:
        package_info = archive.getinfo("package.json")
        if package_info.file_size > 1024 * 1024:
            raise ValueError("任务包清单过大")
        metadata = json.loads(read_package_member(archive, "package.json", 1024 * 1024))
        task_id = metadata.get("task_id")
        package_version = metadata.get("package_version")
        if package_version not in {LEGACY_TASK_PACKAGE_VERSION, TASK_PACKAGE_VERSION} or not isinstance(task_id, str) or not TASK_ID_RE.fullmatch(task_id):
            raise ValueError("任务包格式无效")
        if not encrypted and confirm_plaintext is not None and not confirm_plaintext(task_id):
            raise RuntimeError("未加密的明文交接包未经人工确认，已取消导入")
        raw_manifest = metadata.get("manifest")
        if not isinstance(raw_manifest, dict) or len(raw_manifest) > MAX_PACKAGE_FILES:
            raise ValueError("任务包清单无效或文件过多")
        required = {"task.json", "public.pem", "private.pem.enc", "roster.csv", "invite-index.csv", "state.sqlite3"}
        manifest: dict[str, str] = {}
        for raw_rel, expected_hash in raw_manifest.items():
            if not isinstance(raw_rel, str) or not raw_rel or not isinstance(expected_hash, str) or ":" in raw_rel:
                raise ValueError("任务包包含不安全路径")
            if package_version == TASK_PACKAGE_VERSION and "\\" in raw_rel:
                raise ValueError("任务包包含不安全路径")
            rel = raw_rel.replace("\\", "/")
            rel_path = PurePosixPath(rel)
            if rel_path.is_absolute() or ".." in rel_path.parts:
                raise ValueError("任务包包含不安全路径")
            parts = rel_path.parts
            allowed = (
                rel in required
                or (len(parts) == 2 and parts[0] == "invites" and parts[1].endswith(".html") and INVITE_ID_RE.fullmatch(parts[1][:-5]))
                or (len(parts) == 3 and parts[0] == "submissions" and INVITE_ID_RE.fullmatch(parts[1]) and re.fullmatch(r"v\d{4}_[0-9a-f]{12}\.yintian", parts[2]))
                or rel in {"reports/progress.json", "reports/progress.xlsx"}
            )
            if not allowed or rel in manifest:
                raise ValueError("任务包包含不允许或重复的文件")
            manifest[rel] = expected_hash
        if not required.issubset(manifest):
            raise ValueError("任务包缺少必要文件")
        target = parent / task_id
        if target.exists():
            raise FileExistsError(f"目标任务已存在: {target}")
        expected_prefix = task_id + "/"
        archive_names: dict[str, str] = {}
        for info in archive.infolist():
            if package_version == TASK_PACKAGE_VERSION and "\\" in info.filename:
                raise ValueError("任务包包含不安全路径")
            normalized = info.filename.replace("\\", "/")
            if normalized in archive_names:
                raise ValueError("任务包包含重复条目")
            archive_names[normalized] = info.filename
        names = set(archive_names)
        expected_names = {expected_prefix + rel for rel in manifest}
        if names != expected_names | {"package.json"}:
            raise ValueError("任务包文件不完整或包含额外条目")
        infos = [archive.getinfo(archive_names[name]) for name in expected_names]
        if any(info.file_size > MAX_PACKAGE_MEMBER_BYTES for info in infos) or sum(info.file_size for info in infos) > MAX_PACKAGE_BYTES:
            raise ValueError("任务包解压后超过安全上限")
        target.mkdir(parents=True, mode=0o700)
        try:
            total_size = 0
            for rel, expected_hash in manifest.items():
                destination = (target.joinpath(*PurePosixPath(rel).parts)).resolve()
                if target != destination and target not in destination.parents:
                    raise ValueError("任务包包含越界路径")
                data = read_package_member(archive, archive_names[expected_prefix + rel], MAX_PACKAGE_BYTES - total_size)
                total_size += len(data)
                if sha256_bytes(data) != expected_hash:
                    raise ValueError(f"任务包哈希不匹配: {rel}")
                atomic_write(destination, data)
            _, imported_task = load_task(target)
            if imported_task["task_id"] != task_id:
                raise ValueError("任务包内外 task_id 不一致")
            from cryptography.hazmat.primitives import serialization

            public_key = serialization.load_pem_public_key((target / "public.pem").read_bytes())
            public_der = public_key.public_bytes(serialization.Encoding.DER, serialization.PublicFormat.SubjectPublicKeyInfo)
            if hashlib.sha256(public_der).hexdigest()[:24] != imported_task["key_id"] or not valid_private_key_blob((target / "private.pem.enc").read_bytes()):
                raise ValueError("任务包密钥材料无效")
            roster = read_roster(target / "roster.csv")
            with (target / "invite-index.csv").open("r", encoding="utf-8-sig", newline="") as stream:
                index = list(csv.DictReader(stream))
            if len(index) != len(roster) or not index or any(not INVITE_ID_RE.fullmatch(row.get("invite_id", "")) for row in index):
                raise ValueError("任务包名单或邀请索引无效")
            if status_task(target)["total"] != len(roster):
                raise ValueError("任务包状态数据库与名单不一致")
            if imported_task["format_version"] == FORMAT_VERSION:
                with closing(connect_db(target)) as db, db:
                    columns = {row[1] for row in db.execute("PRAGMA table_info(invites)")}
                if "token_hash" not in columns:
                    raise ValueError("v2 任务缺少邀请认证数据")
        except Exception:
            shutil.rmtree(target, ignore_errors=True)
            raise
    if os.name != "nt":
        for directory in [target, *(path for path in target.rglob("*") if path.is_dir())]:
            os.chmod(directory, 0o700)
    return {"task_id": task_id, "task_dir": str(target)}


def confirm_plaintext_import(task_id: str) -> bool:
    typed = input(f"该交接包未加密且无法认证来源；输入任务 ID {task_id} 确认导入: ").strip()
    return typed == task_id


def cmd_import(args) -> dict[str, Any]:
    require_tty("import-task")
    return import_task(args.package, args.out, confirm_plaintext=confirm_plaintext_import)


def cmd_purge(args) -> dict[str, Any]:
    require_tty("purge")
    root, task = load_task(args.task_dir)
    if not task_expired(task) and not args.allow_early:
        raise RuntimeError("任务尚未到保存期限；如确需提前删除，请增加 --allow-early")
    typed = input(f"输入任务 ID {task['task_id']} 以确认删除任务目录: ").strip()
    if typed != task["task_id"]:
        raise RuntimeError("任务 ID 不匹配，已取消")
    parent = root.parent
    with task_lock(root):
        status = status_task(root)
        summary = {"task_id": task["task_id"], "purged_at": now_iso(), "aggregate_counts": status["counts"], "total": status["total"]}
        shutil.rmtree(root)
    summary_path = parent / f"{task['task_id']}.purged.json"
    dump_json(summary_path, summary)
    return {"summary": str(summary_path)}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="隐填：端到端加密的私密信息收集管理")
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("init-config", help="生成基础身份信息收集配置")
    p.add_argument("--out", required=True); p.add_argument("--force", action="store_true"); p.set_defaults(func=cmd_init_config)
    p = sub.add_parser("create", help="创建任务并批量生成离线邀请")
    p.add_argument("--roster", required=True); p.add_argument("--config", required=True); p.add_argument("--out", required=True)
    p.add_argument("--mode", choices=["directed", "group"], default="directed", help="directed 逐人定向邀请（默认）；group 单份表单群发、无个人令牌，可冒名提交，仅靠复核名单比对兜底")
    p.set_defaults(func=cmd_create)
    p = sub.add_parser("ingest", help="接收 .yintian 密文提交")
    p.add_argument("task_dir"); p.add_argument("submissions_dir"); p.set_defaults(func=cmd_ingest)
    p = sub.add_parser("review", help="本地解密、校验和 OpenVINO OCR 复核")
    p.add_argument("task_dir"); p.set_defaults(func=cmd_review)
    p = sub.add_parser("status", help="查看数据最小化统计")
    p.add_argument("task_dir"); p.set_defaults(func=cmd_status)
    p = sub.add_parser("report", help="生成保留姓名/员工号的数据最小化 XLSX/JSON 报告")
    p.add_argument("task_dir"); p.add_argument("--formats", nargs="+", default=["xlsx", "json"]); p.set_defaults(func=cmd_report)
    p = sub.add_parser("reveal", help="人工终端按需查看一份明文")
    p.add_argument("task_dir"); p.add_argument("invite_id"); p.set_defaults(func=cmd_reveal)
    p = sub.add_parser("export-task", help="导出加密的敏感任务交接包（yintian-task/3，一次性交接密码仅显示一次）")
    p.add_argument("task_dir"); p.add_argument("--out", required=True); p.set_defaults(func=cmd_export)
    p = sub.add_parser("import-task", help="导入任务交接包（加密包需交互输入交接密码；v1/v2 明文旧包需输入任务 ID 二次确认）")
    p.add_argument("package"); p.add_argument("--out", required=True); p.set_defaults(func=cmd_import)
    p = sub.add_parser("purge", help="删除任务目录；不保证擦除备份或磁盘残留")
    p.add_argument("task_dir"); p.add_argument("--allow-early", action="store_true"); p.set_defaults(func=cmd_purge)
    p = sub.add_parser(
        "export-clear",
        help="【仅限人类授权操作员在本人终端运行，禁止 Agent 代跑】把白名单字段的当前有效提交解密导出为明文 XLSX/JSON（需任务密码并输入任务 ID 二次确认）",
        description="仅限人类授权操作员在本人交互终端运行，禁止 Agent 代跑或旁观。明文导出是最后一道出口：只导出 --fields 白名单字段（name/employee_id 默认附带），落盘 0600，系统无法管控后续传播。",
    )
    p.add_argument("task_dir")
    p.add_argument("--out", required=True, help="导出文件路径（多格式时按后缀派生同名 .xlsx/.json）")
    p.add_argument("--fields", required=True, help="逗号分隔的导出字段 id 白名单，仅限非附件字段；name/employee_id 默认附带")
    p.add_argument("--mask", action="append", default=[], metavar="字段id=规则", help="脱敏规则 last4（保留后四位）或 mid4（保留前3后4），可重复；未声明字段明文导出")
    p.add_argument("--purpose", default="", help="本次明文导出的用途，回显确认并写入 JSON 元数据")
    p.add_argument("--recipient", default="", help="明文接收方，回显确认并写入 JSON 元数据")
    p.add_argument("--formats", nargs="+", default=["xlsx"], help="导出格式：xlsx、json（默认 xlsx）")
    p.set_defaults(func=cmd_export_clear)
    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()
    try:
        result = args.func(args)
        if result is not None:
            print(json.dumps(result, ensure_ascii=False, indent=2))
    except Exception as exc:
        print(f"错误: {exc}", file=sys.stderr)
        raise SystemExit(1)


if __name__ == "__main__":
    main()
